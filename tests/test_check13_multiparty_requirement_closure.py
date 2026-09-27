"""CHECK-1.3: close the final multi-party requirement gap.

CHECK-1.2's DOCX content-control fix exposed the real buyer sentence
"Each proposal ... by a Multi-Party Team must include a Multi-Party
Confirmation Form ..." but Fast Analysis run 34 (bid 1360, COMPLETE,
immutable) was extracted before that fix and has no canonical requirement
for it. CHECK-1.3 recovers it deterministically (zero model calls) through
the EXISTING canonical recomputation path
(`full_analysis.build_canonical_package`, fed by `analysis_service.
build_full_analysis_package`) -- appended, never mutating run 34.

Section A is synthetic (always runs). Section B is the REAL acceptance
chain -- real City RFP passage -> canonical requirement -> MULTI_PARTY_FORM
-> real Phoenix / Inquisitive Talent B2 evidence -- using the real run 34
raw snapshot (tests/fixtures/calgary_26_1603_run34_raw_snapshot.json,
byte-for-byte the persisted snapshot; buyer-side public RFP content only)
and the real buyer/bidder ZIPs, skipped (never faked) when the ZIPs are not
on this machine.

Zero provider calls (every Anthropic entry point is poisoned for the whole
module) and zero database I/O (fakes only).
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import os
import sys
import zipfile

import docx
import pytest
from docx.oxml import parse_xml

sys.path.insert(0, os.path.dirname(__file__))

import canonical_procurement as cp  # noqa: E402
import extractor  # noqa: E402
import fast_analysis  # noqa: E402
import full_analysis  # noqa: E402
import procurement_normalization as pn  # noqa: E402
import submission_package as sp  # noqa: E402

_ORIGINAL_RUN_FAST_ANALYSIS = fast_analysis.run_fast_analysis_corpus

HERE = os.path.dirname(__file__)
RUN34 = os.path.join(HERE, "fixtures", "calgary_26_1603_run34_raw_snapshot.json")
#: sha256 of the persisted run 34 raw snapshot as read live (sorted-key
#: JSON) on 2026-09-27, before and after CHECK-1.3's live acceptance.
RUN34_SNAPSHOT_SHA256 = "45543c0a3b517f2f697700d2cd7d81a52e387b7fa0e0a8b2f8406e3f718f2a51"
BIDDER_ZIP = os.getenv("CHECK11_BIDDER_ZIP", r"C:\Users\feras\Downloads\OneDrive_2_9-27-2026.zip")
BUYER_ZIP = os.getenv("CHECK11_BUYER_ZIP", r"C:\Users\feras\Downloads\Calgary RFP.zip")
BID, ORG = 1360, "org-phoenix"

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MP_SENTENCE = ("Each proposal that is submitted on behalf of, and contemplates the provision of the "
               "Deliverables by a Multi-Party Team must include a Multi-Party Confirmation Form completed "
               "and signed by all Team Members.")


@pytest.fixture(autouse=True)
def _no_provider_calls(monkeypatch):
    def refuse(*_a, **_k):
        raise AssertionError("CHECK-1.3 tests must never reach a model provider")
    import config
    for mod in (config, fast_analysis, full_analysis):
        for attr in ("get_anthropic_client", "execute_messages_create"):
            if hasattr(mod, attr):
                monkeypatch.setattr(mod, attr, refuse)
    monkeypatch.setattr(fast_analysis, "run_fast_analysis_corpus", refuse)


def _rfp_docx(*, in_content_control: bool = True, extra: list[str] = ()) -> bytes:
    d = docx.Document()
    d.add_paragraph("APPENDIX F - OTHER ATTACHMENTS", style="Heading 1")
    anchor = d.add_paragraph("B2:  MULTI-PARTY CONFIRMATION FORM")
    para = f'<w:p xmlns:w="{W_NS}"><w:r><w:t>{MP_SENTENCE}</w:t></w:r></w:p>'
    if in_content_control:
        para = (f'<w:sdt xmlns:w="{W_NS}"><w:sdtPr><w:alias w:val="cc"/></w:sdtPr>'
                f'<w:sdtContent>{para}</w:sdtContent></w:sdt>')
    anchor._p.addnext(parse_xml(para))
    d.add_paragraph("THIS FORM TO BE COMPLETED ONLY IF THE PROPOSAL IS SUBMITTED BY A TEAM OF PROPONENTS.")
    for p in extra:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _corpus(name="RFP with Price - V2.5_Services.docx", **kw):
    text, _ = extractor.extract_docx_with_metadata(_rfp_docx(**kw), name)
    return [(name, text)]


EXISTING = [
    {"category": "Mandatory", "source_doc": "RFP.docx",
     "description": "Proponent must provide complete Conflicts of Interest table in Submission Form (Appendix E)."},
    {"category": "Supporting", "source_doc": "Price.xlsx", "description": "Pricing-calculation formula ..."},
]


def _result(requirements):
    return fast_analysis.FastAnalysisResult(requirements=copy.deepcopy(requirements))


# ═══════════════════════════════════════════════════════════════════════
# A. Synthetic: recovered buyer passage -> canonical requirement
# ═══════════════════════════════════════════════════════════════════════

class TestRecoveredPassageBecomesCanonicalRequirement:
    def test_content_control_passage_is_recovered_verbatim(self):
        rec = pn.recover_uncovered_required_form_obligations(_corpus(), EXISTING)
        assert len(rec) == 1
        r = rec[0]
        assert r["description"] == MP_SENTENCE
        assert r["category"] == "Mandatory"
        assert r["semantic_type"] == cp.SEMANTIC_SUBMISSION_REQUIREMENT
        assert r["required_form"] == "Multi-Party Confirmation Form"
        assert r["requirement_origin"] == pn.REQUIREMENT_ORIGIN_REQUIRED_FORM_RECOVERY
        assert r["applicability_condition"].startswith("Each proposal that is submitted on behalf of")

    def test_exact_buyer_source_provenance(self):
        name = "RFP with Price - V2.5_Services.docx"
        r = pn.recover_uncovered_required_form_obligations(_corpus(name), EXISTING)[0]
        assert r["source_doc"] == name
        assert r["source_refs"] == [{"page": None, "sheet": None, "section": "APPENDIX F - OTHER ATTACHMENTS",
                                     "excerpt": MP_SENTENCE, "source_doc": name}]
        loc = r["source_locator"]
        assert loc["heading"] == "B2:  MULTI-PARTY CONFIRMATION FORM"
        assert loc["section"] == "APPENDIX F - OTHER ATTACHMENTS"
        assert _corpus(name)[0][1].splitlines()[loc["line_index"]].strip() == MP_SENTENCE
        assert r["source_context"]["following_text"].startswith("THIS FORM TO BE COMPLETED ONLY IF")
        assert "evidence_id" not in json.dumps(r)

    def test_canonical_semantics_match_the_rest_of_understand(self):
        canon = pn.canonicalize_requirements(
            pn.recover_uncovered_required_form_obligations(_corpus(), EXISTING))
        assert len(canon) == 1
        # same canonical fields every other UNDERSTAND requirement carries
        for key in ("source_variants", "source_docs", "source_refs_all", "duplicate_count",
                    "applicability", "applicable_category_ids", "applicability_basis", "semantic_type"):
            assert key in canon[0]
        # a document VERSION ("V2.5") is not a category; the obligation is submission-wide
        assert canon[0]["applicability"] == cp.APPLICABILITY_SUBMISSION_WIDE
        assert canon[0]["applicable_category_ids"] == []

    def test_expected_role_is_multi_party_form(self):
        r = pn.recover_uncovered_required_form_obligations(_corpus(), EXISTING)[0]
        ee = sp.derive_expected_evidence(r)
        assert ee.roles[0] == sp.ROLE_MULTI_PARTY_FORM
        assert ee.evidence_category == sp.EVIDENCE_CATEGORY_MULTI_PARTY and not ee.flexible

    def test_covered_form_is_never_duplicated_and_recovery_is_idempotent(self):
        covered = EXISTING + [{"category": "Mandatory", "description": "Include the multi-party confirmation form."}]
        assert pn.recover_uncovered_required_form_obligations(_corpus(), covered) == []
        rec = pn.recover_uncovered_required_form_obligations(_corpus(), EXISTING)
        assert pn.recover_uncovered_required_form_obligations(_corpus(), EXISTING + rec) == []

    def test_only_normative_named_form_obligations_are_recovered(self):
        texts = [("x.docx", "The proponent must not include a Marketing Brochure Form.\n"
                             "Proposals must be submitted in electronic form.\n"
                             "A Reference Form may be requested later.\n"
                             "[[SOURCE: x.docx | TABLE]] Each proposal must include a Table Row Form.")]
        assert pn.recover_uncovered_required_form_obligations(texts, []) == []

    def test_version_token_is_not_a_category_but_form_ids_still_are(self):
        assert cp._category_tokens("RFP WITH PRICE - V2.5_DESIGN.DOCX") == set()
        assert cp._category_tokens("ADDENDUM ONE V4.0.PDF") == set()
        assert cp._category_tokens("APPENDIX D1.1 CORPORATE PROFILE") == {"D1"}
        assert cp._category_tokens("APPENDIX C2 - MINIMUM QUALIFICATION.XLSX") == {"C2"}


class TestCanonicalPackageRecomputation:
    def test_recovered_requirement_is_appended_and_existing_ids_unchanged(self):
        base = full_analysis.build_canonical_package(_result(EXISTING))
        withdocs = full_analysis.build_canonical_package(_result(EXISTING), documents=_corpus())
        assert [r["canonical_id"] for r in base.requirements] == ["REQ-0", "REQ-1"]
        assert list(withdocs.requirements[:2]) == list(base.requirements)
        new = withdocs.requirements[2]
        assert new["canonical_id"] == "REQ-2" and new["canonical_id"] in withdocs.canonical_ids
        assert new["requirement_origin"] == pn.REQUIREMENT_ORIGIN_REQUIRED_FORM_RECOVERY
        assert new["description"] == MP_SENTENCE and new["category"] == "Mandatory"

    def test_historical_result_object_is_never_mutated(self):
        result = _result(EXISTING)
        before = copy.deepcopy(result.requirements)
        full_analysis.build_canonical_package(result, documents=_corpus())
        assert result.requirements == before

    def test_no_documents_means_no_recovery_and_unchanged_digest(self):
        a = full_analysis.build_canonical_package(_result(EXISTING))
        b = full_analysis.build_canonical_package(_result(EXISTING), documents=None)
        assert a.package_digest == b.package_digest and len(a.requirements) == 2

    def test_fresh_fast_analysis_step5_recovers_it_deterministically(self, monkeypatch):
        """A FUTURE run's step 5 applies the same recovery. Exercised with
        every extraction call stubbed (canned model output, no provider):
        the deterministic step-5 code path itself is what runs."""
        calls = []

        def fake_extract(name, text, api_key, *, route=None, client=None, telemetry=None):
            calls.append(route)
            return {"doc_metadata": {}, "typed_observations": [], "evaluation_criteria": [],
                    "requirements": copy.deepcopy(EXISTING), "commercial_clauses": []}

        monkeypatch.setattr(fast_analysis, "get_anthropic_client", lambda **_k: object())
        monkeypatch.setattr(fast_analysis, "extract_fast_document", fake_extract)
        monkeypatch.setattr(fast_analysis, "run_focused_task", lambda *a, **k: [])
        result = _ORIGINAL_RUN_FAST_ANALYSIS(_corpus(), "no-key")
        recovered = [r for r in result.requirements if r.get("requirement_origin")]
        assert len(recovered) == 1 and recovered[0]["description"] == MP_SENTENCE
        assert recovered[0]["applicability"] == cp.APPLICABILITY_SUBMISSION_WIDE  # canonicalized
        assert calls  # stubs only; the poisoned provider was never reached


class _FakeDB:
    """Read-only fake: any write attempt is an AttributeError."""

    def __init__(self, snapshot_payload, docs):
        self._payload, self._docs = snapshot_payload, docs
        self.reads = []

    def get_analysis_run(self, run_id):
        self.reads.append("run")
        return {"id": run_id, "bid_id": BID, "status": "COMPLETE", "analysis_mode": "FAST"}

    def get_analysis_result(self, run_id):
        self.reads.append("result")
        return {"fast_analysis_result_snapshot": self._payload}

    def get_documents(self, bid_id):
        return [dict(d) for d in self._docs]

    def download_file(self, path):
        return next(d["_bytes"] for d in self._docs if d["storage_path"] == path)


def _payload(requirements):
    return fast_analysis.serialize_fast_analysis_result(_result(requirements), engine_version="test")


class TestBuyerBidderSeparation:
    def _package_via_service(self, monkeypatch, doc_type):
        import analysis_service
        docs = [{"id": 1, "name": "B2 Multi-Party Confirmation Form_Phoenix.docx", "doc_type": doc_type,
                 "storage_path": "p/b2", "_bytes": _rfp_docx()}]
        fake = _FakeDB(_payload(EXISTING), docs)
        monkeypatch.setattr(analysis_service, "db", fake)
        pkg = analysis_service.build_full_analysis_package(34, include_documents=True)
        return pkg, fake

    def test_bidder_b2_content_cannot_create_a_buyer_requirement(self, monkeypatch):
        """The SAME B2 text, held by a non-buyer document (the bidder's own
        submitted form), is never read into the canonical buyer package --
        only 'RFP / Source' documents are the buyer corpus."""
        pkg, _ = self._package_via_service(monkeypatch, "Proposal")
        assert not [r for r in pkg.requirements if r.get("requirement_origin")]
        # control: identical bytes as a buyer RFP / Source document ARE recovered
        pkg2, _ = self._package_via_service(monkeypatch, "RFP / Source")
        assert len([r for r in pkg2.requirements if r.get("requirement_origin")]) == 1

    def test_recovery_rejects_bidder_evidence_objects(self):
        pkg = sp.build_submission_package([("B2 Multi-Party Confirmation Form.docx", _rfp_docx())],
                                          bid_id=BID, organization_id=ORG)
        with pytest.raises(TypeError):
            pn.recover_uncovered_required_form_obligations(list(pkg.registry), [])
        with pytest.raises(TypeError):
            pn.recover_uncovered_required_form_obligations(list(pkg.documents), [])

    def test_ids_stay_on_their_own_side_and_cross_bid_is_rejected(self):
        cpkg = full_analysis.build_canonical_package(_result(EXISTING), documents=_corpus())
        req = next(r for r in cpkg.requirements if r.get("requirement_origin"))
        bidder = sp.build_submission_package([("B2 Multi-Party Confirmation Form_Phoenix.docx", _rfp_docx())],
                                             bid_id=BID, organization_id=ORG)
        m = sp.map_requirement_to_submission(req, bidder)
        assert m.req_id == req["canonical_id"] and m.candidates
        # buyer id is not bidder evidence ...
        with pytest.raises(KeyError):
            bidder.registry.get(req["canonical_id"], bid_id=BID)
        # ... and bidder evidence ids are not canonical buyer ids
        assert not {c.evidence_id for c in m.candidates} & set(cpkg.canonical_ids)
        # buyer provenance never points at bidder evidence
        assert set(req["source_docs"]) == {"RFP with Price - V2.5_Services.docx"}
        assert not ({d.filename for d in bidder.documents} & set(req["source_docs"]))
        with pytest.raises(sp.CrossBidEvidenceError):
            bidder.registry.get(m.candidates[0].evidence_id, bid_id=BID + 1)
        assert bidder.registry.filter_valid([m.candidates[0].evidence_id], bid_id=BID + 1) == []

    def test_service_path_performs_reads_only(self, monkeypatch):
        pkg, fake = self._package_via_service(monkeypatch, "RFP / Source")
        assert set(fake.reads) == {"run", "result"}   # no write method exists on the fake at all


# ═══════════════════════════════════════════════════════════════════════
# B. REAL Calgary acceptance chain (real run 34 snapshot + real ZIPs)
# ═══════════════════════════════════════════════════════════════════════

def _run34_payload() -> dict:
    with open(RUN34, encoding="utf-8") as fh:
        return json.load(fh)


def test_run34_fixture_is_the_persisted_immutable_snapshot():
    payload = _run34_payload()
    digest = hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()
    assert digest == RUN34_SNAPSHOT_SHA256
    result = fast_analysis.deserialize_fast_analysis_result(payload)
    assert len(result.requirements) == 46
    assert not any("Multi-Party Confirmation Form" in (r.get("description") or "") for r in result.requirements)


needs_real = pytest.mark.skipif(not (os.path.exists(BUYER_ZIP) and os.path.exists(BIDDER_ZIP)),
                                reason="real Calgary buyer/bidder ZIPs not on this machine")


@pytest.fixture(scope="module")
def real_chain():
    payload = _run34_payload()
    result = fast_analysis.deserialize_fast_analysis_result(payload)
    before = copy.deepcopy(payload)
    documents = []
    with zipfile.ZipFile(BUYER_ZIP) as z:
        for n in z.namelist():
            if n.endswith("/"):
                continue
            base = n.rsplit("/", 1)[-1]
            text, _ = extractor.extract_document_with_metadata(z.read(n), base)
            documents.append((base, text))
    package = full_analysis.build_canonical_package(result, bid_id=BID, analysis_run_id=34, documents=documents)
    with open(BIDDER_ZIP, "rb") as fh:
        bidder = sp.build_submission_package([(os.path.basename(BIDDER_ZIP), fh.read())],
                                             bid_id=BID, organization_id=ORG)
    return {"payload_before": before, "payload": payload, "result": result, "documents": documents,
            "package": package, "bidder": bidder}


@needs_real
class TestRealCalgaryAcceptanceChain:
    def test_authoritative_city_passage_to_canonical_requirement(self, real_chain):
        pkg = real_chain["package"]
        recovered = [r for r in pkg.requirements if r.get("requirement_origin")]
        assert len(recovered) == 1 and len(pkg.requirements) == 47
        req = recovered[0]
        assert req["canonical_id"] == "REQ-46"
        assert req["description"] == MP_SENTENCE
        rfp = next(n for n, _ in real_chain["documents"] if n.startswith("S-PT-073-RFP with Price"))
        assert req["source_docs"] == [rfp]
        assert req["source_locator"]["heading"] == "B2:  MULTI-PARTY CONFIRMATION FORM"
        assert req["source_locator"]["section"].startswith("APPENDIX F")
        assert req["source_refs"][0]["excerpt"] == MP_SENTENCE
        assert req["category"] == "Mandatory" and req["semantic_type"] == cp.SEMANTIC_SUBMISSION_REQUIREMENT
        assert req["applicability"] == cp.APPLICABILITY_SUBMISSION_WIDE
        assert req["source_context"]["following_text"].startswith("THIS FORM TO BE COMPLETED ONLY IF")
        # the passage really is content-control-borne in the buyer DOCX
        with zipfile.ZipFile(BUYER_ZIP) as z:
            name = next(n for n in z.namelist() if rfp in n)
            _, meta = extractor.extract_docx_with_metadata(z.read(name), rfp)
        block = next(b for b in meta["blocks"] if (b.get("excerpt") or "").startswith(MP_SENTENCE[:100]))
        assert block["content_control"] is True

    def test_existing_run34_requirements_unchanged_and_run34_not_mutated(self, real_chain):
        pkg, result = real_chain["package"], real_chain["result"]
        assert real_chain["payload"] == real_chain["payload_before"]
        for i, raw in enumerate(result.requirements):
            assert pkg.requirements[i]["canonical_id"] == f"REQ-{i}"
            assert "requirement_origin" not in pkg.requirements[i]
        assert len(result.requirements) == 46

    def test_expected_role_and_real_b2_candidates(self, real_chain):
        req = next(r for r in real_chain["package"].requirements if r.get("requirement_origin"))
        bidder = real_chain["bidder"]
        assert sp.derive_expected_evidence(req).roles[0] == sp.ROLE_MULTI_PARTY_FORM
        m = sp.map_requirement_to_submission(req, bidder)
        b2 = bidder.documents_with_roles([sp.ROLE_MULTI_PARTY_FORM], include_secondary=False)
        assert len(b2) == 1
        b2 = b2[0]
        assert b2.filename == "B2 Multi-Party Confirmation Form_Phoenix Consulting Canada & Inquisitive Talent.pdf"
        assert m.req_id == "REQ-46" and m.artifact_status == sp.ARTIFACT_PRESENT and m.primary_role_present
        assert m.candidates and all(c.submission_document_id == b2.submission_document_id for c in m.candidates)
        for c in m.candidates:
            item = bidder.registry.get(c.evidence_id, bid_id=BID)
            assert item.document_role == sp.ROLE_MULTI_PARTY_FORM
            assert item.provenance["filename"] == b2.filename
            assert item.provenance["content_hash"] == b2.content_hash
            assert item.provenance["package_path"].endswith("Appendices Submission_Phoenix Consulting Canada.zip/" + b2.filename)
        text = " ".join(i.content for i in bidder.registry.for_document(b2.submission_document_id))
        assert "Phoenix Consulting Canada" in text and "Inquisitive Talent" in text
        # candidate mapping only -- no adjudication vocabulary anywhere
        dumped = json.dumps(m.to_dict())
        for verdict in ("ADDRESSED", "PARTIALLY_ADDRESSED", "COMPLIANT", "NON_COMPLIANT"):
            assert f'"{verdict}"' not in dumped

    def test_bidder_b2_never_originates_the_buyer_requirement(self, real_chain):
        bidder = real_chain["bidder"]
        req = next(r for r in real_chain["package"].requirements if r.get("requirement_origin"))
        assert not ({d.filename for d in bidder.documents} & set(req["source_docs"]))
        # the recovery over the real run 34 inputs WITHOUT the buyer RFP yields nothing,
        # even though the bidder's B2 PDF restates the same sentence
        no_rfp = [(n, t) for n, t in real_chain["documents"] if not n.startswith("S-PT-073-RFP with Price")]
        assert pn.recover_uncovered_required_form_obligations(no_rfp, real_chain["result"].requirements) == []
        with pytest.raises(TypeError):
            pn.recover_uncovered_required_form_obligations(list(bidder.registry), [])

    def test_calgary_benchmark_regression(self, real_chain):
        pkg, bidder = real_chain["package"], real_chain["bidder"]
        assert len(bidder.documents) == 13
        members = bidder.member_documents()
        assert len(members) == 4
        assert len(bidder.registry) == 294
        roles = {d.document_role: d.filename for d in members}
        assert roles[sp.ROLE_TECHNICAL_PROPOSAL].startswith("Appendix C")
        assert roles[sp.ROLE_PRICING_FORM].startswith("Appendix D_Price Form_26-1603_Phoenix")
        assert roles[sp.ROLE_SUBMISSION_FORM].startswith("APPENDIX E")
        assert roles[sp.ROLE_MULTI_PARTY_FORM].startswith("B2 Multi-Party Confirmation Form")
        assert len(pkg.scoped_criteria) == 13
        assert pkg.submission_mechanics["submission_deadline"] == "2026-07-16"
        dates = {m["normalized_date_start"] for m in pkg.milestones}
        assert {"2026-07-07", "2026-07-14", "2026-07-16"} <= dates
