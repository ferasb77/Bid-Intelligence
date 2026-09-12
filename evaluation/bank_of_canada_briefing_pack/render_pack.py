from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).resolve().parent / "BANK_OF_CANADA_BID_INTELLIGENCE_BRIEFING_PACK.md"
OUTPUT = ROOT / "output/documents/Bank_of_Canada_Bid_Intelligence_Briefing_Pack.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "5")
        item.set(qn("w:color"), "D9D9D9")
        el.append(item)
    tbl_pr.append(el)


def margins(cell):
    tc = cell._tc.get_or_add_tcPr()
    mar = tc.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc.append(mar)
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def add_inline(paragraph, text, color="262626"):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part.strip("*`") if (part.startswith("**") or part.startswith("`")) else part)
        run.bold = part.startswith("**")
        run.font.name = "Aptos"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for j, value in enumerate(rows[0]):
        cell = table.rows[0].cells[j]
        shade(cell, "17365D")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell)
        p = cell.paragraphs[0]
        add_inline(p, value, "FFFFFF")
        for r in p.runs:
            r.bold = True
    for i, values in enumerate(rows[1:], 1):
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cell = cells[j]
            if i % 2 == 0:
                shade(cell, "F3F6FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(cell)
            add_inline(cell.paragraphs[0], value)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Bid Intelligence Briefing Pack  |  ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(38, 38, 38)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title.element.get_or_add_pPr()
    for existing in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(existing)
    for idx, size in [(1, 19), (2, 14), (3, 11)]:
        st = styles[f"Heading {idx}"]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(12 if idx == 1 else 8)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True
    page_number(sec.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    # Purpose-built cover.
    doc.add_paragraph("BID INTELLIGENCE", style="Subtitle").runs[0].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph("Bank of Canada Bid Intelligence Briefing Pack", style="Title")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("RFP 2026-026")
    r.bold = True; r.font.size = Pt(17); r.font.color.rgb = RGBColor(23, 54, 93)
    doc.add_paragraph("Talent, Learning and Organizational Development Services").runs[0].font.size = Pt(14)
    p = doc.add_paragraph("Executive Opportunity Brief and Buyer Brief")
    p.paragraph_format.space_before = Pt(24)
    p.runs[0].font.size = Pt(12); p.runs[0].bold = True
    doc.add_paragraph("Prepared for proposal-team evaluation\nEvidence cut-off 8 September 2026")
    evaluation = doc.add_paragraph("Evaluation artifact", style="Subtitle")
    evaluation.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_page_break()

    # Skip source cover content through first horizontal rule.
    start = 0
    rules = [i for i, x in enumerate(lines) if x.strip() == "---"]
    if rules:
        start = rules[0] + 1
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line == "---":
            i += 1
            continue
        if line.startswith("# "):
            text = line[2:]
            if text.startswith("Volume") or text.startswith("Appendix"):
                doc.add_page_break()
            doc.add_paragraph(text, style="Heading 1")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 2")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 3")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in vals):
                    rows.append(vals)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        # Merge prose lines until a structural boundary.
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|---|\||- |\d+\. )", lines[i]):
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(para))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Bank of Canada Bid Intelligence Briefing Pack"
    doc.core_properties.subject = "Executive Opportunity Brief and Buyer Brief for RFP 2026-026"
    doc.core_properties.author = "Bid Intelligence"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
