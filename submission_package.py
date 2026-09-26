"""
submission_package.py -- CHECK-1: Package-Aware Proposal Assurance Foundation.

THE bidder-side canonical model for CHECK. Pure and deterministic: no
model call, no database/Storage I/O, no writes. It answers, for one
bidder SUBMISSION PACKAGE (never "one proposal document"):

  1. What artifacts were submitted, and what LOGICAL ROLE does each play?
     (technical proposal / pricing form / submission form / multi-party
     form / resume / certificate / ...)                -> SubmissionDocument
  2. What evidence does each artifact contain, at an exact, traceable
     location (page + section for prose, sheet + cell for spreadsheets,
     table row / label->value for forms)?              -> EvidenceItem
  3. ONE shared bidder-evidence registry every future CHECK analyzer
     references by `evidence_id`                        -> SubmissionEvidenceRegistry
  4. For each canonical RFP requirement: WHERE is its evidence expected
     (expected evidence ROLES, or flexible, or PORTAL_NATIVE), is an
     artifact of that role actually in the package, and which evidence
     items are CANDIDATES                              -> RequirementEvidenceMapping

It deliberately does NOT adjudicate. There is no ADDRESSED / PARTIAL /
MISSING verdict, no score, no claims assurance and no recommendation
anywhere in this module (CHECK-1 is ingestion, canonicalization and
candidate mapping only). The one package-level status it does emit --
`artifact_status` -- is about ARTIFACT PRESENCE, not requirement
compliance, and exists precisely to make the known failure class
("Price Form missing" because the narrative has no prices, when the
completed Price Form is a separate submitted file) structurally
impossible: an absence claim is only permitted when the status is
MISSING_FROM_PACKAGE, which by construction requires that NO document of
any expected role is present, readable or not, and that the location is
neither flexible nor possibly portal-native.

Boundaries (constitutional, not stylistic):

  * Buyer-side canonical procurement truth and bidder-side evidence stay
    distinct. Requirements are INPUT ONLY -- deep-copied, never mutated,
    never written anywhere. Nothing a bidder uploads can change a
    requirement (see tests/test_check1_submission_package.py security
    tests).
  * File identity is extractor.build_alignment_submission_package()'s
    own file_id / content_hash (reused, not re-derived): a
    submission_document_id IS that file_id.
  * Every evidence_id is derived from (bid_id, submission_document_id,
    kind, locator) -- an evidence id minted for one bid can never resolve
    under another bid (`SubmissionEvidenceRegistry.get` fails closed).
  * No concatenation of the whole package into a prompt. Structure is
    parsed deterministically per artifact; retrieval is bounded
    (`top_k`) and lexical/structural.

Reuse (task section 14): extractor (ZIP-safe unpacking, identity,
dedup), proposal_intelligence.compute_package_digest (the SAME package
identity migration 015's proposal_package_snapshots already uses),
canonical_procurement (semantic type / commercial topic /
scoped_criterion_map_key -- never a parallel vocabulary), and
section_drafting's evidence-id registry / fail-closed claim
reconciliation via `to_proposal_source_ref` (an EvidenceItem projects
into the exact ProposalSourceRef shape SectionResponseBrief's tier-2
`proposal_source_refs` already carries).
"""
from __future__ import annotations

import copy
import datetime as _dt
import hashlib
import io
import json
import re
from dataclasses import dataclass, field
from typing import Iterable, Optional

import canonical_procurement as cp

CHECK1_CONTRACT_VERSION = "check-1.1.0"

# ═══════════════════════════════════════════════════════════════════════
# 1. Closed vocabularies
# ═══════════════════════════════════════════════════════════════════════

ROLE_TECHNICAL_PROPOSAL = "TECHNICAL_PROPOSAL"
ROLE_PRICING_FORM = "PRICING_FORM"
ROLE_SUBMISSION_FORM = "SUBMISSION_FORM"
ROLE_MULTI_PARTY_FORM = "MULTI_PARTY_FORM"
ROLE_SOCIAL_PROCUREMENT_RESPONSE = "SOCIAL_PROCUREMENT_RESPONSE"
ROLE_CERTIFICATE = "CERTIFICATE"
ROLE_EVIDENCE_ATTACHMENT = "EVIDENCE_ATTACHMENT"
ROLE_RESUME = "RESUME"
ROLE_ORGANIZATION_CHART = "ORGANIZATION_CHART"
ROLE_SUPPORTING_DOCUMENT = "SUPPORTING_DOCUMENT"
ROLE_UNKNOWN = "UNKNOWN"

DOCUMENT_ROLES = (
    ROLE_TECHNICAL_PROPOSAL, ROLE_PRICING_FORM, ROLE_SUBMISSION_FORM,
    ROLE_MULTI_PARTY_FORM, ROLE_SOCIAL_PROCUREMENT_RESPONSE, ROLE_CERTIFICATE,
    ROLE_EVIDENCE_ATTACHMENT, ROLE_RESUME, ROLE_ORGANIZATION_CHART,
    ROLE_SUPPORTING_DOCUMENT, ROLE_UNKNOWN,
)

#: Expected-evidence LOCATION for a response completed directly in a
#: buyer portal (Ariba / MERX / Bonfire / ...) -- never a document role:
#: such a response may legitimately not exist in the uploaded files.
LOCATION_PORTAL_NATIVE = "PORTAL_NATIVE"

ROLE_BASIS_USER_ASSIGNED = "USER_ASSIGNED"
ROLE_BASIS_FILENAME = "FILENAME"
ROLE_BASIS_FILENAME_AND_CONTENT = "FILENAME_AND_CONTENT"
ROLE_BASIS_CONTENT = "CONTENT_SIGNALS"
ROLE_BASIS_NONE = "NO_SIGNAL"

CONFIDENCE_HIGH, CONFIDENCE_MEDIUM, CONFIDENCE_LOW = "HIGH", "MEDIUM", "LOW"

EVIDENCE_KIND_SECTION_TEXT = "SECTION_TEXT"   # prose block within a section/page
EVIDENCE_KIND_TABLE_ROW = "TABLE_ROW"         # one row of a PDF/DOCX table
EVIDENCE_KIND_FORM_FIELD = "FORM_FIELD"       # label -> value pair
EVIDENCE_KIND_CHECKBOX = "CHECKBOX"           # checked / unchecked declaration
EVIDENCE_KIND_SHEET_ROW = "SHEET_ROW"         # one worksheet row (retrieval)
EVIDENCE_KIND_SHEET_CELL = "SHEET_CELL"       # one value/input cell (structured)
EVIDENCE_KINDS = (
    EVIDENCE_KIND_SECTION_TEXT, EVIDENCE_KIND_TABLE_ROW, EVIDENCE_KIND_FORM_FIELD,
    EVIDENCE_KIND_CHECKBOX, EVIDENCE_KIND_SHEET_ROW, EVIDENCE_KIND_SHEET_CELL,
)

# Package-level ARTIFACT presence status for one requirement -- NOT an
# adjudication of whether the requirement is addressed.
ARTIFACT_PRESENT = "ARTIFACT_PRESENT"
MISSING_FROM_PACKAGE = "MISSING_FROM_PACKAGE"
POSSIBLY_PORTAL_NATIVE = "POSSIBLY_PORTAL_NATIVE"
NOT_VERIFIABLE_FROM_FILES = "NOT_VERIFIABLE_FROM_FILES"
FLEXIBLE_LOCATION = "FLEXIBLE_LOCATION"
ARTIFACT_STATUSES = (
    ARTIFACT_PRESENT, MISSING_FROM_PACKAGE, POSSIBLY_PORTAL_NATIVE,
    NOT_VERIFIABLE_FROM_FILES, FLEXIBLE_LOCATION,
)

LOCATION_BASIS_EXPLICIT = "EXPLICIT_FIELD"        # requirement carries expected roles
LOCATION_BASIS_STATED = "STATED_IN_REQUIREMENT"   # text names the artifact/portal
LOCATION_BASIS_INFERRED = "SEMANTIC_INFERENCE"    # derived from semantic type
LOCATION_BASIS_NOT_STATED = "NOT_STATED"          # flexible / unknown

# Every substantive role a flexible requirement may be evidenced in.
_SUBSTANTIVE_ROLES = (
    ROLE_TECHNICAL_PROPOSAL, ROLE_PRICING_FORM, ROLE_SUBMISSION_FORM,
    ROLE_MULTI_PARTY_FORM, ROLE_SOCIAL_PROCUREMENT_RESPONSE, ROLE_CERTIFICATE,
    ROLE_EVIDENCE_ATTACHMENT, ROLE_RESUME, ROLE_ORGANIZATION_CHART,
    ROLE_SUPPORTING_DOCUMENT,
)

MAX_EVIDENCE_CONTENT_CHARS = 2000
DEFAULT_TOP_K = 5


class CrossBidEvidenceError(PermissionError):
    """An evidence id was presented under a bid it does not belong to."""


# ═══════════════════════════════════════════════════════════════════════
# 2. Deterministic document-role classification (section 3)
# ═══════════════════════════════════════════════════════════════════════
#
# Buyer-agnostic procurement vocabulary only -- nothing here names a
# particular buyer, solicitation or appendix letter. "Appendix D" is a
# Price Form in one RFP and a Submission Form in another, so appendix
# letters are NEVER a role signal; the words describing the artifact are.

_TECHNICAL_NAME_RE = re.compile(
    r"\btechnical\s+(?:(?:&|and)\s+(?:commercial|financial)\s+)?(?:proposal|response|submission|offer)\b"
    r"|\bproposal\s+response\b|\brated\s+(?:criteria\s+)?response\b|\bresponse\s+to\s+(?:the\s+)?(?:rfp|rfq|rfsq|tender)\b"
    r"|\b(?:methodology|work\s*plan)\b", re.IGNORECASE)

_ROLE_NAME_RULES: tuple = (
    # (role, filename regex) in PRIORITY order -- the first match is the
    # primary role unless a technical-proposal name also matches.
    (ROLE_MULTI_PARTY_FORM, re.compile(
        r"\bmulti[\s_-]*party\b|\bjoint[\s_-]+(?:venture|bid|submission|proposal|response)\b"
        r"|\bconsortium\b|\bteaming\s+(?:agreement|form)\b|\bsub[\s-]?consultant\s+(?:form|confirmation)\b",
        re.IGNORECASE)),
    (ROLE_SUBMISSION_FORM, re.compile(
        r"\bsubmission\s+form\b|\bproposal\s+submission\s+form\b|\bform\s+of\s+(?:tender|offer|proposal)\b"
        r"|\bbid\s+form\b|\btender\s+form\b|\bdeclarations?\b|\battestations?\b|\bcompliance\s+(?:form|declaration)\b"
        r"|\backnowledg(?:e?ment)\s+form\b", re.IGNORECASE)),
    (ROLE_PRICING_FORM, re.compile(
        r"\bpric(?:e|ing)\s*(?:form|schedule|sheet|proposal|workbook|submission|template)?\b"
        r"|\bcost\s+(?:proposal|form|schedule|breakdown)\b|\brate\s+(?:card|schedule|sheet)\b"
        r"|\bfee\s+(?:schedule|proposal)\b|\b(?:financial|commercial)\s+(?:proposal|offer|response)\b",
        re.IGNORECASE)),
    (ROLE_SOCIAL_PROCUREMENT_RESPONSE, re.compile(
        r"\bsocial\s+(?:procurement|value|benefits?)\b|\bindigenous\s+(?:participation|procurement|benefits?)\b"
        r"|\bcommunity\s+benefits?\b|\bsustainab(?:le|ility)\s+procurement\b", re.IGNORECASE)),
    (ROLE_RESUME, re.compile(
        r"(?:^|[\s_\-(])(?:cvs?|curricul(?:um|a)[\s_]+vitae|r[ée]sum[ée]s?|bios?|biograph(?:y|ies))(?:$|[\s_\-).])",
        re.IGNORECASE)),
    (ROLE_ORGANIZATION_CHART, re.compile(r"\borg(?:ani[sz]ation(?:al)?)?[\s_-]*chart\b", re.IGNORECASE)),
    (ROLE_CERTIFICATE, re.compile(
        r"\bcertificates?\b|\binsurance\b|\bwcb\b|\bworkers[’'`]?\s*comp(?:ensation)?\b|\bgood\s+standing\b"
        r"|\b(?:business\s+)?registration\b|\blicen[cs]es?\b|\baccreditations?\b|\bclearance\s+letter\b",
        re.IGNORECASE)),
    (ROLE_EVIDENCE_ATTACHMENT, re.compile(
        r"\bcase\s+stud(?:y|ies)\b|\breference\s+letters?\b|\bwork\s+samples?\b|\bsamples?\b|\bportfolio\b",
        re.IGNORECASE)),
    (ROLE_SUPPORTING_DOCUMENT, re.compile(
        r"\b(?:attachment|exhibit|supporting|brochure|annex)\b", re.IGNORECASE)),
)

# Content signals: (role, regex). Counted, not merely detected, so a
# document is classified by what it predominantly IS.
_ROLE_CONTENT_RULES: tuple = (
    (ROLE_MULTI_PARTY_FORM, re.compile(
        r"\bmulti[\s-]*party\b|\bjoint\s+venture\b|\bconsortium\b|\blead\s+(?:proponent|member|firm)\b"
        r"|\bmember\s+(?:firm|organi[sz]ation)s?\b|\bjointly\s+and\s+severally\b|\bparticipating\s+part(?:y|ies)\b",
        re.IGNORECASE)),
    (ROLE_SUBMISSION_FORM, re.compile(
        r"\bconflicts?\s+of\s+interest\b|\bauthori[sz]ed\s+(?:signatory|representative|signing\s+officer)\b"
        r"|\blegal\s+name\b|\backnowledg\w*\s+(?:receipt\s+of\s+)?addend\w*|\bhereby\s+(?:declare|certif|confirm|acknowledg)\w*"
        r"|\bdeclar(?:e|ation)s?\b|\battest\w*\b|\bbinding\s+offer\b|\bsignature\b", re.IGNORECASE)),
    (ROLE_PRICING_FORM, re.compile(
        r"\blump\s+sum\b|\bunit\s+(?:price|cost|rate)\b|\btotal\s+(?:weighted\s+)?(?:cost|price)\b|\bhourly\s+rate\b"
        r"|\bper\s+diem\b|\b(?:gst|hst|taxes)\s+(?:excluded|included|extra)\b|\bweighted\s+cost\b|\bprice\b|\bpricing\b",
        re.IGNORECASE)),
    (ROLE_SOCIAL_PROCUREMENT_RESPONSE, re.compile(
        r"\bsocial\s+(?:procurement|value)\b|\bindigenous\b|\bequity[\s-]deserving\b|\bcommunity\s+benefit",
        re.IGNORECASE)),
    (ROLE_RESUME, re.compile(
        r"\bprofessional\s+experience\b|\bemployment\s+history\b|\beducation\b|\bcertifications?\b|\bcurriculum\s+vitae\b",
        re.IGNORECASE)),
    (ROLE_CERTIFICATE, re.compile(
        r"\bcertificate\s+of\s+(?:insurance|registration|status|good\s+standing)\b|\bpolicy\s+number\b|\binsurer\b"
        r"|\bthis\s+is\s+to\s+certify\b|\bclearance\b", re.IGNORECASE)),
    (ROLE_TECHNICAL_PROPOSAL, re.compile(
        r"\bmethodolog\w*\b|\bapproach\b|\bwork\s*plan\b|\bdeliverables?\b|\bimplementation\b|\bproject\s+team\b"
        r"|\bexperience\b|\bgovernance\b|\bquality\s+assurance\b", re.IGNORECASE)),
)

# A section heading inside an artifact that shows it EMBEDS another role
# (e.g. a "Technical & Commercial Proposal" with a PRICING part, or a
# technical proposal with a CONSORTIUM part). Secondary roles only --
# never the primary classification.
_EMBEDDED_ROLE_HEADING_RULES: tuple = (
    (ROLE_PRICING_FORM, re.compile(r"\b(?:pricing|price|cost|fees?|commercial\s+proposal|financial\s+proposal)\b", re.IGNORECASE)),
    (ROLE_MULTI_PARTY_FORM, re.compile(r"\b(?:consortium|joint\s+venture|multi[\s-]*party|partnership\s+structure)\b", re.IGNORECASE)),
    (ROLE_RESUME, re.compile(r"\b(?:resumes?|r[ée]sum[ée]s|cvs|curricul(?:um|a)\s+vitae|team\s+bios?)\b", re.IGNORECASE)),
    (ROLE_ORGANIZATION_CHART, re.compile(r"\borgani[sz]ation(?:al)?\s+chart\b", re.IGNORECASE)),
    (ROLE_SOCIAL_PROCUREMENT_RESPONSE, re.compile(r"\bsocial\s+(?:procurement|value)\b", re.IGNORECASE)),
)

_SPREADSHEET_TYPES = frozenset({"xlsx", "xls", "csv"})


@dataclass(frozen=True)
class RoleClassification:
    role: str
    confidence: str
    basis: str
    secondary_roles: tuple = ()
    signals: tuple = ()

    def to_dict(self) -> dict:
        return {"role": self.role, "confidence": self.confidence, "basis": self.basis,
                "secondary_roles": list(self.secondary_roles), "signals": list(self.signals)}


def _filename_roles(filename: str) -> list[str]:
    stem = re.sub(r"[_]+", " ", filename or "")
    stem = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", stem)
    return [role for role, rx in _ROLE_NAME_RULES if rx.search(stem)]


def _content_role_scores(text: str) -> dict[str, int]:
    body = text or ""
    return {role: len(rx.findall(body)) for role, rx in _ROLE_CONTENT_RULES}


def classify_document_role(
    filename: str,
    text: str = "",
    *,
    file_type: str | None = None,
    section_titles: Iterable[str] = (),
    override: str | None = None,
) -> RoleClassification:
    """Deterministically classify ONE submitted artifact's logical role.

    Precedence:
      1. An explicit human `override` (must be a member of DOCUMENT_ROLES)
         -- recorded as USER_ASSIGNED, never silently second-guessed.
      2. Filename vocabulary. A technical-proposal name wins primary over
         any other name match ("Technical & Commercial Proposal" is a
         TECHNICAL_PROPOSAL that ALSO carries pricing -> secondary
         PRICING_FORM); otherwise the first rule in priority order wins.
      3. Content signals (counted), with one structural rule: a
         spreadsheet whose content is predominantly pricing vocabulary is
         a PRICING_FORM.
      4. UNKNOWN -- never guessed into the nearest bucket.

    Secondary roles come from other filename matches and from section
    headings that show an EMBEDDED part of another role (a "PART 5 --
    PRICING" section inside a technical proposal)."""
    if override:
        if override not in DOCUMENT_ROLES:
            raise ValueError(f"unknown document role override: {override!r}")
        return RoleClassification(override, CONFIDENCE_HIGH, ROLE_BASIS_USER_ASSIGNED,
                                  signals=("user_override",))

    ftype = (file_type or "").lower()
    name_roles = _filename_roles(filename)
    stem = re.sub(r"[_]+", " ", filename or "")
    is_technical_name = bool(_TECHNICAL_NAME_RE.search(stem))
    scores = _content_role_scores(text)

    embedded = []
    for title in section_titles or ():
        for role, rx in _EMBEDDED_ROLE_HEADING_RULES:
            if rx.search(title or "") and role not in embedded:
                embedded.append(role)

    signals: list[str] = []
    primary: str | None = None
    basis = ROLE_BASIS_NONE

    if is_technical_name and ftype not in _SPREADSHEET_TYPES:
        primary = ROLE_TECHNICAL_PROPOSAL
        signals.append("filename:technical_proposal")
    elif name_roles:
        primary = name_roles[0]
        signals.append(f"filename:{primary.lower()}")

    if primary is not None:
        basis = ROLE_BASIS_FILENAME
        if scores.get(primary, 0) > 0:
            basis = ROLE_BASIS_FILENAME_AND_CONTENT
            signals.append(f"content:{primary.lower()}={scores[primary]}")
        confidence = CONFIDENCE_HIGH if basis == ROLE_BASIS_FILENAME_AND_CONTENT else CONFIDENCE_MEDIUM
    else:
        ranked = sorted(((n, r) for r, n in scores.items() if n > 0), key=lambda x: (-x[0], DOCUMENT_ROLES.index(x[1])))
        if ftype in _SPREADSHEET_TYPES and scores.get(ROLE_PRICING_FORM, 0) >= 2:
            primary = ROLE_PRICING_FORM
            signals.append(f"content:pricing_spreadsheet={scores[ROLE_PRICING_FORM]}")
        elif ranked and ranked[0][0] >= 3 and (len(ranked) == 1 or ranked[0][0] >= 2 * ranked[1][0]):
            primary = ranked[0][1]
            signals.append(f"content:{primary.lower()}={ranked[0][0]}")
        if primary is not None:
            basis = ROLE_BASIS_CONTENT
            confidence = CONFIDENCE_MEDIUM if scores.get(primary, 0) >= 5 else CONFIDENCE_LOW
        else:
            primary, confidence = ROLE_UNKNOWN, CONFIDENCE_LOW

    secondary = []
    for r in name_roles + embedded:
        if r != primary and r not in secondary and r not in (ROLE_SUPPORTING_DOCUMENT, ROLE_UNKNOWN):
            secondary.append(r)
    for r in embedded:
        signals.append(f"embedded_section:{r.lower()}")

    return RoleClassification(primary, confidence, basis, tuple(secondary), tuple(dict.fromkeys(signals)))


# ═══════════════════════════════════════════════════════════════════════
# 3. Structured parsing (sections 4 / 8)
# ═══════════════════════════════════════════════════════════════════════

_TOC_LINE_RE = re.compile(r"\.{4,}\s*\d+\s*$|\s{2,}\d+\s*$")
_NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,4})[.)]?\s+([A-Z][^\n]{1,118})$")
_PART_HEADING_RE = re.compile(
    r"^(PART|SECTION|APPENDIX|SCHEDULE|ANNEX|ARTICLE)\s+([A-Z0-9IVXL]{1,4})\b\s*[—–:\-.]?\s*(.{0,110})$",
    re.IGNORECASE)
_CAPS_HEADING_RE = re.compile(r"^[A-Z][A-Z0-9 &/,'\-—–()]{4,79}$")


@dataclass(frozen=True)
class ProposalSection:
    """One heading in a submitted artifact's own structure. `number` is
    the author's own numbering ("3.2", "PART 4") -- never renumbered."""

    section_id: str
    number: Optional[str]
    title: str
    level: int
    parent_id: Optional[str]
    page_start: Optional[int]
    page_end: Optional[int] = None

    def to_dict(self) -> dict:
        return {"section_id": self.section_id, "number": self.number, "title": self.title,
                "level": self.level, "parent_id": self.parent_id,
                "page_start": self.page_start, "page_end": self.page_end}

    @property
    def label(self) -> str:
        return f"{self.number} {self.title}".strip() if self.number and self.number not in self.title else self.title


def detect_heading(line: str) -> tuple[Optional[str], str, int] | None:
    """(number, title, level) when `line` is a heading, else None.
    Table-of-contents lines (dot leaders / trailing page numbers) are
    never headings -- a TOC entry is not the section itself."""
    s = (line or "").strip()
    if not s or len(s) > 130 or _TOC_LINE_RE.search(s):
        return None
    m = _PART_HEADING_RE.match(s)
    if m and (m.group(3) or "").strip()[:1].upper() == (m.group(3) or "").strip()[:1] and not s.endswith("."):
        kind = m.group(1).upper()
        number = f"{kind} {m.group(2).upper()}"
        title = (m.group(3) or "").strip(" —–-:.\t") or number
        return number, title, 1
    m = _NUMBERED_HEADING_RE.match(s)
    if m and not s.endswith((".", ",", ";")) and len(s.split()) <= 14:
        number = m.group(1)
        return number, m.group(2).strip(), number.count(".") + 1
    if _CAPS_HEADING_RE.match(s) and len(s.split()) >= 2 and not re.search(r"\d{3,}", s):
        return None, s, 1
    return None


class _SectionTracker:
    """Builds a heading hierarchy as headings are encountered in reading
    order; parent = nearest previous section of a lower level."""

    def __init__(self, doc_id: str):
        self.doc_id = doc_id
        self.sections: list[dict] = []
        self.current: Optional[dict] = None

    def open(self, number, title, level, page) -> dict:
        parent = None
        for prev in reversed(self.sections):
            if prev["level"] < level:
                parent = prev["section_id"]
                break
        sid = f"S{len(self.sections) + 1}"
        sec = {"section_id": sid, "number": number, "title": title, "level": level,
               "parent_id": parent, "page_start": page, "page_end": None}
        self.sections.append(sec)
        self.current = sec
        return sec

    def touch(self, page):
        if self.current is not None and page is not None:
            self.current["_last_page"] = page

    def finish(self) -> tuple:
        out = []
        for sec in self.sections:
            last = sec.get("_last_page")
            start = sec["page_start"]
            end = max(last, start) if (last is not None and start is not None) else (last if last is not None else start)
            out.append(ProposalSection(sec["section_id"], sec["number"], sec["title"], sec["level"],
                                       sec["parent_id"], sec["page_start"], end))
        return tuple(out)

    def path(self) -> list[str]:
        if self.current is None:
            return []
        by_id = {s["section_id"]: s for s in self.sections}
        chain, node = [], self.current
        while node is not None:
            label = f"{node['number']} {node['title']}".strip() if node["number"] and node["number"] not in node["title"] else node["title"]
            chain.append(label)
            node = by_id.get(node["parent_id"]) if node["parent_id"] else None
        return list(reversed(chain))


def _section_location(tracker: _SectionTracker) -> dict:
    cur = tracker.current
    if cur is None:
        return {}
    return {"section_id": cur["section_id"], "section_number": cur["number"],
            "section_title": cur["title"], "section_path": tracker.path()}


# `\b` after the keyword matters (CHECK-1.1, real Calgary proposal): a real
# answer beginning "Enterprise administration ..." must not read as an
# "enter ..." placeholder.
_PLACEHOLDER_RE = re.compile(r"^[\s_.\-—–]*$|^\[?\s*(?:insert|enter|type|tbd|n/?a\s+if|click\s+here)\b[^\]]*\]?$", re.IGNORECASE)
_LABEL_VALUE_RE = re.compile(r"^([A-Z][^:\n]{1,80}?)\s*:\s*(.*)$")
# Checkbox glyphs only (ballot boxes / bracketed x). Tick marks and
# filled squares are routinely used as list bullets in narrative
# proposals, so they are deliberately NOT treated as checkboxes.
_CHECKED_RE = re.compile(r"[☒☑]|\[\s*[xX]\s*\]")
_UNCHECKED_RE = re.compile(r"☐|\[\s\]")


def _is_completed(value) -> bool:
    if value is None:
        return False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return True
    return not _PLACEHOLDER_RE.match(str(value).strip())


def _jsonable(value):
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _raw_item(kind: str, locator: str, content: str, location: dict, structured: dict | None = None) -> dict:
    return {"kind": kind, "locator": locator, "content": (content or "").strip()[:MAX_EVIDENCE_CONTENT_CHARS],
            "location": location, "structured_value": structured}


def _line_items(lines: list[str], tracker: _SectionTracker, page: Optional[int], prefix: str,
                items: list[dict], block_idx: int, *, max_value_chars: int = 120) -> None:
    """Form-ish lines inside a prose block: 'Label: value' pairs and
    checkbox declarations become their own structured evidence."""
    for li, line in enumerate(lines):
        s = line.strip()
        if not s:
            continue
        loc = {"page": page, **_section_location(tracker)}
        if _CHECKED_RE.search(s) or _UNCHECKED_RE.search(s):
            checked = bool(_CHECKED_RE.search(s))
            label = _UNCHECKED_RE.sub("", _CHECKED_RE.sub("", s)).strip(" :-")
            items.append(_raw_item(EVIDENCE_KIND_CHECKBOX, f"{prefix}:{block_idx}:l{li}", s, loc,
                                   {"label": label, "checked": checked, "completed": checked}))
            continue
        m = _LABEL_VALUE_RE.match(s)
        # A form field is a SHORT label with a short (or blank) value --
        # "About Phoenix: Founded in 2010 ..." is prose, not a form field.
        if m and len(m.group(1).split()) <= 6 and len(m.group(2).strip()) <= max_value_chars:
            value = m.group(2).strip()
            items.append(_raw_item(EVIDENCE_KIND_FORM_FIELD, f"{prefix}:{block_idx}:l{li}", s, loc,
                                   {"label": m.group(1).strip(), "value": value or None,
                                    "completed": _is_completed(value)}))


def _pdf_form_table_fields(rows: list) -> Optional[list[dict]]:
    """CHECK-1.1 (real Calgary Appendix E): a PDF LABEL -> VALUE form table
    (a label column, a value column, labels wrapped over several visual
    lines) is NOT a header-mapped data table. PyMuPDF reports each wrapped
    label line as its own row and marks the value cell of a continuation
    line as None (merged cell) -- that None/'' distinction is what groups
    lines into fields, so no layout guessing is needed. Returns None when
    the table is not a label/value form (header-like first row, no
    'Label:' cell, or no value column), leaving the existing header-mapped
    TABLE_ROW path in charge."""
    if not rows or max(len(r) for r in rows) < 2:
        return None
    norm = [[(c.replace("\n", " ").strip() if isinstance(c, str) else None) for c in r] for r in rows]
    first_non_empty = [c for c in norm[0] if c]
    if len(first_non_empty) >= 3:
        return None  # a header row of a data table
    if not any(any((c or "").endswith((":", "?")) for c in r[:-1]) for r in norm):
        return None
    value_col = None
    for r in norm:
        idx = [i for i, c in enumerate(r) if c]
        if len(idx) >= 2:
            value_col = max(value_col or 0, idx[-1])
    if value_col is None:
        # A wholly BLANK form (labels only): the value column is the last
        # one -- every field must still be recorded, as not completed.
        value_col = max(len(r) for r in norm) - 1
    fields: list[dict] = []
    for ri, r in enumerate(norm):
        labels = [c for c in r[:value_col] if c]
        cell = r[value_col] if value_col < len(r) else None
        if cell is not None or not fields:
            parts: list[str] = []
            for c in labels:  # drop a wrapped fragment repeated in a sibling cell
                if not any(c != o and c in o for o in labels) and c not in parts:
                    parts.append(c)
            fields.append({"label_parts": parts, "value": cell or "", "row_start": ri, "row_end": ri})
        else:
            cur = fields[-1]
            for c in labels:
                if not any(c in p for p in cur["label_parts"]):
                    cur["label_parts"].append(c)
            if labels:
                cur["row_end"] = ri
    out = []
    for f in fields:
        label = " ".join(f["label_parts"]).strip()
        if not label and not f["value"]:
            continue
        out.append({"label": label or "(unlabelled)", "value": f["value"],
                    "row_start": f["row_start"], "row_end": f["row_end"]})
    return out or None


def parse_pdf_structure(file_bytes: bytes) -> dict:
    """Per-page, per-block PDF structure: heading hierarchy, prose blocks
    with page + section, tables (PyMuPDF find_tables) as header-mapped
    rows, and label/value + checkbox lines. Never flattens the document
    into one string -- every item keeps its page."""
    import fitz  # PyMuPDF -- already a hard dependency of extractor.py

    items: list[dict] = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    tracker = _SectionTracker("pdf")
    page_count = len(doc)
    tables_count = 0
    for pno, page in enumerate(doc, 1):
        table_rects = []
        pending_tables: list[tuple] = []   # (y0, ti, rows)
        try:
            found = page.find_tables()
            for ti, table in enumerate(found.tables):
                rows = table.extract() or []
                if not rows:
                    continue
                tables_count += 1
                table_rects.append(fitz.Rect(table.bbox))
                pending_tables.append((fitz.Rect(table.bbox).y0, ti, rows))
        except Exception:
            table_rects, pending_tables = [], []

        def _emit_table(ti: int, rows: list) -> None:
            # Emitted in READING ORDER (CHECK-1.1): a table is attributed to
            # the section heading that precedes it on the page, not to
            # whatever section was open when the page started.
            form_fields = _pdf_form_table_fields(rows)
            if form_fields is not None:
                for fi, fld in enumerate(form_fields):
                    loc = {"page": pno, "table_index": ti, "row_index": fld["row_start"],
                           "row_end": fld["row_end"], **_section_location(tracker)}
                    content = f"{fld['label']}: {fld['value']}" if fld["value"] else f"{fld['label']}: (blank)"
                    items.append(_raw_item(EVIDENCE_KIND_FORM_FIELD, f"p{pno}t{ti}f{fi}", content, loc,
                                           {"label": fld["label"], "value": fld["value"] or None,
                                            "completed": _is_completed(fld["value"] or ""),
                                            "source": "pdf_form_table"}))
                return
            header = [(c or "").replace("\n", " ").strip() for c in rows[0]]
            use_header = sum(1 for h in header if h) >= max(2, len(header) // 2)
            for ri, row in enumerate(rows[1:] if use_header else rows, 1 if use_header else 0):
                cells = [(c or "").replace("\n", " ").strip() for c in row]
                if not any(cells):
                    continue
                mapping = ({(header[i] or f"col{i + 1}"): cells[i] for i in range(min(len(header), len(cells)))}
                           if use_header else {f"col{i + 1}": c for i, c in enumerate(cells)})
                loc = {"page": pno, "table_index": ti, "row_index": ri, **_section_location(tracker)}
                items.append(_raw_item(EVIDENCE_KIND_TABLE_ROW, f"p{pno}t{ti}r{ri}",
                                       " | ".join(c for c in cells if c), loc,
                                       {"columns": mapping, "completed": any(_is_completed(c) for c in cells)}))

        pending_tables.sort(key=lambda t: (t[0], t[1]))

        blocks = page.get_text("blocks") or []
        # Top-to-bottom order (PyMuPDF can emit a page footer before the
        # body); `bi` stays the original block index so locators are stable.
        ordered = sorted(enumerate(blocks), key=lambda ib: (round(ib[1][1]), ib[1][0]) if len(ib[1]) >= 4 else (0, 0))
        for bi, block in ordered:
            if len(block) < 7 or block[6] != 0:
                continue  # image block
            rect = fitz.Rect(block[:4])
            while pending_tables and pending_tables[0][0] <= rect.y0:
                _, t_i, t_rows = pending_tables.pop(0)
                _emit_table(t_i, t_rows)
            if any(rect.intersects(tr) and (rect & tr).get_area() > 0.6 * max(rect.get_area(), 1) for tr in table_rects):
                continue  # already captured as table rows
            lines = [ln.strip() for ln in (block[4] or "").splitlines() if ln.strip()]
            if not lines:
                continue
            # A heading can sit anywhere inside a PyMuPDF block: split the
            # block at every heading so each prose run is attributed to the
            # section it actually belongs to.
            runs: list[list[str]] = [[]]
            for line in lines:
                h = detect_heading(line)
                if h is not None:
                    runs.append(("__HEADING__", h))  # type: ignore[arg-type]
                    runs.append([])
                else:
                    runs[-1].append(line)
            part = 0
            for run in runs:
                if isinstance(run, tuple):
                    h = run[1]
                    tracker.open(h[0], h[1], h[2], pno)
                    continue
                tracker.touch(pno)
                text = " ".join(run).strip()
                if len(text) < 3 or not re.search(r"[A-Za-z0-9]{2}", text):
                    continue
                suffix = f"b{bi}" if part == 0 else f"b{bi}.{part}"
                part += 1
                loc = {"page": pno, "block_index": bi, **_section_location(tracker)}
                items.append(_raw_item(EVIDENCE_KIND_SECTION_TEXT, f"p{pno}{suffix}", text, loc))
                # Wrapped PDF prose makes long "Label: text" lines look like
                # form fields; only short-valued lines count.
                _line_items(run, tracker, pno, f"p{pno}{suffix}", items, bi, max_value_chars=60)
        for _, t_i, t_rows in pending_tables:
            _emit_table(t_i, t_rows)
    sections = tracker.finish()
    return {"parser": "pymupdf", "page_count": page_count, "sections": sections,
            "items": items, "tables_count": tables_count}


def _iter_docx_blocks(document):
    """Paragraphs and tables in true body order (python-docx exposes them
    separately; body order matters for section context)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def _walk(parent):
        for child in parent.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                yield Paragraph(child, document)
            elif tag == "tbl":
                yield Table(child, document)
            elif tag in ("sdt", "sdtContent", "customXml", "smartTag"):
                # CHECK-1.1: real buyer forms (e.g. a Submission Form
                # template) wrap labels/paragraphs/tables in content
                # controls (w:sdt). python-docx's body iteration skips
                # them, which silently dropped every label of the real
                # Appendix E DOCX -- descend, in body order.
                yield from _walk(child)

    yield from _walk(document.element.body)


def parse_docx_structure(file_bytes: bytes) -> dict:
    """DOCX structure: heading styles (and numbered headings) -> hierarchy;
    prose paragraphs; tables as label->value FORM_FIELDs (2-column or
    label-first rows) or header-mapped TABLE_ROWs; checkbox glyph
    declarations. DOCX has no reliable page geometry, so `page` is None --
    honestly absent, never fabricated."""
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(file_bytes))
    tracker = _SectionTracker("docx")
    items: list[dict] = []
    tables_count = 0
    pi = 0
    for block in _iter_docx_blocks(document):
        if isinstance(block, Table):
            ti = tables_count
            tables_count += 1
            rows = []
            for row in block.rows:
                cells, prev = [], None
                for cell in row.cells:
                    t = (cell.text or "").strip()
                    if cell._tc is prev:
                        continue  # merged cell repeated by python-docx
                    prev = cell._tc
                    cells.append(t)
                rows.append(cells)
            if not rows:
                continue
            header = rows[0]
            header_like = (len(header) >= 3 and all(header)
                           and not any(_CHECKED_RE.search(h) or _UNCHECKED_RE.search(h) for h in header))
            for ri, cells in enumerate(rows):
                if not any(cells):
                    continue
                loc = {"page": None, "table_index": ti, "row_index": ri, **_section_location(tracker)}
                joined = " | ".join(c for c in cells if c)
                if any(_CHECKED_RE.search(c) or _UNCHECKED_RE.search(c) for c in cells):
                    checked = any(_CHECKED_RE.search(c) for c in cells)
                    label = " ".join(_UNCHECKED_RE.sub("", _CHECKED_RE.sub("", c)).strip() for c in cells).strip()
                    items.append(_raw_item(EVIDENCE_KIND_CHECKBOX, f"t{ti}r{ri}", joined, loc,
                                           {"label": label, "checked": checked, "completed": checked}))
                elif header_like and ri > 0:
                    mapping = {(header[i] or f"col{i + 1}"): cells[i] for i in range(min(len(header), len(cells)))}
                    items.append(_raw_item(EVIDENCE_KIND_TABLE_ROW, f"t{ti}r{ri}", joined, loc,
                                           {"columns": mapping, "completed": any(_is_completed(c) for c in cells[1:])}))
                elif header_like and ri == 0:
                    continue
                elif len(cells) >= 2 and cells[0] and len(cells[0]) <= 160:
                    value = " | ".join(c for c in cells[1:] if c and c != cells[0]).strip()
                    items.append(_raw_item(EVIDENCE_KIND_FORM_FIELD, f"t{ti}r{ri}", joined, loc,
                                           {"label": cells[0], "value": value or None,
                                            "completed": _is_completed(value)}))
                else:
                    items.append(_raw_item(EVIDENCE_KIND_TABLE_ROW, f"t{ti}r{ri}", joined, loc,
                                           {"columns": {f"col{i + 1}": c for i, c in enumerate(cells)},
                                            "completed": any(_is_completed(c) for c in cells)}))
            continue

        text = (block.text or "").strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        m = re.match(r"(?:Heading|Title)\s*(\d)?", style)
        if m:
            level = int(m.group(1)) if m.group(1) else 1
            det = detect_heading(text)
            number = det[0] if det else None
            title = det[1] if det else text
            tracker.open(number, title, level if not det or det[0] is None else det[2], None)
            continue
        det = detect_heading(text)
        if det is not None and det[0] is not None:
            tracker.open(det[0], det[1], det[2], None)
            continue
        pi += 1
        loc = {"page": None, "paragraph_index": pi, **_section_location(tracker)}
        items.append(_raw_item(EVIDENCE_KIND_SECTION_TEXT, f"para{pi}", text, loc))
        _line_items([text], tracker, None, "para", items, pi)
    return {"parser": "python-docx", "page_count": None, "sections": tracker.finish(),
            "items": items, "tables_count": tables_count}


def _is_yellow(rgb) -> bool:
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    h = rgb[-6:]
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return False
    return r >= 0xE0 and g >= 0xC0 and b <= 0xA0


def parse_xlsx_structure(file_bytes: bytes) -> dict:
    """XLSX structure, never flattened into untraceable text: every
    non-empty row becomes a SHEET_ROW (retrieval), and every VALUE cell
    (numeric / date / text value with a label to its left) and every
    INPUT cell (yellow-filled, or unlocked on a protected sheet) --
    completed OR BLANK -- becomes a SHEET_CELL carrying sheet, cell
    coordinate, row, column, row label, column header, is_formula and
    completed. Formula cells record their cached value; a formula with no
    cached value is `completed: None` (unknown), never assumed filled."""
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb_values = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    try:
        wb_formulas = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
    except Exception:
        wb_formulas = None
    items: list[dict] = []
    sheets = []
    for ws in wb_values.worksheets:
        wsf = wb_formulas[ws.title] if wb_formulas is not None and ws.title in wb_formulas.sheetnames else None
        protected = bool(getattr(ws.protection, "sheet", False))
        grid: dict[tuple, object] = {}
        inputs: set = set()
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                is_input = False
                try:
                    fill = c.fill
                    if fill is not None and fill.fill_type == "solid" and _is_yellow(getattr(fill.fgColor, "rgb", None)):
                        is_input = True
                except Exception:
                    pass
                if protected and getattr(c.protection, "locked", True) is False:
                    is_input = True
                if is_input:
                    inputs.add((c.row, c.column))
                if v is not None and str(v).strip() != "":
                    grid[(c.row, c.column)] = v
        sheets.append({"sheet": ws.title, "state": ws.sheet_state, "input_cells": len(inputs),
                       "non_empty_cells": len(grid)})

        def _is_text(v):
            return isinstance(v, str) and not re.fullmatch(r"[\s$€£%.,\d()\-]+", v)

        def _row_label(r, c):
            for cc in range(c - 1, 0, -1):
                v = grid.get((r, cc))
                if _is_text(v):
                    return str(v).strip()
            return None

        def _col_header(r, c):
            for rr in range(r - 1, 0, -1):
                v = grid.get((rr, c))
                if _is_text(v):
                    return str(v).strip()
            return None

        formulas: dict[tuple, str] = {}
        if wsf is not None:
            for row in wsf.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas[(c.row, c.column)] = c.value
        rows_idx = sorted({r for r, _ in grid} | {r for r, _ in inputs} | {r for r, _ in formulas})
        for r in rows_idx:
            cols = sorted(c for (rr, c) in grid if rr == r)
            if cols:
                cells = [{"cell": f"{get_column_letter(c)}{r}", "value": _jsonable(grid[(r, c)])} for c in cols]
                content = f"{ws.title} row {r}: " + " | ".join(str(x["value"]) for x in cells)
                items.append(_raw_item(EVIDENCE_KIND_SHEET_ROW, f"{ws.title}!R{r}", content,
                                       {"sheet": ws.title, "row": r}, {"cells": cells}))
            for c in sorted({c for (rr, c) in grid if rr == r} | {c for (rr, c) in inputs if rr == r}
                            | {c for (rr, c) in formulas if rr == r}):
                v = grid.get((r, c))
                is_input = (r, c) in inputs
                formula = formulas.get((r, c))
                label = _row_label(r, c)
                header = _col_header(r, c)
                is_value_cell = v is not None and (not _is_text(v) or label is not None)
                if not (is_input or formula or (is_value_cell and not (label is None and header is None))):
                    continue
                if _is_text(v) and not is_input and label is None:
                    continue  # a pure label/heading cell, captured in SHEET_ROW
                coord = f"{get_column_letter(c)}{r}"
                completed = None if (formula and v is None) else _is_completed(v)
                content = f"{ws.title}!{coord}"
                if label:
                    content += f" [{label}]"
                if header and header != label:
                    content += f" ({header})"
                content += f" = {v if v is not None else '(blank)'}"
                items.append(_raw_item(EVIDENCE_KIND_SHEET_CELL, f"{ws.title}!{coord}", content,
                                       {"sheet": ws.title, "cell": coord, "row": r, "column": get_column_letter(c)},
                                       {"value": _jsonable(v), "label": label, "column_header": header,
                                        "is_input_cell": is_input, "is_formula": bool(formula),
                                        "formula": formula, "completed": completed}))
    return {"parser": "openpyxl", "page_count": None, "sections": (), "items": items,
            "sheets": sheets, "tables_count": 0}


def parse_text_rows_structure(text: str, *, sheet_hint: str | None = None) -> dict:
    """Fallback for CSV / XLS / TXT: reuses the extractor's already-marked
    'Row N: a | b' text (so row coordinates stay the physical ones) or
    blank-line paragraphs."""
    items: list[dict] = []
    current_sheet = sheet_hint
    tracker = _SectionTracker("text")
    para = 0
    for raw in (text or "").split("\n"):
        line = raw.strip()
        if not line:
            continue
        sm = re.match(r"\[\[SOURCE:[^|\]]*(?:\|\s*SHEET:\s*([^|\]]+))?", line)
        if sm:
            if sm.group(1):
                current_sheet = sm.group(1).strip()
            continue
        rm = re.match(r"Row\s+(\d+):\s*(.*)$", line)
        if rm:
            r = int(rm.group(1))
            items.append(_raw_item(EVIDENCE_KIND_SHEET_ROW, f"{current_sheet or 'rows'}!R{r}",
                                   f"{current_sheet + ' ' if current_sheet else ''}row {r}: {rm.group(2)}",
                                   {"sheet": current_sheet, "row": r},
                                   {"cells": [{"value": v.strip()} for v in rm.group(2).split("|")]}))
            continue
        det = detect_heading(line)
        if det is not None:
            tracker.open(det[0], det[1], det[2], None)
            continue
        para += 1
        items.append(_raw_item(EVIDENCE_KIND_SECTION_TEXT, f"para{para}", line,
                               {"page": None, "paragraph_index": para, **_section_location(tracker)}))
        _line_items([line], tracker, None, "para", items, para)
    return {"parser": "marked-text", "page_count": None, "sections": tracker.finish(),
            "items": items, "tables_count": 0}


def parse_document_structure(file_bytes: bytes, file_type: str, extracted_text: str = "") -> dict:
    ftype = (file_type or "").lower()
    try:
        if ftype == "pdf":
            return parse_pdf_structure(file_bytes)
        if ftype == "docx":
            return parse_docx_structure(file_bytes)
        if ftype == "xlsx":
            return parse_xlsx_structure(file_bytes)
    except Exception as exc:  # structured parse failed -> honest fallback, never silent loss
        out = parse_text_rows_structure(extracted_text)
        out["parser"] = f"marked-text (structured {ftype} parse failed: {type(exc).__name__})"
        return out
    return parse_text_rows_structure(extracted_text)


# ═══════════════════════════════════════════════════════════════════════
# 4. Evidence registry (section 9)
# ═══════════════════════════════════════════════════════════════════════

def derive_evidence_id(bid_id: int, submission_document_id: str, kind: str, locator: str) -> str:
    """Deterministic, bid-bound evidence id. Folding bid_id in means the
    same bytes uploaded under a different bid yield DIFFERENT ids -- an id
    can never be replayed across a tenancy boundary."""
    raw = f"{bid_id}\x00{submission_document_id}\x00{kind}\x00{locator}"
    return "EV-" + hashlib.sha256(raw.encode("utf-8", "surrogateescape")).hexdigest()[:20]


@dataclass(frozen=True)
class EvidenceItem:
    evidence_id: str
    bid_id: int
    submission_document_id: str
    document_role: str
    kind: str
    location: dict
    content: str
    structured_value: Optional[dict]
    provenance: dict

    def to_dict(self) -> dict:
        return {"evidence_id": self.evidence_id, "bid_id": self.bid_id,
                "submission_document_id": self.submission_document_id,
                "document_role": self.document_role, "kind": self.kind,
                "location": self.location, "content": self.content,
                "structured_value": self.structured_value, "provenance": self.provenance}

    def citation(self) -> str:
        """Human-traceable location, e.g.
        'Technical Proposal.pdf, page 12, section 3.2 Six-Stage ...' or
        'Price Form.xlsx, sheet Pricing, cell E7'."""
        return format_citation(self.provenance.get("filename"), self.location)


def format_citation(filename: str | None, location: dict) -> str:
    parts = [filename or "(unnamed file)"]
    loc = location or {}
    if loc.get("sheet"):
        parts.append(f"sheet {loc['sheet']}")
    if loc.get("cell"):
        parts.append(f"cell {loc['cell']}")
    elif loc.get("row") is not None and loc.get("sheet") is not None:
        parts.append(f"row {loc['row']}")
    if loc.get("page") is not None:
        parts.append(f"page {loc['page']}")
    if loc.get("section_title"):
        num = loc.get("section_number")
        title = loc["section_title"]
        parts.append(f"section {num} {title}".replace("  ", " ") if num and num not in title else f"section {title}")
    if loc.get("table_index") is not None:
        parts.append(f"table {loc['table_index'] + 1} row {loc.get('row_index')}")
    return ", ".join(parts)


class SubmissionEvidenceRegistry:
    """THE single bidder-evidence registry for a package. Future CHECK
    analyzers reference items ONLY by evidence_id; `get`/`filter_valid`
    fail closed on unknown or cross-bid ids."""

    def __init__(self, bid_id: int, items: Iterable[EvidenceItem] = ()):
        self.bid_id = bid_id
        self._items: dict[str, EvidenceItem] = {}
        for it in items:
            if it.bid_id != bid_id:
                raise CrossBidEvidenceError(f"evidence {it.evidence_id} belongs to bid {it.bid_id}, not {bid_id}")
            self._items[it.evidence_id] = it

    def __len__(self):
        return len(self._items)

    def __iter__(self):
        return iter(self._items.values())

    def __contains__(self, evidence_id):
        return evidence_id in self._items

    def get(self, evidence_id: str, *, bid_id: int) -> EvidenceItem:
        if bid_id != self.bid_id:
            raise CrossBidEvidenceError(f"registry for bid {self.bid_id} cannot resolve evidence for bid {bid_id}")
        item = self._items.get(evidence_id)
        if item is None:
            raise KeyError(evidence_id)
        return item

    def filter_valid(self, evidence_ids: Iterable, *, bid_id: int) -> list[str]:
        """Fail-closed filter (same discipline as section_drafting's
        reconcile_material_claims): unknown / malformed / cross-bid ids
        are dropped, never trusted."""
        if bid_id != self.bid_id:
            return []
        return [e for e in (evidence_ids or []) if isinstance(e, str) and e in self._items]

    def for_document(self, submission_document_id: str) -> list[EvidenceItem]:
        return [i for i in self._items.values() if i.submission_document_id == submission_document_id]

    def for_roles(self, roles: Iterable[str], package: "SubmissionPackage") -> list[EvidenceItem]:
        doc_ids = {d.submission_document_id for d in package.documents_with_roles(roles)}
        return [i for i in self._items.values() if i.submission_document_id in doc_ids]


# ═══════════════════════════════════════════════════════════════════════
# 5. Canonical submission package (section 3)
# ═══════════════════════════════════════════════════════════════════════

@dataclass(frozen=True)
class SubmissionDocument:
    submission_document_id: str
    content_hash: str
    filename: str
    package_path: str
    file_type: str
    lifecycle_status: str
    included: bool
    role: RoleClassification
    parse_status: str                       # PARSED / FAILED / UNSUPPORTED / REJECTED / DUPLICATE
    parser: Optional[str]
    page_count: Optional[int]
    sheets: tuple = ()
    sections: tuple = ()
    evidence_count: int = 0
    duplicate_of: Optional[str] = None
    unusable_reason: Optional[str] = None
    # CHECK-1.1 -- logical-artifact identity. Several FILES can be ONE
    # submitted item (a DOCX source + its submitted PDF, a re-saved copy of
    # the same completed workbook, a byte-identical copy inside a nested
    # archive, an earlier draft). Exactly one member per logical artifact is
    # AUTHORITATIVE; only it contributes evidence and package membership.
    logical_artifact_id: Optional[str] = None
    representation_relationship: str = "AUTHORITATIVE"
    representation_of: Optional[str] = None
    representation_basis: Optional[str] = None

    @property
    def document_role(self) -> str:
        return self.role.role

    @property
    def all_roles(self) -> tuple:
        return (self.role.role,) + tuple(self.role.secondary_roles)

    @property
    def readable(self) -> bool:
        return self.parse_status == "PARSED"

    @property
    def is_authoritative(self) -> bool:
        return self.representation_relationship == REP_AUTHORITATIVE

    def to_dict(self) -> dict:
        return {
            "submission_document_id": self.submission_document_id, "content_hash": self.content_hash,
            "filename": self.filename, "package_path": self.package_path, "file_type": self.file_type,
            "lifecycle_status": self.lifecycle_status, "included": self.included,
            "document_role": self.role.role, "role_classification": self.role.to_dict(),
            "parse_status": self.parse_status, "parser": self.parser, "page_count": self.page_count,
            "sheets": list(self.sheets), "sections": [s.to_dict() for s in self.sections],
            "evidence_count": self.evidence_count, "duplicate_of": self.duplicate_of,
            "unusable_reason": self.unusable_reason,
            "logical_artifact_id": self.logical_artifact_id,
            "representation_relationship": self.representation_relationship,
            "representation_of": self.representation_of,
            "representation_basis": self.representation_basis,
        }


@dataclass(frozen=True)
class SubmissionPackage:
    bid_id: int
    organization_id: str
    package_digest: str
    documents: tuple
    registry: SubmissionEvidenceRegistry
    manifest: tuple
    contract_version: str = CHECK1_CONTRACT_VERSION

    def document(self, submission_document_id: str) -> SubmissionDocument:
        for d in self.documents:
            if d.submission_document_id == submission_document_id:
                return d
        raise KeyError(submission_document_id)

    def member_documents(self) -> list[SubmissionDocument]:
        """Package members that stand for a submitted artifact: included,
        not a byte-identical duplicate, not a rejected ZIP entry, and the
        AUTHORITATIVE representation of its logical artifact (an alternate
        DOCX/PDF/re-saved representation is never double-counted)."""
        return [d for d in self.documents
                if d.included and d.lifecycle_status not in ("duplicate", "rejected") and d.is_authoritative]

    def documents_with_roles(self, roles: Iterable[str], *, include_secondary: bool = True) -> list[SubmissionDocument]:
        wanted = set(roles)
        out = []
        for d in self.member_documents():
            have = set(d.all_roles) if include_secondary else {d.document_role}
            if have & wanted:
                out.append(d)
        return out

    def roles_present(self) -> dict[str, list[str]]:
        out: dict[str, list[str]] = {}
        for d in self.member_documents():
            for r in d.all_roles:
                out.setdefault(r, []).append(d.submission_document_id)
        return out

    def logical_artifacts(self) -> list[dict]:
        """One entry per LOGICAL submitted artifact: its authoritative
        document plus every linked representation (never double-counted)."""
        groups: dict[str, dict] = {}
        for d in self.documents:
            key = d.logical_artifact_id or d.submission_document_id
            g = groups.setdefault(key, {"logical_artifact_id": key, "authoritative": None, "representations": []})
            if d.is_authoritative and g["authoritative"] is None:
                g["authoritative"] = d
            else:
                g["representations"].append(d)
        return list(groups.values())

    def to_dict(self) -> dict:
        return {"bid_id": self.bid_id, "package_digest": self.package_digest,
                "contract_version": self.contract_version,
                "documents": [d.to_dict() for d in self.documents],
                "evidence_count": len(self.registry)}


# ── Logical-artifact / representation linking (CHECK-1.1) ──────────────

REP_AUTHORITATIVE = "AUTHORITATIVE"
REP_ALTERNATE = "ALTERNATE_REPRESENTATION"            # same item, other format / re-saved identical content
REP_DRAFT = "SUPERSEDED_OR_DRAFT_VARIANT"             # same item, differing (earlier/other) completed content
REP_TEMPLATE = "TEMPLATE_VARIANT"                     # same item, blank / uncompleted copy
REP_BYTE_DUPLICATE = "BYTE_IDENTICAL_DUPLICATE"       # extractor-detected identical bytes
REPRESENTATION_RELATIONSHIPS = (REP_AUTHORITATIVE, REP_ALTERNATE, REP_DRAFT, REP_TEMPLATE, REP_BYTE_DUPLICATE)

# Final submitted formats rank ahead of editable sources. A spreadsheet is
# its own final format (a pricing workbook is submitted as the workbook).
_FORMAT_RANK = {"pdf": 0, "xlsx": 0, "xls": 0, "csv": 0, "docx": 1, "txt": 2, "md": 2}
_STEM_SPLIT_RE = re.compile(r"[^a-z0-9]+")


def _stem_tokens(filename: str) -> tuple:
    stem = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", (filename or "").lower())
    return tuple(t for t in _STEM_SPLIT_RE.split(stem) if t)


def _content_tokens(items: list[dict]) -> set:
    toks: set = set()
    for it in items:
        toks |= set(re.findall(r"[a-z0-9]{3,}", (it.get("content") or "").lower()))
    return toks


def _sheet_signature(items: list[dict]) -> Optional[str]:
    """Digest of every non-empty worksheet cell VALUE (sheet + coordinate +
    value). Two workbooks with the same signature carry identical content
    even when their bytes differ (re-saved / renamed copies)."""
    rows = sorted((it["locator"], json.dumps(it.get("structured_value"), sort_keys=True, default=str))
                  for it in items if it["kind"] == EVIDENCE_KIND_SHEET_ROW)
    if not rows:
        return None
    return hashlib.sha256(json.dumps(rows).encode("utf-8")).hexdigest()


def _completed_count(items: list[dict]) -> int:
    return sum(1 for it in items
               if it["kind"] in (EVIDENCE_KIND_FORM_FIELD, EVIDENCE_KIND_CHECKBOX, EVIDENCE_KIND_TABLE_ROW,
                                 EVIDENCE_KIND_SHEET_CELL)
               and (it.get("structured_value") or {}).get("completed"))


def _authority_key(rec: dict) -> tuple:
    return (0 if rec["parse_status"] == "PARSED" else 1,
            _FORMAT_RANK.get(rec["file_type"], 3),
            -len(rec["stem"]),                       # the specific (bidder-named) copy over a template-named one
            -rec["package_path"].count("/"),         # the copy assembled into a (nested) submission archive
            rec["package_path"])


def link_representations(records: list[dict]) -> dict[str, dict]:
    """Deterministic logical-artifact grouping over per-file records
    ({file_id, content_hash, filename, package_path, file_type,
    lifecycle_status, parse_status, duplicate_of, items}). Returns
    file_id -> {logical_artifact_id, relationship, representation_of,
    basis}. Linking rules (never a model, never a guess about content the
    files do not share):
      * extractor byte-identical duplicate -> BYTE_IDENTICAL_DUPLICATE;
      * same filename stem, different format (DOCX source + PDF);
      * spreadsheets with an identical cell-value signature;
      * one stem a >=2-token prefix of the other AND content-token
        Jaccard >= 0.6 (an unnamed/earlier copy of a named form).
    The authoritative member is chosen by `_authority_key` (readable,
    final format, most specific name, packaged copy)."""
    base = [r for r in records if r["lifecycle_status"] in ("extracted", "failed")]
    for r in records:
        r["stem"] = _stem_tokens(r["filename"])
    parent = {r["file_id"]: r["file_id"] for r in base}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[max(ra, rb)] = min(ra, rb)

    sig = {r["file_id"]: (_sheet_signature(r["items"]) if r["file_type"] in _SPREADSHEET_TYPES else None) for r in base}
    toks = {r["file_id"]: _content_tokens(r["items"]) for r in base}
    for i, a in enumerate(base):
        for b in base[i + 1:]:
            if a["stem"] and a["stem"] == b["stem"] and a["file_type"] != b["file_type"]:
                union(a["file_id"], b["file_id"])
                continue
            if sig[a["file_id"]] and sig[a["file_id"]] == sig[b["file_id"]]:
                union(a["file_id"], b["file_id"])
                continue
            short, long_ = sorted((a["stem"], b["stem"]), key=len)
            if len(short) >= 2 and short != long_ and long_[:len(short)] == short:
                ta, tb = toks[a["file_id"]], toks[b["file_id"]]
                if ta and tb and len(ta & tb) / len(ta | tb) >= 0.6:
                    union(a["file_id"], b["file_id"])

    groups: dict[str, list[dict]] = {}
    for r in base:
        groups.setdefault(find(r["file_id"]), []).append(r)
    out: dict[str, dict] = {}
    for members in groups.values():
        auth = sorted(members, key=_authority_key)[0]
        la_id = "LA-" + hashlib.sha256(auth["content_hash"].encode("utf-8")).hexdigest()[:16]
        out[auth["file_id"]] = {"logical_artifact_id": la_id, "relationship": REP_AUTHORITATIVE,
                                "representation_of": None,
                                "basis": "sole representation" if len(members) == 1 else
                                "authoritative: readable, final format, most specific filename, packaged copy"}
        for m in members:
            if m is auth:
                continue
            if m["stem"] == auth["stem"] and m["file_type"] != auth["file_type"]:
                rel, why = REP_ALTERNATE, f"same filename stem, {m['file_type']} representation of the {auth['file_type']}"
            elif sig.get(m["file_id"]) and sig.get(m["file_id"]) == sig.get(auth["file_id"]):
                rel, why = REP_ALTERNATE, "identical worksheet cell values (re-saved / renamed copy)"
            elif _completed_count(m["items"]) == 0 < _completed_count(auth["items"]):
                rel, why = REP_TEMPLATE, "no completed fields while the authoritative copy is completed"
            else:
                rel, why = REP_DRAFT, "same logical form with differing completed content"
            out[m["file_id"]] = {"logical_artifact_id": la_id, "relationship": rel,
                                 "representation_of": auth["file_id"], "basis": why}
    for r in records:
        if r["lifecycle_status"] == "duplicate" and r.get("duplicate_of") in out:
            tgt = out[r["duplicate_of"]]
            auth_id = r["duplicate_of"] if tgt["relationship"] == REP_AUTHORITATIVE else tgt["representation_of"]
            out[r["file_id"]] = {"logical_artifact_id": tgt["logical_artifact_id"], "relationship": REP_BYTE_DUPLICATE,
                                 "representation_of": auth_id,
                                 "basis": f"byte-identical to {r['duplicate_of']}"}
        elif r["file_id"] not in out:
            out[r["file_id"]] = {"logical_artifact_id": "LA-" + hashlib.sha256(
                (r["content_hash"] + r["file_id"]).encode("utf-8")).hexdigest()[:16],
                "relationship": REP_AUTHORITATIVE, "representation_of": None,
                "basis": f"standalone {r['lifecycle_status']} entry"}
    return out


def build_submission_package(
    raw_files: list[tuple[str, bytes]],
    *,
    bid_id: int,
    organization_id: str,
    role_overrides: Optional[dict] = None,
) -> SubmissionPackage:
    """Canonicalize one bidder submission package. Reuses
    extractor.build_alignment_submission_package for ZIP-safe discovery
    (including ONE level of nested archive, CHECK-1.1), per-occurrence
    file identity, byte-identical dedup and text extraction; then parses
    each readable artifact's STRUCTURE, classifies its role, links
    alternate representations into logical artifacts, and mints bid-bound
    evidence ids for AUTHORITATIVE representations only. `role_overrides`
    maps a filename OR submission_document_id to a DOCUMENT_ROLES member."""
    import extractor
    import proposal_intelligence

    overrides = dict(role_overrides or {})
    base = extractor.build_alignment_submission_package(raw_files, include_bytes=True, expand_nested_zips=True)
    files = base["files"]
    package_digest = proposal_intelligence.compute_package_digest(extractor.build_report_manifest(files))

    records: list[dict] = []
    role_by_file: dict[str, RoleClassification] = {}
    for f in files:
        fid = f["file_id"]
        ftype = f.get("file_type") or "unknown"
        override = overrides.get(fid) or overrides.get(f.get("filename")) or overrides.get(f.get("package_path"))
        parse_status = {"extracted": "PARSED", "failed": "FAILED", "unsupported": "UNSUPPORTED",
                        "rejected": "REJECTED", "duplicate": "DUPLICATE"}.get(f["lifecycle_status"], "FAILED")
        structure = {"parser": None, "page_count": None, "sections": (), "items": [], "sheets": []}
        if parse_status == "PARSED":
            structure = parse_document_structure(f.get("_bytes") or b"", ftype, f.get("text") or "")
        titles = [s.title for s in structure["sections"]]
        role = classify_document_role(f.get("filename") or "", f.get("text") or "", file_type=ftype,
                                      section_titles=titles, override=override)
        role_by_file[fid] = role
        records.append({"file_id": fid, "content_hash": f["content_hash"], "filename": f.get("filename") or "",
                        "package_path": f.get("package_path") or "", "file_type": ftype,
                        "lifecycle_status": f["lifecycle_status"], "parse_status": parse_status,
                        "duplicate_of": f.get("duplicate_of_file_id"), "items": structure["items"],
                        "structure": structure, "file": f, "override": override})

    links = link_representations(records)

    documents: list[SubmissionDocument] = []
    all_items: list[EvidenceItem] = []
    for rec in records:
        f, fid, structure = rec["file"], rec["file_id"], rec["structure"]
        link = links[fid]
        role = role_by_file[fid]
        auth_id = link["representation_of"]
        if auth_id and not rec["override"]:
            # A representation IS the same logical artifact: it carries the
            # authoritative member's role (a byte duplicate always did).
            role = role_by_file[auth_id]
        authoritative = link["relationship"] == REP_AUTHORITATIVE
        provenance_base = {
            "file_id": fid, "content_hash": f["content_hash"], "filename": f.get("filename"),
            "package_path": f.get("package_path"), "file_type": rec["file_type"], "parser": structure["parser"],
            "logical_artifact_id": link["logical_artifact_id"], "contract_version": CHECK1_CONTRACT_VERSION,
        }
        count = 0
        if authoritative:
            seen_ids = set()
            for raw in structure["items"]:
                eid = derive_evidence_id(bid_id, fid, raw["kind"], raw["locator"])
                if eid in seen_ids:
                    continue
                seen_ids.add(eid)
                all_items.append(EvidenceItem(
                    evidence_id=eid, bid_id=bid_id, submission_document_id=fid, document_role=role.role,
                    kind=raw["kind"], location={k: v for k, v in raw["location"].items() if v not in (None, [], "")} or {},
                    content=raw["content"], structured_value=raw["structured_value"],
                    provenance=dict(provenance_base, locator=raw["locator"]),
                ))
                count += 1
        unusable = f.get("unusable_reason")
        if not authoritative and rec["lifecycle_status"] != "duplicate":
            unusable = (f"{link['relationship']} of {auth_id} ({link['basis']}); evidence is registered from "
                        f"the authoritative representation only")
        documents.append(SubmissionDocument(
            submission_document_id=fid, content_hash=f["content_hash"], filename=f.get("filename") or "",
            package_path=f.get("package_path") or "", file_type=rec["file_type"],
            lifecycle_status=f["lifecycle_status"],
            included=bool(f.get("included")) and authoritative, role=role, parse_status=rec["parse_status"],
            parser=structure["parser"], page_count=structure.get("page_count"),
            sheets=tuple(s["sheet"] for s in structure.get("sheets") or []),
            sections=tuple(structure["sections"]), evidence_count=count,
            duplicate_of=f.get("duplicate_of_file_id"), unusable_reason=unusable,
            logical_artifact_id=link["logical_artifact_id"], representation_relationship=link["relationship"],
            representation_of=auth_id, representation_basis=link["basis"],
        ))

    manifest = tuple(extractor.build_report_manifest(files))
    return SubmissionPackage(bid_id=bid_id, organization_id=organization_id, package_digest=package_digest,
                             documents=tuple(documents), registry=SubmissionEvidenceRegistry(bid_id, all_items),
                             manifest=manifest)


def package_from_persisted_rows(bid_id: int, organization_id: str, package_digest: str,
                                document_rows: list[dict], evidence_rows: list[dict],
                                manifest: Iterable = ()) -> SubmissionPackage:
    """Rebuild a SubmissionPackage from migration-021 rows (a FRESH read,
    no in-memory state). Pure. Every row must belong to `bid_id` -- a
    foreign row fails closed (CrossBidEvidenceError)."""
    docs = []
    for r in sorted(document_rows, key=lambda x: x["id"] if "id" in x else 0):
        if int(r["bid_id"]) != int(bid_id):
            raise CrossBidEvidenceError(f"document row belongs to bid {r['bid_id']}, not {bid_id}")
        role = RoleClassification(r["document_role"], r["role_confidence"], r["role_basis"],
                                  tuple(r.get("secondary_roles") or ()))
        secs = tuple(ProposalSection(s["section_id"], s.get("number"), s["title"], s["level"], s.get("parent_id"),
                                     s.get("page_start"), s.get("page_end")) for s in (r.get("sections") or []))
        docs.append(SubmissionDocument(
            submission_document_id=r["submission_document_id"], content_hash=r["content_hash"],
            filename=r["filename"], package_path=r["package_path"], file_type=r["file_type"],
            lifecycle_status=r["lifecycle_status"], included=bool(r["included"]), role=role,
            parse_status=r["parse_status"], parser=None, page_count=r.get("page_count"),
            sheets=tuple(r.get("sheets") or ()), sections=secs, evidence_count=int(r.get("evidence_count") or 0),
            duplicate_of=r.get("duplicate_of"), unusable_reason=r.get("unusable_reason"),
            logical_artifact_id=r.get("logical_artifact_id"),
            representation_relationship=r.get("representation_relationship") or REP_AUTHORITATIVE,
            representation_of=r.get("representation_of"), representation_basis=r.get("representation_basis")))
    items = []
    for r in evidence_rows:
        items.append(EvidenceItem(
            evidence_id=r["evidence_id"], bid_id=int(r["bid_id"]), submission_document_id=r["submission_document_id"],
            document_role=r["document_role"], kind=r["kind"], location=r.get("location") or {},
            content=r.get("content") or "", structured_value=r.get("structured_value"),
            provenance=r.get("provenance") or {}))
    return SubmissionPackage(bid_id=bid_id, organization_id=organization_id, package_digest=package_digest,
                             documents=tuple(docs), registry=SubmissionEvidenceRegistry(bid_id, items),
                             manifest=tuple(manifest or ()))


# ═══════════════════════════════════════════════════════════════════════
# 6. Expected evidence location (sections 5 / 6)
# ═══════════════════════════════════════════════════════════════════════

_PORTAL_NATIVE_RE = re.compile(
    r"\b(?:ariba|merx|bonfire|biddingo|bids\s*&\s*tenders|jaggaer|coupa|buyandsell|canadabuys|sap\s+business\s+network)\b"
    r"|\b(?:complete[ds]?|submit(?:ted)?|answer(?:ed)?|respond(?:ed)?|enter(?:ed)?)\s+(?:directly\s+)?(?:in|into|within|through|via|on)\s+(?:the\s+)?"
    r"(?:\w+\s+){0,2}(?:portal|platform|e-?procurement\s+system|online\s+(?:form|questionnaire))\b"
    r"|\b(?:digiti[sz]ed|online|electronic|portal)\s+questionnaire\b|\bportal[\s-]native\b",
    re.IGNORECASE)
# CHECK-1.1 (real Calgary 26-1603): "Complete the 'Social Procurement
# Questionnaire' as requested" -- a buyer-NAMED questionnaire with no
# document anchor (no appendix / attachment / template) is, in an
# e-procurement event (SAP Ariba here), answered inside the portal. It is
# therefore a POSSIBLE portal-native response: never provable present or
# absent from uploaded files, never MISSING_FROM_PACKAGE.
_NAMED_QUESTIONNAIRE_RE = re.compile(
    r"\bcomplet\w*\s+(?:the\s+|your\s+|a\s+)?[‘'\"“]?(?:[\w&/-]+\s+){0,5}questionnaire\b", re.IGNORECASE)
_DOCUMENT_ANCHOR_RE = re.compile(
    r"\b(?:appendix|attachment|attached|schedule|annex|exhibit|template|spreadsheet|workbook|excel)\b", re.IGNORECASE)
_PRICING_LOCATION_RE = re.compile(
    r"\bpric(?:e|ing)\s+(?:form|schedule|sheet|workbook|template)\b|\bexcel\s+(?:spreadsheet|pricing|workbook)\b"
    r"|\b(?:yellow|input)\s+cells?\b|\bpricing\s+formula\b|\brate\s+card\b|\bcost\s+(?:breakdown\s+)?form\b",
    re.IGNORECASE)
_PRICING_TOPIC_RE = re.compile(
    r"\bpric(?:e|es|ing)\b|\blump\s+sum\b|\brates?\b|\bcosts?\b|\bfees?\b|\btravel\s+(?:costs?|expenses?)\b", re.IGNORECASE)
_SUBMISSION_FORM_RE = re.compile(
    r"\bsubmission\s+form\b|\bform\s+of\s+(?:tender|offer|proposal)\b|\bdeclar\w+|\battest\w*|\bconflicts?\s+of\s+interest\b"
    r"|\backnowledg\w*\s+(?:of\s+|receipt\s+of\s+)?addend\w*|\bauthori[sz]ed\s+(?:signatory|representative|to\s+bind)\b"
    r"|\bpower\s+to\s+contract\b|\bbinding\s+offer\b|\bsign(?:ed|ature)\b|\bcertif(?:y|ies|ication\s+that)\b", re.IGNORECASE)
_MULTI_PARTY_RE = re.compile(
    r"\bmulti[\s-]*party\b|\bjoint\s+(?:venture|bid|submission|proposal)\b|\bconsorti(?:um|a)\b|\bteaming\b"
    r"|\blead\s+proponent\b|\bsubmitted\s+jointly\b|\bmore\s+than\s+one\s+(?:party|proponent|legal\s+entit)",
    re.IGNORECASE)
# Deliberately NOT multi-party signals: "subcontractor"/"sub-consultant"
# mentions in general conditions (assignment, substitution, named
# individuals) describe contract performance, not a joint-bid structure.
_SOCIAL_RE = re.compile(r"\bsocial\s+(?:procurement|value|benefit)\b|\bindigenous\b|\bcommunity\s+benefit", re.IGNORECASE)
_CERT_RE = re.compile(
    r"\bcertificates?\s+of\s+insurance\b|\bproof\s+of\s+(?:insurance|registration|coverage)\b|\bwcb\b"
    r"|\bworkers[’']?\s+compensation\b|\bgood\s+standing\b|\bregist(?:ered|ration)\b[^.]{0,40}\b(?:registr(?:y|ies)|corporations?\s+act)\b"
    r"|\blicen[cs]e\b|\baccredit\w*", re.IGNORECASE)
_RESUME_RE = re.compile(r"\br[ée]sum[ée]s?\b|\bcvs?\b|\bcurricul(?:um|a)\s+vitae\b|\bkey\s+personnel\b|\bnamed\s+(?:individuals?|resources?)\b", re.IGNORECASE)
_ORG_CHART_RE = re.compile(r"\borgani[sz]ation(?:al)?\s+chart\b|\breporting\s+structure\b", re.IGNORECASE)
_TECHNICAL_RE = re.compile(
    r"\bmethodolog\w*|\bapproach\b|\bwork\s*plan\b|\bexperience\b|\bqualifications?\b|\bdeliver\w*\b|\bdesign\b"
    r"|\bfacilitat\w*|\bprogram(?:me)?s?\b|\bcohorts?\b|\bproposal\s+(?:format|length|submission\s+length)\b"
    r"|\bpages?\b|\bgoals?\b|\bobjectives?\b|\bsuccess\s+measures?\b|\baccount\s+manager\b|\bteam\b", re.IGNORECASE)
_POST_AWARD_RE = re.compile(
    r"\b(?:during\s+the\s+term|upon\s+receipt\s+of\s+written\s+instructions|before\s+(?:the\s+)?end\s+of\s+(?:the\s+)?term"
    r"|post[\s-]termination|on[\s-]site|immediately\s+notify|within\s+\d+\s+(?:hours|business\s+days)|invoices?"
    r"|hold\s*back|set\s+off|written\s+amendment|change\s+requests?|found\s+substances?|safety\s+briefing)\b",
    re.IGNORECASE)

EVIDENCE_CATEGORY_PRICING = "PRICING"
EVIDENCE_CATEGORY_DECLARATION = "DECLARATION"
EVIDENCE_CATEGORY_MULTI_PARTY = "MULTI_PARTY_STRUCTURE"
EVIDENCE_CATEGORY_SOCIAL = "SOCIAL_PROCUREMENT"
EVIDENCE_CATEGORY_CERTIFICATION = "CERTIFICATION_OR_PROOF"
EVIDENCE_CATEGORY_PERSONNEL = "PERSONNEL"
EVIDENCE_CATEGORY_TECHNICAL = "TECHNICAL_RESPONSE"
EVIDENCE_CATEGORY_CONTRACT_OBLIGATION = "CONTRACT_OBLIGATION"
EVIDENCE_CATEGORY_PORTAL = "PORTAL_RESPONSE"
EVIDENCE_CATEGORY_UNKNOWN = "UNKNOWN"


@dataclass(frozen=True)
class ExpectedEvidence:
    """Where a requirement's evidence is expected -- a set of ROLES (never
    one fixed filename), whether that location is flexible, and whether
    it may be answered natively in a buyer portal."""

    roles: tuple
    evidence_category: str
    location_basis: str
    flexible: bool
    portal_native: bool
    rationale: str

    def to_dict(self) -> dict:
        return {"roles": list(self.roles), "evidence_category": self.evidence_category,
                "location_basis": self.location_basis, "flexible": self.flexible,
                "portal_native": self.portal_native, "rationale": self.rationale}


def derive_expected_evidence(requirement: dict) -> ExpectedEvidence:
    """Deterministic expected-evidence roles for ONE canonical RFP
    requirement. Order of `roles` is preference order (primary first).
    Never invents a location: when nothing in the requirement supports a
    specific artifact, the result is FLEXIBLE (all substantive roles,
    basis NOT_STATED)."""
    req = requirement or {}
    explicit = req.get("expected_evidence_roles")
    text = " ".join(str(req.get(k) or "") for k in ("description", "evidence", "rfso_ref", "notes"))
    category = (req.get("category") or "").strip().lower()
    portal = (bool(_PORTAL_NATIVE_RE.search(text)) or bool(req.get("portal_native"))
              or bool(_NAMED_QUESTIONNAIRE_RE.search(text) and not _DOCUMENT_ANCHOR_RE.search(text)))

    if isinstance(explicit, (list, tuple)) and explicit:
        roles = tuple(r for r in explicit if r in DOCUMENT_ROLES and r != ROLE_UNKNOWN)
        if roles:
            return ExpectedEvidence(roles, EVIDENCE_CATEGORY_UNKNOWN, LOCATION_BASIS_EXPLICIT,
                                    False, portal, "requirement carries explicit expected evidence roles")

    semantic = cp.classify_semantic_type(req.get("description") or "")

    def _ee(roles, cat, basis, flexible, why):
        if portal and cat != EVIDENCE_CATEGORY_PORTAL:
            why += "; the requirement text indicates a portal-native response"
        return ExpectedEvidence(tuple(dict.fromkeys(roles)), cat, basis, flexible, portal, why)

    if portal and not _PRICING_LOCATION_RE.search(text):
        # No document role is expected: the response lives in the portal.
        # Candidates are still searched across every artifact (an export
        # may have been uploaded), but no uploaded file's presence or
        # absence can prove the portal response was or was not completed.
        return ExpectedEvidence((), EVIDENCE_CATEGORY_PORTAL, LOCATION_BASIS_STATED, True, True,
                                "response is completed natively in the buyer portal; the uploaded files may not "
                                "contain it and cannot prove whether it was completed")
    if _MULTI_PARTY_RE.search(text) and not re.search(r"\bnot\s+(?:assign|subcontract)", text, re.IGNORECASE):
        return _ee((ROLE_MULTI_PARTY_FORM, ROLE_SUBMISSION_FORM, ROLE_TECHNICAL_PROPOSAL),
                   EVIDENCE_CATEGORY_MULTI_PARTY, LOCATION_BASIS_STATED, False,
                   "joint/consortium structure is evidenced by the multi-party form (or submission form)")
    if _PRICING_LOCATION_RE.search(text):
        return _ee((ROLE_PRICING_FORM,), EVIDENCE_CATEGORY_PRICING, LOCATION_BASIS_STATED, False,
                   "requirement names the pricing form / spreadsheet as the pricing location")
    if category == "financial" and _PRICING_TOPIC_RE.search(text) and not _POST_AWARD_RE.search(text):
        return _ee((ROLE_PRICING_FORM, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_PRICING, LOCATION_BASIS_INFERRED, False,
                   "financial/pricing requirement -- pricing evidence belongs in the pricing form")
    if _SOCIAL_RE.search(text):
        return _ee((ROLE_SOCIAL_PROCUREMENT_RESPONSE, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_SOCIAL,
                   LOCATION_BASIS_STATED, False, "social procurement response")
    if _CERT_RE.search(text) and not _POST_AWARD_RE.search(text):
        return _ee((ROLE_CERTIFICATE, ROLE_EVIDENCE_ATTACHMENT, ROLE_SUBMISSION_FORM), EVIDENCE_CATEGORY_CERTIFICATION,
                   LOCATION_BASIS_INFERRED, True,
                   "proof/registration/certification -- may be an attachment or declared on the submission form")
    if (_SUBMISSION_FORM_RE.search(text) and not _POST_AWARD_RE.search(text)):
        return _ee((ROLE_SUBMISSION_FORM, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_DECLARATION,
                   LOCATION_BASIS_STATED if re.search(r"submission\s+form|declar", text, re.IGNORECASE) else LOCATION_BASIS_INFERRED,
                   False, "declaration/attestation -- expected on the submission form")
    if _RESUME_RE.search(text):
        return _ee((ROLE_RESUME, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_PERSONNEL, LOCATION_BASIS_INFERRED, True,
                   "personnel evidence -- resumes or the technical proposal team section")
    if _ORG_CHART_RE.search(text):
        return _ee((ROLE_ORGANIZATION_CHART, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_PERSONNEL,
                   LOCATION_BASIS_INFERRED, True, "organization chart")
    if (semantic in (cp.SEMANTIC_CONTRACTUAL_OBLIGATION, cp.SEMANTIC_COMMERCIAL_OBLIGATION)
            or _POST_AWARD_RE.search(text) or category in ("financial", "supporting")) and not category == "rated":
        return _ee((ROLE_SUBMISSION_FORM, ROLE_TECHNICAL_PROPOSAL), EVIDENCE_CATEGORY_CONTRACT_OBLIGATION,
                   LOCATION_BASIS_INFERRED, True,
                   "contract/performance obligation -- at submission typically evidenced only by acceptance of "
                   "terms; the RFP does not name a location")
    if category == "rated" or semantic in (cp.SEMANTIC_RESPONSE_PROMPT, cp.SEMANTIC_EVALUATION_CRITERION) or _TECHNICAL_RE.search(text):
        return _ee((ROLE_TECHNICAL_PROPOSAL,), EVIDENCE_CATEGORY_TECHNICAL, LOCATION_BASIS_INFERRED, False,
                   "technical/rated response -- expected in the technical proposal narrative")
    return _ee(_SUBSTANTIVE_ROLES, EVIDENCE_CATEGORY_UNKNOWN, LOCATION_BASIS_NOT_STATED, True,
               "the requirement does not state where evidence belongs -- location flexible/unknown")


# ═══════════════════════════════════════════════════════════════════════
# 7. RFP -> submission candidate mapping (section 7) -- NO adjudication
# ═══════════════════════════════════════════════════════════════════════

_STOPWORDS = frozenset("""
a an and are as at be been by can for from has have if in into is it its may must not of on or per
shall should such that the their them there these this those to under upon was were will with within
without any all each other than also including include includes provide provided proponent proponents
consultant consultants city bidder bidders supplier suppliers required requirement requirements
""".split())


def _terms(text: str) -> set[str]:
    return {w[:7] for w in re.findall(r"[a-z][a-z0-9]{2,}", (text or "").lower()) if w not in _STOPWORDS}


@dataclass(frozen=True)
class EvidenceCandidate:
    evidence_id: str
    submission_document_id: str
    document_role: str
    kind: str
    score: float
    match_basis: str
    role_expected: bool
    matched_terms: tuple
    citation: str

    def to_dict(self) -> dict:
        return {"evidence_id": self.evidence_id, "submission_document_id": self.submission_document_id,
                "document_role": self.document_role, "kind": self.kind, "score": self.score,
                "match_basis": self.match_basis, "role_expected": self.role_expected,
                "matched_terms": list(self.matched_terms), "citation": self.citation}


@dataclass(frozen=True)
class RequirementEvidenceMapping:
    """Candidates only. `artifact_status` describes ARTIFACT PRESENCE for
    the expected roles; it is NOT an ADDRESSED/PARTIAL/MISSING verdict."""

    req_id: str
    requirement_id: Optional[int]
    category: Optional[str]
    expected: ExpectedEvidence
    artifact_status: str
    expected_role_documents: tuple
    candidates: tuple
    other_artifact_candidates: tuple
    notes: tuple = ()
    primary_role_present: bool = False

    @property
    def absence_claim_permitted(self) -> bool:
        return self.artifact_status == MISSING_FROM_PACKAGE

    def to_dict(self) -> dict:
        return {"req_id": self.req_id, "requirement_id": self.requirement_id, "category": self.category,
                "expected": self.expected.to_dict(), "artifact_status": self.artifact_status,
                "absence_claim_permitted": self.absence_claim_permitted,
                "primary_role_present": self.primary_role_present,
                "expected_role_documents": list(self.expected_role_documents),
                "candidates": [c.to_dict() for c in self.candidates],
                "other_artifact_candidates": [c.to_dict() for c in self.other_artifact_candidates],
                "notes": list(self.notes)}


def artifact_status_for(expected: ExpectedEvidence, package: SubmissionPackage) -> tuple[str, tuple, tuple]:
    """(status, expected_role_document_ids, notes). MISSING_FROM_PACKAGE
    is emitted ONLY when: the location is specific (not flexible), not
    possibly portal-native, no member document of ANY expected role
    (primary or embedded/secondary) exists -- readable or not -- and no
    member document is UNKNOWN-role or unreadable (either could be the
    artifact)."""
    docs = package.documents_with_roles(expected.roles)
    notes = []
    readable = [d for d in docs if d.readable]
    if readable:
        embedded = [d for d in readable if d.document_role not in expected.roles]
        if embedded and len(embedded) == len(readable):
            notes.append("expected-role content is present only as an embedded section of another artifact")
        primary = expected.roles[0]
        if not any(primary in d.all_roles for d in readable):
            notes.append(f"primary expected artifact ({primary}) is not in the package; only an alternate "
                         f"expected role is present")
        return ARTIFACT_PRESENT, tuple(d.submission_document_id for d in docs), tuple(notes)
    if docs:
        notes.append("an artifact of the expected role is present but could not be read")
        return NOT_VERIFIABLE_FROM_FILES, tuple(d.submission_document_id for d in docs), tuple(notes)
    if expected.portal_native:
        notes.append("may have been completed natively in the buyer portal; not provable from uploaded files")
        return POSSIBLY_PORTAL_NATIVE, (), tuple(notes)
    members = package.member_documents()
    if any(d.document_role == ROLE_UNKNOWN or not d.readable for d in members):
        notes.append("package contains unclassified or unreadable artifacts that could hold this evidence")
        return NOT_VERIFIABLE_FROM_FILES, (), tuple(notes)
    if expected.flexible:
        return FLEXIBLE_LOCATION, (), ("the RFP does not fix an evidence location",)
    return MISSING_FROM_PACKAGE, (), ("no submitted artifact of any expected role is in the package",)


def _score_item(q: set, item: EvidenceItem, expected_roles: set, is_pricing: bool) -> tuple[float, str, tuple]:
    text_terms = _terms(item.content)
    sv = item.structured_value or {}
    for k in ("label", "column_header", "value"):
        if isinstance(sv.get(k), str):
            text_terms |= _terms(sv[k])
    loc = item.location or {}
    heading_terms = _terms(" ".join(loc.get("section_path") or []) + " " + (loc.get("sheet") or ""))
    matched = (q & text_terms) | (q & heading_terms)
    if not q:
        return 0.0, "NONE", ()
    score = len(q & text_terms) / len(q) + 0.5 * len(q & heading_terms) / len(q)
    basis = "LEXICAL" if matched else "NONE"
    if item.document_role in expected_roles:
        score += 0.25
    else:
        # Evidence inside an EMBEDDED part of an expected role (e.g. the
        # "PART 5 -- PRICING" section of a technical & commercial
        # proposal) earns the same role credit as a standalone artifact.
        heading_text = " ".join(loc.get("section_path") or [])
        if any(role in expected_roles and rx.search(heading_text) for role, rx in _EMBEDDED_ROLE_HEADING_RULES):
            score += 0.25
            basis = "LEXICAL+EMBEDDED_SECTION" if matched else "EMBEDDED_SECTION"
    if is_pricing and item.kind == EVIDENCE_KIND_SHEET_CELL and sv.get("is_input_cell"):
        score += 0.35
        basis = "LEXICAL+STRUCTURAL" if matched else "STRUCTURAL_PRICING_INPUT"
    elif is_pricing and item.kind in (EVIDENCE_KIND_SHEET_CELL, EVIDENCE_KIND_TABLE_ROW) and isinstance(sv.get("value"), (int, float)):
        score += 0.15
        basis = "LEXICAL+STRUCTURAL" if matched else "STRUCTURAL_PRICING_VALUE"
    return round(score, 4), basis, tuple(sorted(matched))


def map_requirement_to_submission(requirement: dict, package: SubmissionPackage, *,
                                  top_k: int = DEFAULT_TOP_K, other_k: int = 3) -> RequirementEvidenceMapping:
    req = copy.deepcopy(requirement or {})
    expected = derive_expected_evidence(req)
    status, doc_ids, notes = artifact_status_for(expected, package)
    q = _terms(" ".join(str(req.get(k) or "") for k in ("description", "evidence")))
    expected_set = set(expected.roles)
    is_pricing = expected.evidence_category == EVIDENCE_CATEGORY_PRICING
    expected_doc_ids = set(doc_ids)

    in_role, other = [], []
    for item in package.registry:
        in_expected_doc = item.submission_document_id in expected_doc_ids
        score, basis, matched = _score_item(q, item, expected_set, is_pricing and in_expected_doc)
        if basis == "NONE":
            continue
        doc = package.document(item.submission_document_id)
        cand = EvidenceCandidate(item.evidence_id, item.submission_document_id, item.document_role, item.kind,
                                 score, basis, in_expected_doc, matched, format_citation(doc.filename, item.location))
        (in_role if in_expected_doc else other).append(cand)
    key = lambda c: (-c.score, c.evidence_id)
    in_role.sort(key=key)
    other.sort(key=key)
    return RequirementEvidenceMapping(
        req_id=req.get("req_id") or "", requirement_id=req.get("id"), category=req.get("category"),
        expected=expected, artifact_status=status, expected_role_documents=doc_ids,
        candidates=tuple(in_role[:top_k]), other_artifact_candidates=tuple(other[:other_k]), notes=notes,
        primary_role_present=bool(expected.roles) and any(
            package.document(did).readable and expected.roles[0] in package.document(did).all_roles
            for did in doc_ids))


def map_requirements_to_submission(requirements: list[dict], package: SubmissionPackage, *,
                                   top_k: int = DEFAULT_TOP_K) -> list[RequirementEvidenceMapping]:
    """Bounded candidate mapping for every requirement. Requirements are
    deep-copied and never mutated -- buyer-side canonical truth is input
    only."""
    return [map_requirement_to_submission(r, package, top_k=top_k) for r in copy.deepcopy(list(requirements or []))]


@dataclass(frozen=True)
class CriterionSectionMapping:
    criterion_key: str
    criterion_label: str
    category: Optional[str]
    weight: Optional[object]
    candidate_sections: tuple   # ({submission_document_id, section_id, label, page_start, page_end, score, matched_terms})
    candidate_evidence: tuple   # EvidenceCandidate

    def to_dict(self) -> dict:
        return {"criterion_key": self.criterion_key, "criterion_label": self.criterion_label,
                "category": self.category, "weight": self.weight,
                "candidate_sections": [dict(s) for s in self.candidate_sections],
                "candidate_evidence": [c.to_dict() for c in self.candidate_evidence]}


def map_evaluation_criteria_to_sections(criteria: list[dict], package: SubmissionPackage, *,
                                        top_k: int = 3) -> list[CriterionSectionMapping]:
    """Evaluation criteria -> technical proposal SECTIONS (heading match
    weighted above body match). Criterion identity uses CI-1.1's
    scoped_criterion_map_key -- never the label alone."""
    tech_docs = package.documents_with_roles([ROLE_TECHNICAL_PROPOSAL], include_secondary=False)
    out = []
    for crit in copy.deepcopy(list(criteria or [])):
        label = crit.get("criterion_label") or crit.get("criterion") or crit.get("label") or crit.get("description") or ""
        prompt = " ".join(str(crit.get(k) or "") for k in ("prompt", "response_prompt", "description", "guidance"))
        q_head = _terms(label)
        q_all = q_head | _terms(prompt)
        scored_sections = []
        for d in tech_docs:
            items = package.registry.for_document(d.submission_document_id)
            body_terms: dict[str, set] = {}
            for it in items:
                sid = (it.location or {}).get("section_id")
                if sid:
                    body_terms[sid] = body_terms.get(sid, set()) | _terms(it.content)
            for s in d.sections:
                h = _terms(s.title)
                hm, bm = q_head & h, q_all & body_terms.get(s.section_id, set())
                if not hm and not bm:
                    continue
                score = (2.0 * len(hm) / max(len(q_head), 1)) + (len(bm) / max(len(q_all), 1))
                scored_sections.append({"submission_document_id": d.submission_document_id,
                                        "section_id": s.section_id, "label": s.label,
                                        "page_start": s.page_start, "page_end": s.page_end,
                                        "score": round(score, 4), "matched_terms": sorted(hm | bm)})
        scored_sections.sort(key=lambda x: (-x["score"], x["submission_document_id"], x["section_id"]))
        mapping = map_requirement_to_submission(
            {"req_id": label, "description": f"{label} {prompt}", "category": "Rated"}, package, top_k=top_k)
        out.append(CriterionSectionMapping(
            criterion_key=cp.scoped_criterion_map_key(crit.get("category"), label), criterion_label=label,
            category=crit.get("category"), weight=crit.get("weight"),
            candidate_sections=tuple(scored_sections[:top_k]), candidate_evidence=mapping.candidates))
    return out


# ═══════════════════════════════════════════════════════════════════════
# 8. Known-bad regression guard: absence claims vs. package reality
# ═══════════════════════════════════════════════════════════════════════

_ABSENCE_RE = re.compile(
    r"\b(?:missing|absent|not\s+(?:provided|included|found|present|submitted|attached|completed)|no\s+evidence\s+of"
    r"|does\s+not\s+(?:include|contain|provide)|could\s+not\s+(?:locate|find)|omitted|lacks?)\b", re.IGNORECASE)
_ARTIFACT_MENTION_RULES: tuple = (
    (ROLE_PRICING_FORM, re.compile(r"\bpric(?:e|ing)\s+(?:form|schedule|sheet|workbook)\b|\bpricing\b|\bprice\b|\bcost\s+(?:form|proposal|breakdown)\b|\blump\s+sum\b", re.IGNORECASE)),
    (ROLE_MULTI_PARTY_FORM, re.compile(r"\bmulti[\s-]*party\b|\bconsortium\b|\bjoint\s+(?:venture|bid|submission)\b|\bpartner(?:ship)?\s+(?:form|confirmation|structure)\b", re.IGNORECASE)),
    (ROLE_SUBMISSION_FORM, re.compile(r"\bsubmission\s+form\b|\bdeclarations?\b|\battestations?\b|\bconflict\s+of\s+interest\b|\backnowledg\w+\s+of\s+addend\w*|\bsigned\s+form\b|\badministrative\b", re.IGNORECASE)),
    (ROLE_RESUME, re.compile(r"\br[ée]sum[ée]s?\b|\bcvs?\b", re.IGNORECASE)),
    (ROLE_CERTIFICATE, re.compile(r"\bcertificates?\b|\bproof\s+of\s+insurance\b", re.IGNORECASE)),
    (ROLE_ORGANIZATION_CHART, re.compile(r"\borg(?:ani[sz]ation(?:al)?)?\s+chart\b", re.IGNORECASE)),
)

ABSENCE_CONTRADICTED_BY_PACKAGE = "CONTRADICTED_BY_PACKAGE"
ABSENCE_NOT_CONTRADICTED = "NOT_CONTRADICTED_BY_PACKAGE"
ABSENCE_NOT_AN_ARTIFACT_CLAIM = "NOT_AN_ARTIFACT_ABSENCE_CLAIM"


def screen_absence_claim(claim, package: SubmissionPackage) -> dict:
    """Structural guard for any absence/"missing" claim (from a legacy
    review, a future CHECK analyzer, or a model): if the claim names an
    artifact role that IS present in the package (as a primary or
    embedded role), the claim is CONTRADICTED_BY_PACKAGE and must not be
    surfaced as a missing-artifact finding. This never judges content
    quality -- a present Price Form may still be incomplete; that is a
    CHECK-2 question, answered from the Price Form's own evidence."""
    text = claim if isinstance(claim, str) else " ".join(
        str((claim or {}).get(k) or "") for k in ("title", "issue", "message", "finding"))
    if not _ABSENCE_RE.search(text):
        return {"claim": text, "verdict": ABSENCE_NOT_AN_ARTIFACT_CLAIM, "roles": [], "present_documents": []}
    roles = [role for role, rx in _ARTIFACT_MENTION_RULES if rx.search(text)]
    if not roles:
        return {"claim": text, "verdict": ABSENCE_NOT_AN_ARTIFACT_CLAIM, "roles": [], "present_documents": []}
    present = package.documents_with_roles(roles)
    return {
        "claim": text,
        "verdict": ABSENCE_CONTRADICTED_BY_PACKAGE if present else ABSENCE_NOT_CONTRADICTED,
        "roles": roles,
        "present_documents": [{"submission_document_id": d.submission_document_id, "filename": d.filename,
                               "document_role": d.document_role} for d in present],
    }


# ═══════════════════════════════════════════════════════════════════════
# 9. Reuse bridges (section 14) and persistence payload (section 15)
# ═══════════════════════════════════════════════════════════════════════

def to_proposal_source_ref(item: EvidenceItem) -> dict:
    """Project an EvidenceItem into the ProposalSourceRef shape
    (analyst._build_proposal_source_ref: file_id/content_hash/filename/
    package_path/file_type/section) that SectionResponseBrief's tier-2
    `proposal_source_refs` already carries -- plus the CHECK-1
    evidence_id and exact page/sheet/cell, so section_drafting's
    evidence_id_registry / reconcile_material_claims apply unchanged to a
    HUMAN-written submission's evidence."""
    prov = item.provenance or {}
    loc = item.location or {}
    ref = {k: prov[k] for k in ("file_id", "content_hash", "filename", "package_path", "file_type") if prov.get(k)}
    if loc.get("section_title"):
        num = loc.get("section_number")
        ref["section"] = f"{num} {loc['section_title']}" if num and num not in loc["section_title"] else loc["section_title"]
    for k in ("page", "sheet", "cell", "row"):
        if loc.get(k) is not None:
            ref[k] = loc[k]
    ref["evidence_id"] = item.evidence_id
    ref["document_role"] = item.document_role
    return ref


def candidate_proposal_source_refs(mapping: RequirementEvidenceMapping, package: SubmissionPackage) -> list[dict]:
    return [to_proposal_source_ref(package.registry.get(c.evidence_id, bid_id=package.bid_id))
            for c in mapping.candidates]


def build_persistence_payload(package: SubmissionPackage) -> dict:
    """Rows in the exact shape of migrations/021_submission_evidence_
    registry.sql's `submission_documents` / `submission_evidence_items`
    (written by `create_submission_evidence_bundle`). Pure -- building
    the payload writes nothing."""
    docs = []
    for d in package.documents:
        docs.append({
            "submission_document_id": d.submission_document_id, "content_hash": d.content_hash,
            "filename": d.filename, "package_path": d.package_path, "file_type": d.file_type,
            "lifecycle_status": d.lifecycle_status, "included": d.included,
            "document_role": d.document_role, "role_confidence": d.role.confidence,
            "role_basis": d.role.basis, "secondary_roles": list(d.role.secondary_roles),
            "parse_status": d.parse_status, "page_count": d.page_count, "sheets": list(d.sheets),
            "sections": [s.to_dict() for s in d.sections], "evidence_count": d.evidence_count,
            "duplicate_of": d.duplicate_of, "unusable_reason": d.unusable_reason,
            "logical_artifact_id": d.logical_artifact_id or d.submission_document_id,
            "representation_relationship": d.representation_relationship,
            "representation_of": d.representation_of, "representation_basis": d.representation_basis,
        })
    items = []
    for it in package.registry:
        items.append({
            "evidence_id": it.evidence_id, "submission_document_id": it.submission_document_id,
            "document_role": it.document_role, "kind": it.kind, "location": it.location,
            "content": it.content, "structured_value": it.structured_value, "provenance": it.provenance,
        })
    return {"bid_id": package.bid_id, "package_digest": package.package_digest,
            "contract_version": package.contract_version, "documents": docs, "evidence_items": items}


def package_fingerprint(package: SubmissionPackage) -> str:
    """Deterministic digest of the canonical package + registry (the same
    sorted-key JSON / sha256 convention as compute_package_digest)."""
    payload = build_persistence_payload(package)
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


__all__ = [n for n in dir() if n.isupper() or n in {
    "RoleClassification", "classify_document_role", "detect_heading", "ProposalSection",
    "parse_pdf_structure", "parse_docx_structure", "parse_xlsx_structure", "parse_text_rows_structure",
    "parse_document_structure", "derive_evidence_id", "EvidenceItem", "format_citation",
    "SubmissionEvidenceRegistry", "CrossBidEvidenceError", "SubmissionDocument", "SubmissionPackage",
    "build_submission_package", "ExpectedEvidence", "derive_expected_evidence", "EvidenceCandidate",
    "RequirementEvidenceMapping", "artifact_status_for", "map_requirement_to_submission",
    "map_requirements_to_submission", "CriterionSectionMapping", "map_evaluation_criteria_to_sections",
    "screen_absence_claim", "to_proposal_source_ref", "candidate_proposal_source_refs",
    "build_persistence_payload", "package_fingerprint", "link_representations",
    "package_from_persisted_rows"}]
