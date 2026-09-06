"""
Bid Intelligence — Staged Document & Package Extraction Engine.
Implements:
1. True Procurement Package Ingestion (PDF, DOCX, XLSX, CSV, TXT, ZIP).
2. Safe ZIP Archive Extraction with Path-Traversal Rejection.
3. Deterministic Source Marker Preprocessing with real row coordinates.
4. Genuinely Staged Pipeline:
   - Stage A: Document Fact Extraction (requirements, dates, criteria, rules, clauses)
   - Stage B: Package Normalization (deduplication, source_refs aggregation)
   - Stage C: Cross-Document Reconciliation & Conflict Detection (6 conflict types)
   - Stage D: Executive Bid Brief Synthesis (generated from normalized model)
5. Strict Source Provenance Validation (rejects hallucinated pages, sheets, rows, or files).
"""
import io
import os
import re
import csv
import json
import zipfile
import xml.etree.ElementTree as ET
import anthropic
from config import get_anthropic_client
from requirement_semantics import (
    resolve_requirement_type,
    is_supplier_qualification,
    normalize_requirement_type,
    merge_requirement_types,
    resolve_candidate_requirement_types,
    SPECIFIC_REQUIREMENT_TYPES,
    TYPE_SUPPLIER_QUALIFICATION,
    ALLOWED_REQUIREMENT_TYPES,
)


# ── PROMPTS FOR STAGED EXTRACTION ─────────────────────────────────────────────

STAGE_A_FACT_EXTRACTION_PROMPT = """You are a senior procurement intelligence analyst.
Extract factual procurement data from this single procurement document.
The text contains deterministic source markers:
- [[SOURCE: <filename> | PAGE: <page_no>]]
- [[SOURCE: <filename> | SHEET: <sheet_name> | ROWS: <range>]]
- [[SOURCE: <filename> | SECTION: <heading>]]
- [[SOURCE: <filename> | ROWS: <range>]]

CRITICAL SOURCE TRACEABILITY RULES:
1. For every requirement or factual item, include "source_refs" pointing to the actual markers present in the text.
2. DO NOT hallucinate or guess page numbers, sheet names, or sections. Use only markers in the text.
3. "source_refs" schema:
   [
     {
       "source_doc": "<filename>",
       "page": <int or null>,
       "sheet": "<str or null>",
       "section": "<str or null>",
       "excerpt": "<1-2 sentence verbatim quote from text>"
     }
   ]
CRITICAL REQUIREMENT CLASSIFICATION RULES:
- "category" describes whether compliance is compulsory: "Mandatory|Rated|Financial|Supporting".
- "requirement_type" describes what KIND of requirement it is (controlled enum):
    "Supplier Qualification"   -> Pass/fail condition concerning whether the bidder itself is eligible, qualified, authorized, or capable of participating (e.g. legal eligibility, conditions of participation, financial standing, certifications, mandatory experience thresholds, OEM authorization as eligibility).
    "Technical Specification"   -> Property or characteristic of the offered solution, product, technology, or technical response.
    "Submission Compliance"    -> Forms, signatures, packaging, portal, document format, checklist, or bid-submission mechanics.
    "Delivery / SLA"           -> Implementation, service levels, response times, maintenance, support, uptime, or delivery performance.
    "Commercial / Contractual" -> Pricing mechanics, payment, insurance maintained during contract, liability, IP, termination, or contract terms.
    "Evaluation / Scored"      -> Competitive scored/rated criteria or scoring mechanics.
    "General Compliance"       -> Mandatory procurement obligations that do not safely fit one of the above.
- CRITICAL RULE: "Mandatory" describes whether compliance is compulsory. "Supplier Qualification" describes what KIND of requirement it is. They are independent concepts. Do NOT classify a requirement as "Supplier Qualification" merely because its category is "Mandatory"!
  Examples:
  - Mandatory + Supplier Qualification -> bidder must hold required registration to participate
  - Mandatory + Technical Specification -> proposed system must meet specified technical characteristic
  - Mandatory + Submission Compliance -> bidder must submit signed response form
  - Mandatory + Delivery / SLA -> supplier must respond to critical incidents within four hours
  - Mandatory + Commercial / Contractual -> supplier must accept stated liability provision
  - Rated + Evaluation / Scored -> experience response scored at 15%

CRITICAL EVALUATION HIERARCHY & WEIGHTING RULES:
- Preserve parent headings and child/subcriteria when stated or clearly structured by the source.
- Record "parent_stage" ONLY when the source structure or wording explicitly establishes that hierarchy (e.g. sub-table or indented/numbered section).
- Do NOT flatten nested evaluation tables.
- Preserve the raw weight string exactly as stated in the text (e.g. "40%", "25 points"). Bare numbers without units are NOT percentages.
- "weight_unit": "Percent" (e.g. 40%), "Points" (e.g. 25 pts), "Other", or "None".
- "weight_basis": record "Overall" if stated or structured as contributing to the entire procurement total; record "Within Parent" if stated as weighting within its parent criterion; record "Unknown" if ambiguous.
- "evaluation_role": controlled role: "Award Criterion" (scored criteria), "Subcriterion" (nested subcriteria), "Qualification / Gate" (mandatory conditions/gates), "Scoring Scale" (points scale 10/7/5/3/0), "Process / Methodology" (procedural stages), "Structural Container" (award criteria / scoring model parent heading), or "Unknown".
- Do NOT reinterpret a child percentage as an overall percentage unless the source makes that clear.
- Do NOT invent missing weights. If a weight is not stated, leave "weight" as null.
- Do NOT force evaluation totals to 100%. If the source states numbers that do not sum to 100%, record them faithfully.
- Do NOT combine points and percentages into one number.
- "threshold": minimum passing score or threshold (e.g. "70% minimum"). Thresholds are NOT weights.
- Include "source_refs" for each evaluation criterion.

CRITICAL SUBMISSION RULE ORTHOGONAL DIMENSIONS:
- "item": What is to be submitted or what the rule describes. Keep the official name (e.g. "Annex 3 (Supplier Response)", "Pricing Approach", "Signed Declaration").
- "artifact_type": Controlled category:
    "Proposal / Response"       -> Main technical proposal or response narrative document.
    "Form / Annex"              -> Standard forms, schedules, annexes, questionnaires, templates.
    "Pricing / Financial"       -> Costing workbooks, fee tables, pricing forms, rate cards.
    "Declaration / Certification" -> Signed declarations, compliance certifications, attestations.
    "Evidence / Attachment"     -> Audited financial accounts, resumes, case studies, sample files.
    "Process Instruction"       -> Pure delivery or process rules without an independent artifact (e.g. "Submit via portal", "Page limit 20 pages").
    "Other Artifact"            -> Other independent submission items.
    "Unknown"                   -> If unclassifiable.
- "file_format": Physical file format requested (e.g. "PDF", "DOCX", "XLSX", "Spreadsheet", "Online Form", "Hard Copy", "Multiple / Mixed", "Unspecified").
    CRITICAL: Portal and Email are DELIVERY CHANNELS, not file formats! Never write "Portal" or "Email" as the file_format. If no format is stated, use "Unspecified".
- "submission_channel": Delivery or transmission mechanism (e.g. "Portal", "Email", "E-procurement", "Courier", "Physical", "Upload", "Unspecified").
- "details": Packaging rules, page limits, signature requirements, or submission portal URLs.
- "mandatory": 1 if compulsory, 0 if optional/if-applicable, null if unspecified.
- CRITICAL PRINCIPLE: Artifact identity, file format, and submission channel are independent concepts.
  A required document submitted via a portal or email is STILL a required artifact (e.g. item="Annex 3 (Supplier Response)", file_format="Unspecified", submission_channel="Portal").
  Do NOT drop or merge away genuine submission artifacts merely because a portal or email delivery instruction is mentioned!
  Do NOT invent documents from pure process instructions (e.g. "Submit via portal" is a Process Instruction, not an artifact).

Return ONLY valid JSON with this exact schema:
{
  "doc_metadata": {
    "title": "Title of tender or null",
    "client": "Issuing organization or null",
    "file_number": "Solicitation number or null",
    "submission_deadline": "YYYY-MM-DD or null",
    "clarification_deadline": "YYYY-MM-DD or null",
    "notes": "Brief factual summary of document purpose"
  },
  "requirements": [
    {
      "req_id": "M1 or R1",
      "category": "Mandatory|Rated|Financial|Supporting",
      "requirement_type": "Supplier Qualification|Technical Specification|Submission Compliance|Delivery / SLA|Commercial / Contractual|Evaluation / Scored|General Compliance",
      "description": "Full requirement text",
      "rfso_ref": "Section reference",
      "weight": null,
      "evidence": "Required proof or submission item",
      "source_refs": []
    }
  ],
  "dates": [
    {"milestone": "Milestone name", "date": "YYYY-MM-DD", "source_doc": "<filename>"}
  ],
  "evaluation_criteria": [
    {
      "stage": "Evaluation stage or criterion title",
      "parent_stage": "Parent heading title or null",
      "hierarchy_level": 1,
      "evaluation_role": "Award Criterion|Subcriterion|Qualification / Gate|Scoring Scale|Process / Methodology|Structural Container|Unknown",
      "weight": "e.g. 75 points or 25% or null",
      "weight_unit": "Percent|Points|Other|None",
      "weight_basis": "Overall|Within Parent|Unknown",
      "threshold": "e.g. 70% or null",
      "notes": "Scoring rules or details",
      "source_refs": []
    }
  ],
  "submission_rules": [
    {
      "item": "Submission component name",
      "artifact_type": "Proposal / Response|Form / Annex|Pricing / Financial|Declaration / Certification|Evidence / Attachment|Process Instruction|Other Artifact|Unknown",
      "file_format": "PDF|DOCX|XLSX|Spreadsheet|Online Form|Hard Copy|Multiple / Mixed|Unspecified",
      "submission_channel": "Portal|Email|E-procurement|Courier|Physical|Upload|Unspecified",
      "format": "Physical format or delivery note (backward compatible string)",
      "details": "Packaging / page limit / portal instruction",
      "mandatory": 1,
      "source_refs": []
    }
  ],
  "deliverables": [
    {"title": "Deliverable output", "description": "What must be produced", "category": "Core|Optional"}
  ],
  "commercial_clauses": [
    {"topic": "Panel Maximums|Rate Caps|Extension Options|Liability", "details": "Commercial mechanism details"}
  ],
  "contract_risks": [
    {"risk": "AI Restrictions|IP Ownership|Liability|Subcontractors", "severity": "High|Medium|Low", "details": "Contractual risk details"}
  ],
  "typed_observations": [
    {
      "family": "IDENTITY|MILESTONE|CONTRACT_TERM|MONETARY|PROCUREMENT_MECHANIC|DOCUMENT_ROLE",
      "semantic_kind": "Controlled kind from the instructions below",
      "original_value": "Exact source value or wording",
      "source_doc": "<filename>",
      "source_refs": [],
      "scope": {"component": null, "lot": null, "category": null},
      "date": null, "time": null, "timezone": null, "precision": null,
      "amount": null, "currency": null, "tax_basis": null,
      "guarantee_status": null, "period_basis": null,
      "duration": null, "unit": null, "option_count": null, "optional": null,
      "document_role": null, "document_role_basis": null,
      "supersession": null
    }
  ]
}

TYPED OBSERVATION RULES:
- Emit one atomic observation for each source occurrence in these families. Do not repeat
  the same fact in multiple typed observations and do not choose package-wide winners.
- Identity kinds: OPPORTUNITY_TITLE, DOCUMENT_TITLE, BUYER_NAME, ISSUING_AUTHORITY,
  SOLICITATION_NUMBER, DOCUMENT_REFERENCE_NUMBER.
- Milestone kinds: SUBMISSION_DEADLINE, CLARIFICATION_DEADLINE, INTENT_TO_BID_DEADLINE,
  SITE_VISIT, BRIEFING, PRESENTATION_OR_DEMO, AWARD_DATE, CONTRACT_START, CONTRACT_END,
  IMPLEMENTATION_DEADLINE, DELIVERY_DEADLINE, BID_VALIDITY_END, PUBLICATION_DATE,
  AMENDMENT_DATE, PAYMENT_MILESTONE, OTHER, UNKNOWN.
- Contract-term kinds: INITIAL_DURATION, COMMENCEMENT_DATE, END_DATE, EXTENSION_OPTION,
  RENEWAL_OPTION, MAXIMUM_TERM, TERMINATION_CONDITION, TERM_STATEMENT.
- Monetary kinds include ESTIMATED_CONTRACT_VALUE, MAXIMUM_CONTRACT_VALUE,
  FRAMEWORK_CEILING, BUDGET, ANNUAL_VALUE, LOT_VALUE, MINIMUM_SPEND, GUARANTEED_SPEND,
  FORECAST_VALUE, EVALUATION_SCENARIO_VALUE, RATE_CAP, UNIT_RATE, MILESTONE_PAYMENT.
- Procurement mechanics are atomic: RFP, ITT, RFQ, SINGLE_SUPPLIER_AWARD,
  MULTIPLE_SUPPLIER_AWARD, STANDING_OFFER, FRAMEWORK, PANEL, CALL_OFF, LOTS,
  FIXED_PRICE, RATE_CARD, HOURLY_RATE, DAILY_RATE, MILESTONE_PAYMENT, SUBSCRIPTION,
  RETAINER, RATE_CAP, ESCALATION_CAP, NO_GUARANTEED_VOLUME, EXTENSION_PRICING,
  EVALUATED_SCENARIO.
- Do not infer currency, authority, amendment precedence, optionality, time, or timezone.
- Every typed observation needs an exact excerpt and only source coordinates present in
  the input markers. Document role needs cited structural wording; filename alone is not proof.
- Use supersession only for an explicit old-to-new statement and quote that statement.
"""

LEGACY_STAGE_D_SYNTHESIS_PROMPT = """You are an executive bid director synthesizing a Bid Brief from normalized procurement facts.
You MUST synthesize ONLY from the supplied normalized facts and detected cross-document conflicts.
Do NOT invent facts not present in the normalized data model.

CONTRACT WITH THE CONTEXT:
- The supplied context is the authoritative normalized procurement model. All requirements, dates, evaluation
  criteria, submission rules, and conflicts have been deterministically extracted and verified.
- Mandatory requirements are compulsory procurement obligations.
- Mandatory does NOT automatically mean supplier qualification.
- qualification_gates contains ONLY true pass/fail bidder eligibility or qualification conditions (category == Mandatory AND requirement_type == Supplier Qualification).
- Technical specifications, submission instructions, delivery/SLA obligations, commercial terms, and general mandatory compliance requirements MUST NOT be represented as qualification gates merely because they are Mandatory.
- Rated/scored requirements MUST NOT become qualification gates.
- All requirements remain available in the complete normalized context and requirements register.
- Distinguish procurement requirements (what the buyer demands) from bidder capability (what the bidder holds).
  Do not convert UNKNOWN bidder capability into PASS.
- Detected conflicts are unresolved unless the source model explicitly resolves them. Do not silently merge
  or dismiss TRUE_CONFLICT or REVIEW_ITEM records.
- Evaluation Hierarchy: preserve parent-child relationships and weight basis. Do NOT treat parent and child weights as independent overall weights. Do NOT normalize totals to 100% if source totals differ. Preserve source discrepancies and mixed units. Distinguish thresholds from weights.
- Do not invent submission documents, certifications, languages, security clearances, insurance coverage,
  or pricing facts unless they are explicitly present in the supplied normalized facts.
- Do not infer organizational capacity or qualifications from blank or internal fields.

Return ONLY valid JSON with this exact schema:
{
  "bid": {
    "title": "Concise opportunity title",
    "client": "Issuing buyer organization name",
    "file_number": "Solicitation number or null",
    "owner": null,
    "sensitivity": "Standard",
    "submission_deadline": "YYYY-MM-DD or null",
    "clarification_deadline": "YYYY-MM-DD or null",
    "value_cad": null,
    "notes": "2-3 sentence executive synthesis of what is being procured"
  },
  "brief": {
    "executive_summary": "Plain-language executive summary answering: What is the buyer actually procuring?",
    "opportunity_type": "Services RFP|Standing Offer|Panel Agreement|Software/Systems|Advisory",
    "contract_term": "e.g. 3 years with 2 optional 1-year extensions, or Not stated",
    "procurement_model": "Single Contract|Standing Offer Panel|Multi-vendor Call-off",
    "scope_categories": ["Category 1", "Category 2"],
    "deliverables_summary": [
      {"title": "Deliverable title", "description": "Scope details", "category": "Core|Optional"}
    ],
    "qualification_gates": [
      {"requirement": "Pass/fail bidder eligibility condition", "type": "Supplier Qualification", "rfp_ref": "Ref", "disqualification_risk": "High"}
    ],
    "evaluation_breakdown": [
      {"stage": "Technical / Price stage", "parent_stage": null, "weight": "75 points / 25%", "weight_basis": "Overall", "threshold": "Threshold or null", "notes": "Scoring rules"}
    ],
    "commercial_structure": [
      {"topic": "Commercial topic", "details": "Details"}
    ],
    "contract_risks": [
      {"risk": "Risk title", "severity": "High|Medium|Low", "details": "Details"}
    ],
    "submission_requirements": [
      {"item": "Document name", "format": "Format", "details": "Details"}
    ],
    "key_dates": [
      {"milestone": "Milestone name", "date": "YYYY-MM-DD"}
    ],
    "source_citations": {
      "mandatory_ref": "Ref string",
      "evaluation_ref": "Ref string",
      "sow_ref": "Ref string"
    }
  },
  "outline": [
    {"sort_order": 0, "section_num": "1", "title": "Proposal Section Title", "owner": null, "word_limit": null, "status": "Not Started", "notes": "Guidance"}
  ]
}
"""




from stage_d_projection import SYNTHESIS_PROMPT as STAGE_D_SYNTHESIS_PROMPT


# ── DOCUMENT TEXT EXTRACTORS & DETERMINISTIC MARKERS ─────────────────────────

def extract_pdf_with_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Extract PDF text with per-page [[SOURCE: filename | PAGE: N]] markers and metadata."""
    meta = {"page_count": 0, "pages": {}}
    
    # 1. Try PyMuPDF (fitz)
    try:
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages = []
        meta["page_count"] = len(doc)
        for i, page in enumerate(doc, 1):
            t = page.get_text().strip()
            if t:
                meta["pages"][i] = t
                pages.append(f"[[SOURCE: {filename} | PAGE: {i}]]\n{t}")
        if pages:
            return "\n\n".join(pages), meta
    except Exception:
        pass

    # 2. Try pypdf fallback
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages = []
        meta["page_count"] = len(reader.pages)
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                meta["pages"][i] = t
                pages.append(f"[[SOURCE: {filename} | PAGE: {i}]]\n{t}")
        if pages:
            return "\n\n".join(pages), meta
    except Exception:
        pass

    raw_text = file_bytes.decode("utf-8", errors="ignore")
    meta["page_count"] = 1
    meta["pages"][1] = raw_text
    return f"[[SOURCE: {filename} | PAGE: 1]]\n" + raw_text, meta


def extract_docx_with_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Extract DOCX text with heading / section markers and metadata."""
    meta = {"sections": [], "tables_count": 0}
    
    # 1. Try python-docx
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        current_heading = "Document Body"
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            if p.style.name.startswith("Heading"):
                current_heading = text
                meta["sections"].append(current_heading)
                paragraphs.append(f"\n[[SOURCE: {filename} | SECTION: {current_heading}]]")
            else:
                paragraphs.append(text)
        
        for t in doc.tables:
            meta["tables_count"] += 1
            for row in t.rows:
                row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_vals:
                    paragraphs.append(f"[[SOURCE: {filename} | TABLE]] " + " | ".join(row_vals))

        if paragraphs:
            return f"[[SOURCE: {filename} | SECTION: Header]]\n" + "\n".join(paragraphs), meta
    except Exception:
        pass

    # 2. Native XML zipfile fallback for DOCX
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            if "word/document.xml" in z.namelist():
                xml_data = z.read("word/document.xml")
                root = ET.fromstring(xml_data)
                texts = []
                for node in root.iter():
                    if node.tag.endswith("t") and node.text:
                        texts.append(node.text)
                if texts:
                    meta["sections"].append("Document Body")
                    return f"[[SOURCE: {filename} | SECTION: Document Body]]\n" + " ".join(texts), meta
    except Exception:
        pass

    return f"[[SOURCE: {filename}]]\n[DOCX parsing failed: unreadable content]", meta


def extract_xlsx_with_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Extract XLSX text preserving sheet names and real worksheet row coordinates (even with blank rows)."""
    meta = {"sheets": [], "rows_per_sheet": {}}
    
    # 1. Try openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheets_text = []
        meta["sheets"] = list(wb.sheetnames)
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            rows_text = []
            non_empty_indices = []
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                clean_row = [str(val).strip() for val in row if val is not None and str(val).strip()]
                if clean_row:
                    non_empty_indices.append(r_idx)
                    rows_text.append(f"Row {r_idx}: " + " | ".join(clean_row))
            
            if rows_text:
                min_r = non_empty_indices[0]
                max_r = non_empty_indices[-1]
                meta["rows_per_sheet"][sheetname] = (min_r, max_r, set(non_empty_indices))
                chunk = f"[[SOURCE: {filename} | SHEET: {sheetname} | ROWS: {min_r}-{max_r}]]\n" + "\n".join(rows_text)
                sheets_text.append(chunk)
            else:
                meta["rows_per_sheet"][sheetname] = (0, 0, set())
        
        if sheets_text:
            return "\n\n".join(sheets_text), meta
    except Exception:
        pass

    # 2. Native XML zipfile fallback for XLSX sharedStrings & sheets
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
            shared_strings = []
            if "xl/sharedStrings.xml" in z.namelist():
                sst_root = ET.fromstring(z.read("xl/sharedStrings.xml"))
                for si in sst_root.iter():
                    if si.tag.endswith("t") and si.text:
                        shared_strings.append(si.text.strip())
            
            if shared_strings:
                meta["sheets"] = ["Workbook Data"]
                meta["rows_per_sheet"]["Workbook Data"] = (1, len(shared_strings), set(range(1, len(shared_strings)+1)))
                return f"[[SOURCE: {filename} | SHEET: Workbook Data | ROWS: 1-{len(shared_strings)}]]\n" + "\n".join(shared_strings), meta
    except Exception:
        pass

    return f"[[SOURCE: {filename}]]\n[XLSX parsing failed: unreadable spreadsheet]", meta


def extract_csv_with_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Extract CSV text with real row coordinate markers."""
    meta = {"rows_count": 0, "non_empty_rows": set()}
    try:
        text_content = file_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text_content))
        rows_text = []
        for r_idx, row in enumerate(reader, 1):
            clean_row = [str(val).strip() for val in row if val is not None and str(val).strip()]
            if clean_row:
                meta["non_empty_rows"].add(r_idx)
                rows_text.append(f"Row {r_idx}: " + " | ".join(clean_row))
        
        meta["rows_count"] = len(rows_text)
        if rows_text:
            min_r = min(meta["non_empty_rows"])
            max_r = max(meta["non_empty_rows"])
            return f"[[SOURCE: {filename} | ROWS: {min_r}-{max_r}]]\n" + "\n".join(rows_text), meta
    except Exception:
        pass

    return f"[[SOURCE: {filename}]]\n" + file_bytes.decode("utf-8", errors="ignore"), meta


def extract_document_with_metadata(file_bytes: bytes, filename: str) -> tuple[str, dict]:
    """Extract text with deterministic markers and return rich parsing metadata."""
    name_lower = filename.lower()
    if name_lower.endswith(".pdf"):
        return extract_pdf_with_metadata(file_bytes, filename)
    elif name_lower.endswith(".docx"):
        return extract_docx_with_metadata(file_bytes, filename)
    elif name_lower.endswith(".xlsx"):
        return extract_xlsx_with_metadata(file_bytes, filename)
    elif name_lower.endswith(".csv"):
        return extract_csv_with_metadata(file_bytes, filename)
    elif name_lower.endswith(".txt") or name_lower.endswith(".md"):
        txt = file_bytes.decode("utf-8", errors="ignore")
        return f"[[SOURCE: {filename}]]\n" + txt, {"type": "text", "len": len(txt)}
    else:
        # Unsupported format (e.g. .doc or .xls)
        return f"[[SOURCE: {filename}]]\n[Unsupported file format. Please upload PDF, DOCX, XLSX, CSV, or TXT]", {"unsupported": True}


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Standard multi-format document parser returning marked text."""
    text, _ = extract_document_with_metadata(file_bytes, filename)
    return text


# ── PACKAGE UNPACKING WITH SECURITY CONTROLS ──────────────────────────────────

def unpack_procurement_package(raw_files: list[tuple[str, bytes]]) -> tuple[list[tuple[str, bytes]], list[str]]:
    """
    Unpack uploaded files, safely extracting ZIP archives.
    Rejects path traversal (e.g. ../ or absolute paths), filters unsupported binaries (.exe, .bin)
    and legacy unsupported office formats (.doc, .xls), returning (supported_files, warnings).
    """
    supported_extensions = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}
    unsupported_legacy = {".doc", ".xls"}
    unpacked = []
    warnings = []

    for fname, fbytes in raw_files:
        fname_clean = os.path.basename(fname.strip())
        lower_name = fname_clean.lower()

        if lower_name.endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(fbytes)) as z:
                    for entry in z.infolist():
                        if entry.is_dir():
                            continue
                        
                        entry_name = entry.filename.replace("\\", "/")
                        
                        # Security Check: Reject path traversal
                        if entry_name.startswith("/") or ".." in entry_name.split("/"):
                            warnings.append(f"Security Alert: Skipped unsafe file path in archive: {entry_name}")
                            continue

                        # Skip hidden OS files
                        base_entry = os.path.basename(entry_name)
                        if base_entry.startswith(".") or entry_name.startswith("__MACOSX"):
                            continue

                        ext = os.path.splitext(base_entry)[1].lower()
                        if ext in supported_extensions:
                            unpacked.append((base_entry, z.read(entry.filename)))
                        elif ext in unsupported_legacy:
                            warnings.append(f"Skipped legacy format '{base_entry}'. Convert to modern .docx / .xlsx format.")
                        else:
                            warnings.append(f"Ignored unsupported file in archive: {base_entry}")
            except Exception as e:
                warnings.append(f"Failed to extract ZIP archive {fname_clean}: {e}")
        else:
            ext = os.path.splitext(fname_clean)[1].lower()
            if ext in supported_extensions:
                unpacked.append((fname_clean, fbytes))
            elif ext in unsupported_legacy:
                warnings.append(f"Skipped legacy format '{fname_clean}'. Convert to modern .docx / .xlsx format.")
            else:
                warnings.append(f"Ignored unsupported file: {fname_clean}")

    return unpacked, warnings


# ── SOURCE PROVENANCE VALIDATION ENGINE ───────────────────────────────────────

def validate_source_refs(source_refs: list, package_metadata: dict) -> list[dict]:
    """
    Strict validation of AI-returned source references against physical document parse metadata.
    Validates:
    1. Source document existence in package.
    2. Page bounds (1 <= page <= max_pages).
    3. Sheet name existence in workbook.
    4. Sheet row ranges against actual non-empty rows.
    5. Excerpt presence/plausibility.
    
    Returns list of verified/annotated source references. Invalid locations are rejected or flagged.
    """
    if not isinstance(source_refs, list):
        return []

    valid_files = set(package_metadata.get("files", []))
    valid_files_lower = {f.lower(): f for f in valid_files}
    doc_meta_map = package_metadata.get("doc_metadata", {})
    doc_texts = package_metadata.get("doc_texts", {})

    verified_refs = []
    for ref in source_refs:
        if not isinstance(ref, dict):
            continue

        raw_doc = str(ref.get("source_doc") or "").strip()
        matched_doc = valid_files_lower.get(raw_doc.lower())

        if not matched_doc:
            # Rejection 1: Nonexistent document in package
            verified_refs.append({
                "source_doc": raw_doc or "Unknown",
                "page": None,
                "sheet": None,
                "section": None,
                "excerpt": ref.get("excerpt", ""),
                "verified": False,
                "validation_error": f"Document '{raw_doc}' not found in procurement package"
            })
            continue

        doc_meta = doc_meta_map.get(matched_doc, {})
        page = ref.get("page")
        sheet = ref.get("sheet")
        section = ref.get("section")
        excerpt = ref.get("excerpt", "")

        is_valid = True
        error_reasons = []

        # Check page bounds
        if page is not None:
            try:
                page_int = int(page)
                max_pages = doc_meta.get("page_count", 0)
                if max_pages > 0 and (page_int < 1 or page_int > max_pages):
                    is_valid = False
                    error_reasons.append(f"Page {page_int} out of bounds (document has {max_pages} pages)")
                    page = None
            except (ValueError, TypeError):
                is_valid = False
                error_reasons.append(f"Invalid page value '{page}'")
                page = None

        # Check sheet name
        if sheet is not None:
            sheet_str = str(sheet).strip()
            valid_sheets = doc_meta.get("sheets", [])
            if valid_sheets and sheet_str not in valid_sheets:
                # Check case-insensitive match
                sheet_match = next((s for s in valid_sheets if s.lower() == sheet_str.lower()), None)
                if sheet_match:
                    sheet = sheet_match
                else:
                    is_valid = False
                    error_reasons.append(f"Sheet '{sheet_str}' does not exist in workbook")
                    sheet = None

        # Check rows in sheet
        if sheet and "rows_per_sheet" in doc_meta:
            row_info = doc_meta["rows_per_sheet"].get(sheet)
            if row_info:
                min_r, max_r, nonempty_rows = row_info
                # If ref has row range or rows
                rows_spec = ref.get("rows")
                if rows_spec:
                    try:
                        if isinstance(rows_spec, int) and rows_spec not in nonempty_rows:
                            is_valid = False
                            error_reasons.append(f"Row {rows_spec} is blank or outside data range")
                    except Exception:
                        pass

        # Check excerpt plausibility
        if excerpt and len(excerpt) > 15:
            full_doc_text = doc_texts.get(matched_doc, "")
            # Check if excerpt words appear in doc
            sample_words = [w for w in re.findall(r'\w+', excerpt.lower()) if len(w) > 3][:6]
            if sample_words and not any(w in full_doc_text.lower() for w in sample_words):
                is_valid = False
                error_reasons.append("Excerpt text not found in source document")

        clean_ref = {
            "source_doc": matched_doc,
            "page": page,
            "sheet": sheet,
            "section": section,
            "excerpt": excerpt,
            "verified": is_valid
        }
        if not is_valid:
            clean_ref["validation_error"] = "; ".join(error_reasons)

        verified_refs.append(clean_ref)

    return verified_refs


# ── DETERMINISTIC CONFLICT & AMBIGUITY DETECTION ENGINE ───────────────────────

DATE_EVENT_PATTERNS = [
    # 1. Question / Clarification deadlines (must evaluate before general deadline keywords)
    ("QUESTION_DEADLINE", [
        r"question", r"clarification", r"enquir", r"inquir", r"rfi", r"q&a", r"q\s*&\s*a",
        r"queries", r"questions\s+due", r"last\s+date\s+for\s+questions", r"deadline\s+for\s+questions"
    ]),
    # 2. Intent to Bid / Confirmation
    ("INTENT_TO_BID_DATE", [
        r"intent\s+to\s+bid", r"bid\s+intent", r"intention\s+to\s+bid", r"confirmation\s+of\s+intent",
        r"intent\s+deadline", r"intent\s+date", r"opt-in", r"acknowledgement\s+deadline"
    ]),
    # 3. Publication / Issue Dates
    ("PUBLICATION_DATE", [
        r"solicitation\s+publication", r"publication\s+date", r"date\s+issued", r"posting\s+date",
        r"date\s+of\s+issue", r"published\s+date", r"rfp\s+issue"
    ]),
    # 4. Amendment / Addenda Dates
    ("AMENDMENT_DATE", [
        r"amendment\s+no", r"amendment\s+published", r"amendment\s+issued", r"addendum\s+published",
        r"addendum\s+issued", r"bulletin\s+issued", r"bulletin\s+date"
    ]),
    # 5. Experience timeframe / Reference period (not a procurement event)
    ("EXPERIENCE_TIMEFRAME", [
        r"recent\s+engagements", r"experience\s+timeframe", r"reference\s+period", r"project\s+period",
        r"prior\s+experience"
    ]),
    # 6. Site visits / Conferences / Walkthroughs
    ("SITE_VISIT_DATE", [
        r"site\s+visit", r"site\s+walkthrough", r"bidders\s+conference", r"information\s+session",
        r"proponents\s+conference", r"mandatory\s+meeting", r"optional\s+meeting", r"walkthrough",
        r"site\s+inspection", r"site\s+meeting"
    ]),
    # 7. Presentations / Demos / Interviews
    ("PRESENTATION_DATE", [
        r"presentation", r"interview", r"demonstration", r"demo", r"shortlist", r"oral",
        r"oral\s+presentation"
    ]),
    # 8. Award / Selection
    ("AWARD_DATE", [
        r"award", r"selection", r"intent\s+to\s+award", r"notification\s+of\s+award",
        r"anticipated\s+award", r"target\s+award"
    ]),
    # 9. Contract Start
    ("CONTRACT_START", [
        r"contract\s+start", r"commencement", r"start\s+date", r"service\s+start",
        r"effective\s+date", r"work\s+start"
    ]),
    # 10. Contract End / Expiry
    ("CONTRACT_END", [
        r"contract\s+end", r"completion", r"expiry", r"expiration", r"end\s+date",
        r"termination\s+date", r"work\s+end"
    ]),
    # 11. Proposal Validity
    ("VALIDITY_DATE", [
        r"validity", r"valid\s+until", r"proposal\s+valid", r"bid\s+valid",
        r"irrevocable\s+for"
    ]),
    # 12. Submission / Closing Deadlines (Standard Proposal Closing Event)
    ("SUBMISSION_DEADLINE", [
        r"submission", r"closing", r"close", r"due\s+date", r"rfp\s+due", r"tender\s+close",
        r"bid\s+close", r"deadline", r"closing\s+date", r"closing\s+time", r"submission\s+date"
    ]),
]


def classify_date_milestone(milestone_str: str) -> str:
    """Classify procurement milestone string into semantic event category."""
    m = (milestone_str or "").lower().strip()
    for event_type, patterns in DATE_EVENT_PATTERNS:
        if any(re.search(pat, m) for pat in patterns):
            return event_type
    return "OTHER_DATE"


SUBMISSION_DIMENSION_PATTERNS = [
    ("ENVELOPE_STRUCTURE", [
        r"separate\s+envelope", r"two-envelope", r"two\s+envelope", r"separate\s+financial",
        r"separate\s+technical", r"combined\s+proposal", r"single\s+package", r"combined\s+envelope",
        r"single\s+pdf", r"single\s+combined", r"combined\s+technical\s+and\s+financial",
        r"separate\s+files\s+per\s+envelope", r"envelope\s+1", r"envelope\s+2"
    ]),
    ("PAGE_LIMIT", [
        r"page\s+limit", r"max\s+pages", r"maximum\s+pages", r"pages\s+maximum", r"page\s+count",
        r"pages\s*\(\s*excluding", r"page\s+size",
        r"\b(?:exceed|maximum|limit\s+of|max)\s*(?:of\s*)?\d+\s*pages?\b",
        r"\b\d+\s*pages?\s*(?:max|limit)\b"
    ]),
    ("PORTAL_REQUIREMENT", [
        r"supplier\s+registration", r"portal\s+account", r"digital\s+key", r"eprocurement\s+setup",
        r"merx\s+registration", r"registration\s+required", r"maintain\s+a\s+(?:merx|portal|buyandsell)\s+account",
        r"(?:merx|portal|buyandsell)\s+account", r"vendor\s+registration", r"account\s+registration"
    ]),
    ("SUBMISSION_CHANNEL", [
        r"upload\s+bid", r"upload\s+proposal", r"submit\s+(?:bid|proposal)\s+(?:through|via|on)",
        r"electronic\s+bid\s+submission", r"electronic\s+submission",
        r"portal\s+upload", r"email\s+submission", r"email\s+only", r"\bemail\b", r"courier",
        r"hardcopy", r"hard\s+copy", r"physical\s+delivery", r"in\s+person", r"mail\s+submission",
        r"electronic\s+tender", r"\bmerx\b", r"\bbuyandsell\b"
    ]),
    ("SIGNATURE_REQUIREMENT", [
        r"authorized\s+signature", r"docusign", r"digital\s+signature", r"signed\s+form",
        r"duly\s+executed", r"signing\s+officer"
    ]),
    ("FILE_FORMAT", [
        r"excel\s+spreadsheet", r"excel\s+workbook", r"xlsx", r"docx", r"searchable\s+pdf",
        r"pdf\s+format", r"word\s+format", r"spreadsheet", r"excel", r"\bpdf\b"
    ]),
    ("DOCUMENT_REQUIREMENT", [
        r"appendix", r"submission\s+form", r"questionnaire", r"rate\s+card", r"pricing\s+form",
        r"certificate\s+of\s+insurance"
    ]),
]


def classify_submission_rule_dimension(item_str: str, format_str: str = "", details_str: str = "") -> str:
    """Classify a submission rule into its operational dimension with strict precedence."""
    combined = f"{item_str} {format_str} {details_str}".lower()
    for dimension, patterns in SUBMISSION_DIMENSION_PATTERNS:
        if any(re.search(pat, combined) for pat in patterns):
            return dimension
    return "OTHER_SUBMISSION_RULE"


INSURANCE_CLASS_PATTERNS = [
    ("COMMERCIAL_GENERAL_LIABILITY", [
        r"commercial\s+general\s+liability", r"\bcgl\b", r"general\s+liability", r"public\s+liability",
        r"comprehensive\s+general\s+liability"
    ]),
    ("PROFESSIONAL_LIABILITY", [
        r"professional\s+liability", r"errors\s+and\s+omissions", r"errors\s*&\s*omissions",
        r"\be&o\b", r"\be\s*and\s*o\b", r"professional\s+indemnity"
    ]),
    ("CYBER_LIABILITY", [
        r"cyber", r"data\s+breach", r"network\s+security", r"technology\s+errors"
    ]),
    ("AUTOMOBILE_LIABILITY", [
        r"automobile", r"motor\s+vehicle", r"fleet\s+liability", r"auto\s+liability"
    ]),
    ("WORKERS_COMPENSATION", [
        r"workers\s+compensation", r"wsib", r"employer\s+liability", r"worker'?s\s+comp"
    ]),
]


def classify_insurance_class(text: str) -> str:
    """Classify insurance policy requirement into standardized coverage class."""
    t = (text or "").lower()
    for ins_class, patterns in INSURANCE_CLASS_PATTERNS:
        if any(re.search(pat, t) for pat in patterns):
            return ins_class
    return "OTHER_INSURANCE"


def classify_commercial_topic(topic_str: str, details_str: str = "") -> str:
    """Classify commercial clause into standardized topic."""
    combined = f"{topic_str} {details_str}".lower()
    if any(k in combined for k in ["insurance", "liability", "indemnity", "cgl", "e&o", "wsib"]):
        return "INSURANCE_REQUIREMENT"
    if any(k in combined for k in ["panel vendor", "maximum vendors", "maximum number of suppliers", "panel size", "number of standing offers", "vendor cap", "maximum awards", "number of awards", "maximum panel", "panel maximum", "panel cap", "panel"]):
        return "PANEL_VENDOR_CAP"
    if any(k in combined for k in ["rate cap", "maximum rate", "daily rate cap", "hourly rate cap", "maximum per diem", "rate ceiling", "per diem cap"]):
        return "RATE_CAP"
    if any(k in combined for k in ["escalation", "annual rate increase", "price increase cap", "inflation adjustment cap", "annual escalation"]):
        return "ANNUAL_ESCALATION_CAP"
    if any(k in combined for k in ["contract value cap", "maximum contract value", "total expenditure cap", "funding ceiling", "overall contract maximum"]):
        return "CONTRACT_VALUE_CAP"
    return "OTHER_COMMERCIAL_TERM"


def extract_commercial_limit(topic_type: str, text: str) -> float | None:
    """Extract normalized numeric limit for a commercial topic."""
    t = (text or "").lower()
    if topic_type == "PANEL_VENDOR_CAP":
        m = re.search(r'\b(?:up\s+to|maximum\s+of|max\s+of|limit\s+of|maximum|max)?\s*(\d+)\s*(?:vendors?|suppliers?|proponents?|standing\s+offers?|awards?|firms?)\b', t)
        if m:
            return float(m.group(1))
        m_num = re.search(r'(?:panel\s+vendors?|maximum\s+panel|panel\s+size|panel\s+maximums?)\s*(?:=|:|\sis\s)?\s*(\d+)\b', t)
        if m_num:
            return float(m_num.group(1))
        m_any = re.search(r'(\d+)\s*(?:vendors?|suppliers?|proponents?|standing\s+offers?|awards?|firms?)\b', t)
        if m_any:
            return float(m_any.group(1))
    elif topic_type == "ANNUAL_ESCALATION_CAP":
        m = re.search(r'(\d+(?:\.\d+)?)\s*%', t)
        if m:
            return float(m.group(1))
    elif topic_type in ["RATE_CAP", "CONTRACT_VALUE_CAP"]:
        return extract_monetary_amount(text)
    return None


def extract_monetary_amount(text: str) -> float | None:
    """
    Extract normalized monetary amount (in CAD/currency units) from text.
    Protects against interpreting unrelated years or numeric identifiers (e.g. 2026) as amounts.
    """
    if not text or not isinstance(text, str):
        return None
    t = text.lower()

    # 1. Millions: $2M, 2m, 2.5m, 2 million, $2 million, 2M coverage
    m_mil = re.search(r"(?:cad|usd|c\$|\$)?\s*(\d+(?:\.\d+)?)\s*(?:m\b|million)", t)
    if m_mil:
        return float(m_mil.group(1)) * 1_000_000.0

    # 2. Thousands: $500K, 500k, 500 thousand
    m_k = re.search(r"(?:cad|usd|c\$|\$)?\s*(\d+(?:\.\d+)?)\s*(?:k\b|thousand)", t)
    if m_k:
        return float(m_k.group(1)) * 1_000.0

    # 3. Explicit currency prefix + comma/number: $2,000,000 or CAD 2,000,000 or $5000000
    m_curr = re.search(r"(?:cad|usd|c\$|\$)\s*(\d{1,3}(?:,\d{3})+|\d{4,})\b", t)
    if m_curr:
        return float(m_curr.group(1).replace(",", ""))

    # 4. Comma-formatted with explicit insurance/limit context (e.g. "limit of 2,000,000" or "coverage 2,000,000")
    if any(k in t for k in ["limit", "coverage", "policy", "insurance", "liability", "aggregate", "occurrence"]):
        m_comma = re.search(r"\b(\d{1,3}(?:,\d{3})+)\b", t)
        if m_comma:
            return float(m_comma.group(1).replace(",", ""))

    # Plain integers without currency symbols or explicit million/thousand markers are NOT treated as monetary limits (protects against years like 2026)
    return None


def select_opposing_pair(
    records: list,
    value_fn,
    source_fn=None,
    require_different_sources: bool = False
) -> tuple | None:
    """
    Find two records record_a and record_b such that value_fn(record_a) != value_fn(record_b).
    If require_different_sources=True and source_fn is provided, also enforces source_fn(record_a) != source_fn(record_b).
    Returns (record_a, record_b) or None if no such pair exists.
    """
    if not records or len(records) < 2:
        return None

    valid_records = [r for r in records if value_fn(r) is not None and str(value_fn(r)).strip() != ""]
    if len(valid_records) < 2:
        return None

    if require_different_sources and source_fn:
        for i, r_a in enumerate(valid_records):
            val_a = value_fn(r_a)
            src_a = source_fn(r_a)
            for r_b in valid_records[i + 1:]:
                val_b = value_fn(r_b)
                src_b = source_fn(r_b)
                if val_a != val_b and src_a != src_b:
                    return r_a, r_b
        return None
    else:
        val_map = {}
        for r in valid_records:
            val_map.setdefault(value_fn(r), []).append(r)
        unique_vals = list(val_map.keys())
        if len(unique_vals) < 2:
            return None
        return val_map[unique_vals[0]][0], val_map[unique_vals[1]][0]


def classify_security_clearance(text: str) -> str | None:
    """
    Classify security clearance level with strict priority:
    1. TOP_SECRET (checked first)
    2. SECRET (checked second)
    3. RELIABILITY (checked third)
    """
    if not text or not isinstance(text, str):
        return None
    t = text.lower()
    if re.search(r"top\s+secret", t):
        return "TOP_SECRET"
    if re.search(r"\bsecret\b", t):
        return "SECRET"
    if re.search(r"reliability", t):
        return "RELIABILITY"
    return None


def _extract_eval_criterion_identity(ec: dict) -> str | None:
    """
    Extract normalized evaluation criterion identity, prioritizing explicit IDs before broad keywords.
    Extracts explicit IDs from stage / criterion / title / notes / criterion_id.
    """
    if not isinstance(ec, dict):
        return None
    stage = (ec.get("stage") or "").strip()
    title = (ec.get("criterion") or ec.get("title") or ec.get("name") or ec.get("item") or "").strip()
    desc = (ec.get("description") or ec.get("notes") or "").strip()
    cid = (ec.get("criterion_id") or ec.get("id") or "").strip()

    combined = f"{cid} {stage} {title} {desc}".lower()

    # 1. Explicit criterion identifier (e.g. R1, R2, CR1, CR2, TC1, TC2, RT1, PR1)
    # Check cid field first
    if cid and re.match(r'^(?:R|CR|TC|RT|PR|M)\d+$', cid, re.IGNORECASE):
        return f"CRITERION_{cid.upper()}"

    # Search for explicit criterion marker at word boundary in cid, stage, or title
    id_match = re.search(r'\b(R\d+|CR\d+|TC\d+|RT\d+|PR\d+)\b', f"{cid} {stage} {title}", re.IGNORECASE)
    if id_match:
        return f"CRITERION_{id_match.group(1).upper()}"

    # 2. Overall technical vs financial ratio / score
    if any(k in combined for k in ["overall technical weight", "overall ratio", "technical / financial", "tech / fin", "technical ratio", "weighting ratio", "70/30", "80/20", "75/25", "technical score", "technical weight", "rated score", "technical points"]):
        return "OVERALL_TECHNICAL_WEIGHT"
    if "overall" in stage.lower() and any(k in stage.lower() for k in ["tech", "rated"]):
        return "OVERALL_TECHNICAL_WEIGHT"
    if "overall" in stage.lower() and any(k in stage.lower() for k in ["finan", "price", "cost", "commercial"]):
        return "OVERALL_FINANCIAL_WEIGHT"
    if "overall technical" in title.lower() or title.lower() in ["overall technical weight", "technical score", "technical weight"]:
        return "OVERALL_TECHNICAL_WEIGHT"

    # 3. Specific semantic criteria identity (only when no explicit identifier exists)
    if any(k in combined for k in ["methodology", "technical approach", "work plan", "approach"]):
        return "CRITERION_METHODOLOGY"
    if any(k in combined for k in ["team", "key personnel", "staff experience", "team experience", "resource qualifications"]):
        return "CRITERION_TEAM_EXPERIENCE"
    if any(k in combined for k in ["corporate experience", "firm experience", "past performance", "firm track record", "company experience"]):
        return "CRITERION_CORPORATE_EXPERIENCE"
    if any(k in combined for k in ["pricing", "cost", "financial weight", "commercial proposal", "per-diem", "rate card weight", "financial score"]):
        return "CRITERION_FINANCIAL_WEIGHT"
    if any(k in combined for k in ["interview", "oral presentation", "presentation"]):
        return "CRITERION_PRESENTATION_WEIGHT"
    if any(k in combined for k in ["indigenous", "procurement strategy for indigenous", "psib"]):
        return "CRITERION_INDIGENOUS_WEIGHT"
    if any(k in combined for k in ["esg", "sustainability", "environmental"]):
        return "CRITERION_ESG_WEIGHT"

    # 4. Specific title if not generic
    cand_title = title or stage
    cand_lower = cand_title.lower()
    if cand_lower and len(cand_lower) > 3 and not any(cand_lower == g for g in ["rated criteria", "technical criteria", "mandatory criteria", "evaluation", "rated", "technical", "stage"]):
        clean_title = re.sub(r'[^a-z0-9]+', '_', cand_lower).strip('_').upper()
        return f"CRITERION_TITLE_{clean_title}"

    return None


def _extract_operational_scope(text: str, source_doc: str = "") -> str:
    """
    Extract explicit operational scope marker from procurement text and source document name.

    Recognises only explicit procurement scope markers such as:
        Category 1, Category A, Cat 2, Cat B
        Stream 1, Workstream 2
        Lot 1, Lot A
        Work Package 3
        Service Category 2

    Does NOT infer scope from subject-matter keywords (e.g. 'HR Advisory',
    'Facilitation', 'Learning & Development'). Those are GENERAL_SCOPE
    unless an explicit marker is also present.

    Bare substrings like 'd1', 'd2', 'd3' are ignored because they appear
    commonly in filenames/identifiers and are not reliable scope indicators.
    """
    combined = f"{text} {source_doc}".lower()

    # ── Explicit category markers ─────────────────────────────────────────────
    # Must be preceded by 'category', 'cat', 'service category' etc. so that
    # bare 'd1' or 'b2' in a filename or criterion ID are not mismatched.
    m = re.search(
        r'\b(?:service\s+)?cat(?:egory)?\s+([a-z0-9]+)\b',
        combined
    )
    if m:
        marker = m.group(1).upper()
        # Only accept single-digit / single-letter markers (Category 1, Cat A)
        if re.match(r'^[0-9A-Z]$', marker):
            return f"CATEGORY_{marker}"
        # Also accept two-digit markers (Category 10)
        if re.match(r'^[0-9]{1,2}$', marker):
            return f"CATEGORY_{marker}"

    # ── Explicit stream / workstream markers ─────────────────────────────────
    m = re.search(r'\b(?:work\s*)?stream\s+([a-z0-9]+)\b', combined)
    if m:
        marker = m.group(1).upper()
        if re.match(r'^[0-9A-Z]{1,2}$', marker):
            return f"STREAM_{marker}"

    # ── Explicit lot markers ──────────────────────────────────────────────────
    m = re.search(r'\blot\s+([a-z0-9]+)\b', combined)
    if m:
        marker = m.group(1).upper()
        if re.match(r'^[0-9A-Z]{1,2}$', marker):
            return f"LOT_{marker}"

    # ── Explicit work-package markers ─────────────────────────────────────────
    m = re.search(r'\bwork\s+package\s+([a-z0-9]+)\b', combined)
    if m:
        marker = m.group(1).upper()
        if re.match(r'^[0-9A-Z]{1,2}$', marker):
            return f"WORK_PACKAGE_{marker}"

    return "GENERAL_SCOPE"


def _extract_evaluation_scope(ec: dict, s_doc: str) -> str:
    """Wrapper: derive operational scope for an evaluation criterion record."""
    stage = (ec.get("stage") or "") if isinstance(ec, dict) else ""
    criterion = (ec.get("criterion") or ec.get("title") or "") if isinstance(ec, dict) else ""
    notes = (ec.get("notes") or "") if isinstance(ec, dict) else ""
    return _extract_operational_scope(f"{stage} {criterion} {notes}", s_doc)


def _extract_requirement_scope(req: dict, s_doc: str) -> str:
    """Wrapper: derive operational scope for a mandatory requirement record."""
    cat = (req.get("category") or "") if isinstance(req, dict) else ""
    rfso_ref = (req.get("rfso_ref") or "") if isinstance(req, dict) else ""
    desc = (req.get("description") or "") if isinstance(req, dict) else ""
    return _extract_operational_scope(f"{desc} {cat} {rfso_ref}", s_doc)


def _extract_deliverable_scope(d: dict, s_doc: str) -> str:
    """Wrapper: derive operational scope for a deliverable record."""
    title = (d.get("title") or d.get("item") or d.get("name") or "") if isinstance(d, dict) else ""
    desc = (d.get("description") or d.get("details") or "") if isinstance(d, dict) else ""
    return _extract_operational_scope(f"{title} {desc}", s_doc)




def _extract_requirement_subject(req: dict) -> str:
    """
    Extract normalized role / subject / criterion identity for mandatory requirement.
    Preserves role precedence, prioritizes explicit requirement IDs and subject domains
    before falling back to generic corporate subject.
    """
    if not isinstance(req, dict):
        return "SUBJECT_GENERAL"

    desc = (req.get("description") or "").lower()
    title = (req.get("title") or req.get("item") or req.get("name") or "").lower()
    rfso = (req.get("rfso_ref") or "").lower()
    req_id = (req.get("req_id") or "").lower()
    combined = f"{req_id} {title} {rfso} {desc}"

    # 1. Role identities take semantic precedence where present
    if any(k in combined for k in ["project manager", "project lead", "engagement lead", "team lead", "pm role"]):
        return "ROLE_PROJECT_MANAGER"
    if any(k in combined for k in ["facilitator", "session facilitator", "lead facilitator"]):
        return "ROLE_FACILITATOR"
    if any(k in combined for k in ["executive coach", "coaching resource", "coach"]):
        return "ROLE_EXECUTIVE_COACH"
    if any(k in combined for k in ["senior advisor", "senior consultant", "advisor"]):
        return "ROLE_SENIOR_ADVISOR"
    if any(k in combined for k in ["consultant", "junior consultant", "analyst"]):
        return "ROLE_CONSULTANT"
    if any(k in combined for k in ["instructional designer", "learning specialist"]):
        return "ROLE_INSTRUCTIONAL_DESIGNER"

    # 2. Explicit requirement ID or RFSO reference marker (e.g. M1, M2, M3, M4, M5, CR1)
    m_id = re.search(r'\b(m\d+|cr\d+|mandatory\s+\d+|r\d+)\b', f"{req_id} {rfso} {title} {desc[:40]}")
    if m_id:
        return f"REF_{m_id.group(1).upper()}"

    if req_id and len(req_id) <= 8 and re.match(r'^[a-z0-9_\-]+$', req_id):
        return f"REF_{req_id.upper()}"

    # 3. Normalized requirement subject/domain
    if any(k in combined for k in ["consulting", "advisory", "management consulting"]):
        return "SUBJECT_CONSULTING_EXPERIENCE"
    if any(k in combined for k in ["training", "learning", "curriculum", "course delivery"]):
        return "SUBJECT_TRAINING_EXPERIENCE"
    if any(k in combined for k in ["coaching", "executive development"]):
        return "SUBJECT_COACHING_EXPERIENCE"
    if any(k in combined for k in ["audit", "accounting", "financial review"]):
        return "SUBJECT_AUDIT_EXPERIENCE"
    if any(k in combined for k in ["software", "application development", "it systems"]):
        return "SUBJECT_IT_EXPERIENCE"
    if any(k in combined for k in ["cybersecurity", "information security"]):
        return "SUBJECT_CYBERSECURITY"
    if any(k in combined for k in ["financial capability", "annual revenue", "financial stability", "bank reference"]):
        return "SUBJECT_FINANCIAL_CAPABILITY"
    if any(k in combined for k in ["litigation", "legal proceeding", "insolvency"]):
        return "SUBJECT_LITIGATION"
    if any(k in combined for k in ["bilingual", "french"]):
        return "SUBJECT_BILINGUAL"
    if any(k in combined for k in ["iso", "quality assurance", "quality management"]):
        return "SUBJECT_QUALITY_MANAGEMENT"

    # 4. Specific title if not generic
    if title and len(title) > 3 and not any(title == g for g in ["mandatory", "mandatory criteria", "mandatory requirement", "qualification"]):
        clean_title = re.sub(r'[^a-z0-9]+', '_', title).strip('_').upper()
        return f"SUBJECT_TITLE_{clean_title}"

    # 5. General corporate fallback
    if any(k in combined for k in ["bidder", "proponent", "firm", "corporate", "vendor track record", "organization"]):
        return "SUBJECT_BIDDER_CORPORATE"

    return "SUBJECT_GENERAL"



def _extract_deliverable_identity(d: dict) -> str | None:
    """Extract normalized identity for a scope deliverable."""
    title = (d.get("title") or d.get("item") or d.get("name") or "").lower() if isinstance(d, dict) else ""
    desc = (d.get("description") or d.get("details") or "").lower() if isinstance(d, dict) else ""
    combined = f"{title} {desc}"

    if any(k in combined for k in ["leadership cohort", "training cohort", "leadership development cohort", "leadership program", "cohort"]):
        return "DELIVERABLE_COHORT"
    if any(k in combined for k in ["executive coaching", "coaching session", "coaching hours", "coaching"]):
        return "DELIVERABLE_EXECUTIVE_COACHING"
    if any(k in combined for k in ["workshop", "training session", "facilitation session"]):
        return "DELIVERABLE_WORKSHOP"
    if any(k in combined for k in ["advisory report", "assessment report", "evaluation report", "roadmap"]):
        return "DELIVERABLE_ADVISORY_REPORT"
    return None


def _extract_deliverable_quantity(d: dict) -> tuple[float, str] | None:
    """Extract normalized quantity and unit for a deliverable."""
    desc = (d.get("description") or d.get("details") or "").lower() if isinstance(d, dict) else ""
    title = (d.get("title") or d.get("item") or "").lower() if isinstance(d, dict) else ""
    combined = f"{title} {desc}"

    # Strip category / stream / appendix prefixes so category numbers (e.g. Category 1) are not mistaken for quantities
    clean_combined = re.sub(r'\b(?:category|cat|stream|appendix|annex|envelope|item|req|m|cr|r)\s*[a-z0-9]+\b', '', combined, flags=re.IGNORECASE)

    m = re.search(r'\b(\d+(?:\.\d+)?)\s*(?:training\s+|leadership\s+|executive\s+|facilitation\s+)*(cohorts?|sessions?|participants?|hours?|workshops?|deliverables?|reports?)\b', clean_combined)
    if m:
        qty = float(m.group(1))
        unit = m.group(2).rstrip('s')
        return qty, unit
    return None


def _is_physical_file(doc_name: str, package_files: list[str]) -> bool:
    """Check if a cited document name corresponds to a real physical file in the procurement package."""
    if not doc_name or not isinstance(doc_name, str):
        return False
    doc_clean = doc_name.strip()
    if doc_clean in ["General RFP Overview", "Summary", "Executive Brief", "Model Interpretation", "Package Overview", "Doc A", "Doc B"]:
        return False
    doc_base = doc_clean.split("/")[-1].split("\\")[-1]
    for pf in package_files:
        pf_base = pf.split("/")[-1].split("\\")[-1]
        if doc_clean == pf or doc_base == pf_base or doc_clean.lower() == pf.lower():
            return True
    return False


def validate_conflict_source_validity(source_a: dict, source_b: dict, package_files: list[str]) -> str:
    """
    Verify whether both cited filenames resolve to physical package files in package_files
    (PHYSICAL_BOTH, PHYSICAL_PARTIAL, or SYNTHESIZED).
    Note: Validates physical filename presence in procurement package; does not perform coordinate-level source_refs validation.
    """
    doc_a = source_a.get("doc", "") if isinstance(source_a, dict) else ""
    doc_b = source_b.get("doc", "") if isinstance(source_b, dict) else ""

    is_a_phys = _is_physical_file(doc_a, package_files)
    is_b_phys = _is_physical_file(doc_b, package_files)

    if is_a_phys and is_b_phys:
        return "PHYSICAL_BOTH"
    elif is_a_phys or is_b_phys:
        return "PHYSICAL_PARTIAL"
    return "SYNTHESIZED"


def _clean_optional_text(value) -> str:
    """Return stripped text if value is a non-empty string, else empty string."""
    return value.strip() if isinstance(value, str) else ""


def detect_document_conflicts(normalized_facts: dict, package_files: list[str]) -> list[dict]:
    """
    Refined Stage C Deterministic Cross-Document Reconciliation Engine.
    Distinguishes:
      - TRUE_CONFLICT: Incompatible physical statements across distinct physical documents within the SAME semantic event / dimension / scope / subject.
      - REVIEW_ITEM: Unresolved scope nuance, internal document discrepancy, or ambiguity.
      - Selects actual opposing source pairs (source_a != source_b) for all detected conflicts.
    """
    conflicts = []
    conflict_idx = 1

    # ── 1. DATE CONFLICTS ─────────────────────────────────────────────────────
    dates = normalized_facts.get("dates", [])
    if len(dates) >= 2:
        actionable_milestones = {
            "SUBMISSION_DEADLINE", "QUESTION_DEADLINE", "SITE_VISIT_DATE",
            "PRESENTATION_DATE", "AWARD_DATE", "CONTRACT_START", "CONTRACT_END", "VALIDITY_DATE"
        }
        milestone_map = {}
        for d in dates:
            m_key = d.get("milestone", "")
            event_type = classify_date_milestone(m_key)
            if event_type in actionable_milestones:
                milestone_map.setdefault(event_type, []).append(d)

        for event_type, d_list in milestone_map.items():
            doc_vals = {}
            for d in d_list:
                s_doc = d.get("source_doc", "Doc")
                dt = _clean_optional_text(d.get("date"))
                if dt:
                    doc_vals.setdefault(s_doc, set()).add(dt)

            # Check internal inconsistencies
            for s_doc, vals in doc_vals.items():
                if len(vals) > 1:
                    internal_list = [d for d in d_list if d.get("source_doc") == s_doc]
                    pair = select_opposing_pair(internal_list, lambda x: _clean_optional_text(x.get("date")))
                    if pair:
                        d_a, d_b = pair
                        src_a = {"doc": s_doc, "ref": d_a.get("milestone", ""), "text": d_a.get("date", "")}
                        src_b = {"doc": s_doc, "ref": d_b.get("milestone", ""), "text": d_b.get("date", "")}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        conflicts.append({
                            "conflict_id": f"CONF-DATE-{conflict_idx}",
                            "conflict_type": "DATE_CONFLICT",
                            "classification": "REVIEW_ITEM",
                            "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                            "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing dates for {event_type.replace('_', ' ').title()}.",
                            "source_validity": sv,
                            "topic": f"Internal Discrepancy for {event_type.replace('_', ' ').title()} in {s_doc}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Conflicting target dates detected within {s_doc}: {', '.join(sorted(vals))}.",
                            "recommended_action": "Verify if an addendum or revision formally clarifies the authoritative date."
                        })
                        conflict_idx += 1

            # Check cross-document conflicts
            single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
            if len(set(single_val_docs.values())) > 1:
                cross_candidates = [d for d in d_list if d.get("source_doc") in single_val_docs]
                cross_pair = select_opposing_pair(cross_candidates, lambda x: _clean_optional_text(x.get("date")), source_fn=lambda x: x.get("source_doc"), require_different_sources=True)
                if cross_pair:
                    d_a, d_b = cross_pair
                    src_a = {"doc": d_a.get("source_doc", "Doc A"), "ref": d_a.get("milestone", ""), "text": d_a.get("date", "")}
                    src_b = {"doc": d_b.get("source_doc", "Doc B"), "ref": d_b.get("milestone", ""), "text": d_b.get("date", "")}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                    unique_dates = sorted(set(single_val_docs.values()))
                    conflicts.append({
                        "conflict_id": f"CONF-DATE-{conflict_idx}",
                        "conflict_type": "DATE_CONFLICT",
                        "classification": classification,
                        "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                        "reason": f"Conflicting dates detected for the same semantic milestone ({event_type.replace('_', ' ').title()}) across documents.",
                        "source_validity": sv,
                        "topic": f"Differing dates for {event_type.replace('_', ' ').title()}",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Conflicting target dates detected: {', '.join(unique_dates)}.",
                        "recommended_action": "Verify if an addendum or revision formally clarifies the authoritative date."
                    })
                    conflict_idx += 1

    # ── 2. EVALUATION CONFLICTS ───────────────────────────────────────────────
    eval_criteria = normalized_facts.get("evaluation_criteria", [])
    if len(eval_criteria) >= 2:
        eval_by_identity = {}
        for ec in eval_criteria:
            s_doc = ec.get("source_doc", "Document")
            ident = _extract_eval_criterion_identity(ec)
            scope = _extract_evaluation_scope(ec, s_doc)
            if ident:
                wt = str(ec.get("weight") or ec.get("points") or "").strip()
                if wt:
                    eval_by_identity.setdefault((ident, scope), []).append((s_doc, wt, ec))

        for (ident, scope), ec_list in eval_by_identity.items():
            doc_vals = {}
            for item in ec_list:
                doc_vals.setdefault(item[0], set()).add(item[1])

            scope_suffix = f" ({scope.replace('_', ' ').title()})" if scope != "GENERAL_SCOPE" else ""
            ident_label = f"{ident.replace('CRITERION_', '').replace('_', ' ').title()}{scope_suffix}"

            # Internal
            for s_doc, vals in doc_vals.items():
                if len(vals) > 1:
                    internal_list = [item for item in ec_list if item[0] == s_doc]
                    pair = select_opposing_pair(internal_list, lambda x: x[1])
                    if pair:
                        src_a = {"doc": s_doc, "ref": ident_label, "text": pair[0][1]}
                        src_b = {"doc": s_doc, "ref": ident_label, "text": pair[1][1]}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        conflicts.append({
                            "conflict_id": f"CONF-EVAL-{conflict_idx}",
                            "conflict_type": "EVALUATION_CONFLICT",
                            "classification": "REVIEW_ITEM",
                            "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                            "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing scoring values for {ident_label}.",
                            "source_validity": sv,
                            "topic": f"Internal Discrepancy for {ident_label} in {s_doc}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Differing scoring values within {s_doc}: {', '.join(sorted(vals))}.",
                            "recommended_action": "Submit clarification to confirm authoritative evaluation weighting formula."
                        })
                        conflict_idx += 1

            # Cross-document
            single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
            if len(set(single_val_docs.values())) > 1:
                cross_candidates = [item for item in ec_list if item[0] in single_val_docs]
                cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                if cross_pair:
                    src_a = {"doc": cross_pair[0][0], "ref": ident_label, "text": cross_pair[0][1]}
                    src_b = {"doc": cross_pair[1][0], "ref": ident_label, "text": cross_pair[1][1]}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                    unique_wts = sorted(set(single_val_docs.values()))
                    conflicts.append({
                        "conflict_id": f"CONF-EVAL-{conflict_idx}",
                        "conflict_type": "EVALUATION_CONFLICT",
                        "classification": classification,
                        "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                        "reason": f"Evaluation breakdown specifies contradictory scoring weights for {ident_label} across documents.",
                        "source_validity": sv,
                        "topic": f"Differing Evaluation Scoring Weights for {ident_label}",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Evaluation breakdown specifies differing scoring weights ({', '.join(unique_wts)}).",
                        "recommended_action": "Submit clarification to confirm authoritative evaluation weighting formula."
                    })
                    conflict_idx += 1

    # ── 3. SUBMISSION RULE CONFLICTS ──────────────────────────────────────────
    sub_rules = normalized_facts.get("submission_rules", [])
    if len(sub_rules) >= 2:
        dim_rules = {}
        for sr in sub_rules:
            dim = classify_submission_rule_dimension(sr.get("item", ""), sr.get("format", ""), sr.get("details", ""))
            dim_rules.setdefault(dim, []).append(sr)

        # Dimension A: ENVELOPE_STRUCTURE (Separate vs Combined)
        env_rules = dim_rules.get("ENVELOPE_STRUCTURE", [])
        if len(env_rules) >= 2:
            doc_env = {}
            for r in env_rules:
                doc = r.get("source_doc", "Doc")
                txt = f"{r.get('item','')} {r.get('format','')} {r.get('details','')}".lower()
                is_sep = "separate" in txt
                is_comb = any(k in txt for k in ["single", "combined", "single package", "single combined", "combined proposal"])
                if is_sep:
                    doc_env.setdefault(doc, {}).setdefault("sep", []).append(r)
                if is_comb:
                    doc_env.setdefault(doc, {}).setdefault("comb", []).append(r)

            # Internal inconsistencies
            for s_doc, env_dict in doc_env.items():
                if "sep" in env_dict and "comb" in env_dict:
                    sep_r = env_dict["sep"][0]
                    comb_r = env_dict["comb"][0]
                    src_a = {"doc": s_doc, "ref": "Submission Rules", "text": str(sep_r.get("format") or sep_r.get("item", ""))}
                    src_b = {"doc": s_doc, "ref": "Submission Rules", "text": str(comb_r.get("format") or comb_r.get("item", ""))}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    conflicts.append({
                        "conflict_id": f"CONF-SUB-{conflict_idx}",
                        "conflict_type": "SUBMISSION_RULE_CONFLICT",
                        "classification": "REVIEW_ITEM",
                        "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                        "reason": f"Internal source inconsistency: same document ({s_doc}) contains contradictory envelope instructions (separate vs combined submission).",
                        "source_validity": sv,
                        "topic": f"Internal Envelope Discrepancy in {s_doc}",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Conflicting envelope instructions detected within {s_doc}.",
                        "recommended_action": "Always separate Financial Envelope from Technical Proposal to prevent mandatory disqualification."
                    })
                    conflict_idx += 1

            # Cross-document
            docs_with_only_sep = [doc for doc, env_dict in doc_env.items() if "sep" in env_dict and "comb" not in env_dict]
            docs_with_only_comb = [doc for doc, env_dict in doc_env.items() if "comb" in env_dict and "sep" not in env_dict]
            if docs_with_only_sep and docs_with_only_comb:
                sep_r = doc_env[docs_with_only_sep[0]]["sep"][0]
                comb_r = doc_env[docs_with_only_comb[0]]["comb"][0]
                src_a = {"doc": sep_r.get("source_doc", "Doc A"), "ref": "Submission Rules", "text": str(sep_r.get("format") or sep_r.get("item", ""))}
                src_b = {"doc": comb_r.get("source_doc", "Doc B"), "ref": "Submission Rules", "text": str(comb_r.get("format") or comb_r.get("item", ""))}
                sv = validate_conflict_source_validity(src_a, src_b, package_files)
                classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"

                conflicts.append({
                    "conflict_id": f"CONF-SUB-{conflict_idx}",
                    "conflict_type": "SUBMISSION_RULE_CONFLICT",
                    "classification": classification,
                    "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                    "reason": "Contradictory envelope separation instructions (separate vs combined submission) across documents.",
                    "source_validity": sv,
                    "topic": "Envelope / Document Separation Contradiction",
                    "source_a": src_a,
                    "source_b": src_b,
                    "assessment": "One document indicates separate technical/financial files while another suggests combined submission.",
                    "recommended_action": "Always separate Financial Envelope from Technical Proposal to prevent mandatory disqualification."
                })
                conflict_idx += 1

        # Dimension B: SUBMISSION_CHANNEL (e.g. MERX/Portal vs Email only)
        chan_rules = dim_rules.get("SUBMISSION_CHANNEL", [])
        if len(chan_rules) >= 2:
            doc_chan = {}
            for r in chan_rules:
                doc = r.get("source_doc", "Doc")
                txt = f"{r.get('item','')} {r.get('format','')} {r.get('details','')}".lower()
                is_portal = any(k in txt for k in ["merx", "buyandsell", "portal", "electronic bid submission", "upload bid"])
                is_email = any(k in txt for k in ["email only", "email submission only", "via email only", "courier only", "hardcopy only"])
                if is_portal:
                    doc_chan.setdefault(doc, {}).setdefault("portal", []).append(r)
                if is_email:
                    doc_chan.setdefault(doc, {}).setdefault("email", []).append(r)

            # Internal inconsistencies
            for s_doc, chan_dict in doc_chan.items():
                if "portal" in chan_dict and "email" in chan_dict:
                    port_r = chan_dict["portal"][0]
                    email_r = chan_dict["email"][0]
                    src_a = {"doc": s_doc, "ref": "Submission Channel", "text": str(port_r.get("format") or port_r.get("details", ""))}
                    src_b = {"doc": s_doc, "ref": "Submission Channel", "text": str(email_r.get("format") or email_r.get("details", ""))}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    conflicts.append({
                        "conflict_id": f"CONF-SUB-{conflict_idx}",
                        "conflict_type": "SUBMISSION_RULE_CONFLICT",
                        "classification": "REVIEW_ITEM",
                        "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                        "reason": f"Internal source inconsistency: same document ({s_doc}) contains conflicting submission channels (portal upload vs direct email/hardcopy).",
                        "source_validity": sv,
                        "topic": f"Internal Submission Channel Discrepancy in {s_doc}",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Conflicting submission channels specified within {s_doc}.",
                        "recommended_action": "Verify authoritative submission channel with contracting authority."
                    })
                    conflict_idx += 1

            # Cross-document
            docs_with_only_portal = [doc for doc, chan_dict in doc_chan.items() if "portal" in chan_dict and "email" not in chan_dict]
            docs_with_only_email = [doc for doc, chan_dict in doc_chan.items() if "email" in chan_dict and "portal" not in chan_dict]
            if docs_with_only_portal and docs_with_only_email:
                port_r = doc_chan[docs_with_only_portal[0]]["portal"][0]
                email_r = doc_chan[docs_with_only_email[0]]["email"][0]
                src_a = {"doc": port_r.get("source_doc", "Doc A"), "ref": "Submission Channel", "text": str(port_r.get("format") or port_r.get("details", ""))}
                src_b = {"doc": email_r.get("source_doc", "Doc B"), "ref": "Submission Channel", "text": str(email_r.get("format") or email_r.get("details", ""))}
                sv = validate_conflict_source_validity(src_a, src_b, package_files)
                classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"

                conflicts.append({
                    "conflict_id": f"CONF-SUB-{conflict_idx}",
                    "conflict_type": "SUBMISSION_RULE_CONFLICT",
                    "classification": classification,
                    "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                    "reason": "Contradictory transmission channels specified across procurement documents.",
                    "source_validity": sv,
                    "topic": "Conflicting Submission Channels (Portal vs Email/Hardcopy)",
                    "source_a": src_a,
                    "source_b": src_b,
                    "assessment": "Documents specify conflicting submission channels (portal upload vs direct email/hardcopy).",
                    "recommended_action": "Verify authoritative submission channel with contracting authority."
                })
                conflict_idx += 1

        # Dimension C: PAGE_LIMIT (e.g. 10 pages vs 15 pages for the same scope)
        page_rules = dim_rules.get("PAGE_LIMIT", [])
        if len(page_rules) >= 2:
            scope_limits = {}
            for pr in page_rules:
                combined = f"{pr.get('item','')} {pr.get('format','')} {pr.get('details','')}".lower()
                doc = pr.get("source_doc", "Doc")

                # Ignore per-sample or per-resume sub-caps
                if any(k in combined for k in ["per sample", "per resume", "per individual", "per profile", "sample of", "samples of"]):
                    continue

                scope = "GENERAL_PROPOSAL"
                if "d1" in doc.lower() or "category 1" in combined:
                    scope = "CATEGORY_1"
                elif "d2" in doc.lower() or "category 2" in combined:
                    scope = "CATEGORY_2"
                elif "d3" in doc.lower() or "category 3" in combined:
                    scope = "CATEGORY_3"

                match = re.search(r'\b(?:exceed|maximum|limit\s+of|max)\s*(?:of\s*)?(\d+)\s*pages?\b', combined) or re.search(r'\b(\d+)\s*pages?\s*(?:max|limit)\b', combined)
                if not match:
                    match = re.search(r'\b(\d+)\s*pages?\b', combined)
                if match:
                    scope_limits.setdefault(scope, []).append((doc, int(match.group(1)), pr))

            for scope, p_list in scope_limits.items():
                doc_vals = {}
                for item in p_list:
                    doc_vals.setdefault(item[0], set()).add(item[1])

                # Internal
                for s_doc, vals in doc_vals.items():
                    if len(vals) > 1:
                        internal_list = [item for item in p_list if item[0] == s_doc]
                        pair = select_opposing_pair(internal_list, lambda x: x[1])
                        if pair:
                            src_a = {"doc": s_doc, "ref": f"Page Limit ({scope.replace('_', ' ').title()})", "text": f"{pair[0][1]} pages"}
                            src_b = {"doc": s_doc, "ref": f"Page Limit ({scope.replace('_', ' ').title()})", "text": f"{pair[1][1]} pages"}
                            sv = validate_conflict_source_validity(src_a, src_b, package_files)
                            conflicts.append({
                                "conflict_id": f"CONF-SUB-{conflict_idx}",
                                "conflict_type": "SUBMISSION_RULE_CONFLICT",
                                "classification": "REVIEW_ITEM",
                                "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                                "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing page limits for {scope.replace('_', ' ').title()}.",
                                "source_validity": sv,
                                "topic": f"Internal Page Limit Discrepancy in {s_doc}",
                                "source_a": src_a,
                                "source_b": src_b,
                                "assessment": f"Conflicting page limits within {s_doc} ({', '.join(str(u) for u in sorted(vals))} pages).",
                                "recommended_action": "Strictly comply with the most restrictive page limit unless clarified."
                            })
                            conflict_idx += 1

                # Cross-document
                single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
                if len(set(single_val_docs.values())) > 1:
                    cross_candidates = [item for item in p_list if item[0] in single_val_docs]
                    cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                    if cross_pair:
                        p_a, p_b = cross_pair
                        src_a = {"doc": p_a[0], "ref": f"Page Limit ({scope.replace('_', ' ').title()})", "text": f"{p_a[1]} pages"}
                        src_b = {"doc": p_b[0], "ref": f"Page Limit ({scope.replace('_', ' ').title()})", "text": f"{p_b[1]} pages"}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                        unique_limits = sorted(set(single_val_docs.values()))
                        conflicts.append({
                            "conflict_id": f"CONF-SUB-{conflict_idx}",
                            "conflict_type": "SUBMISSION_RULE_CONFLICT",
                            "classification": classification,
                            "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                            "reason": f"Contradictory maximum page limits specified for {scope.replace('_', ' ').title()}.",
                            "source_validity": sv,
                            "topic": f"Differing Page Limits for {scope.replace('_', ' ').title()}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Conflicting page limit thresholds detected ({', '.join(str(u) for u in unique_limits)} pages).",
                            "recommended_action": "Strictly comply with the most restrictive page limit unless clarified."
                        })
                        conflict_idx += 1

    # ── 4. MANDATORY REQUIREMENT CONFLICTS ────────────────────────────────────
    reqs = normalized_facts.get("requirements", [])
    mand_reqs = [r for r in reqs if r.get("category") == "Mandatory"]
    if len(mand_reqs) >= 2:
        req_topics = {}
        for r in mand_reqs:
            desc = r.get("description", "")
            s_doc = "Doc"
            for sref in r.get("source_refs", []):
                if isinstance(sref, dict) and sref.get("source_doc"):
                    s_doc = sref.get("source_doc")
                    break

            scope = _extract_requirement_scope(r, s_doc)
            subject = _extract_requirement_subject(r)

            # Check years of experience contradictions
            exp_match = re.search(r'\b(?:minimum\s+)?(\d+)\s+years?\b', desc.lower())
            if exp_match and any(k in desc.lower() for k in ["experience", "advisory", "consulting", "track record"]):
                req_topics.setdefault(("YEARS_OF_EXPERIENCE", scope, subject), []).append((s_doc, int(exp_match.group(1)), desc, r))

            # Check security clearance contradictions
            sec_level = classify_security_clearance(desc)
            if sec_level:
                req_topics.setdefault(("SECURITY_CLEARANCE", scope, subject), []).append((s_doc, sec_level, desc, r))

        for (topic, scope, subject), r_list in req_topics.items():
            doc_vals = {}
            for item in r_list:
                doc_vals.setdefault(item[0], set()).add(item[1])

            scope_label = scope.replace("_", " ").title()
            subj_label = subject.replace("ROLE_", "").replace("SUBJECT_", "").replace("REF_", "").replace("_", " ").title()
            criteria_label = f"{topic.replace('_', ' ').title()} ({scope_label} - {subj_label})"

            # Internal
            for s_doc, vals in doc_vals.items():
                if len(vals) > 1:
                    internal_list = [item for item in r_list if item[0] == s_doc]
                    pair = select_opposing_pair(internal_list, lambda x: x[1])
                    if pair:
                        src_a = {"doc": s_doc, "ref": f"Mandatory Criteria ({criteria_label})", "text": pair[0][2][:120]}
                        src_b = {"doc": s_doc, "ref": f"Mandatory Criteria ({criteria_label})", "text": pair[1][2][:120]}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        conflicts.append({
                            "conflict_id": f"CONF-MAND-{conflict_idx}",
                            "conflict_type": "MANDATORY_REQUIREMENT_CONFLICT",
                            "classification": "REVIEW_ITEM",
                            "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                            "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing requirements for {criteria_label}.",
                            "source_validity": sv,
                            "topic": f"Internal Discrepancy for {criteria_label} in {s_doc}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Differing thresholds within {s_doc} ({', '.join(str(v) for v in sorted(vals))}).",
                            "recommended_action": "Seek authoritative clarification to ensure compliance response aligns with latest standard."
                        })
                        conflict_idx += 1

            # Cross-document
            single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
            if len(set(single_val_docs.values())) > 1:
                cross_candidates = [item for item in r_list if item[0] in single_val_docs]
                cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                if cross_pair:
                    r_a, r_b = cross_pair
                    src_a = {"doc": r_a[0], "ref": f"Mandatory Criteria ({criteria_label})", "text": r_a[2][:120]}
                    src_b = {"doc": r_b[0], "ref": f"Mandatory Criteria ({criteria_label})", "text": r_b[2][:120]}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                    unique_vals = sorted(set(str(v) for v in single_val_docs.values()))
                    conflicts.append({
                        "conflict_id": f"CONF-MAND-{conflict_idx}",
                        "conflict_type": "MANDATORY_REQUIREMENT_CONFLICT",
                        "classification": classification,
                        "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                        "reason": f"Contradictory mandatory criteria detected for {criteria_label} across documents.",
                        "source_validity": sv,
                        "topic": f"Conflicting Mandatory Requirements ({criteria_label})",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Differing mandatory requirement thresholds stated across procurement documents ({', '.join(unique_vals)}).",
                        "recommended_action": "Seek authoritative clarification to ensure compliance response aligns with latest standard."
                    })
                    conflict_idx += 1

    # ── 5. COMMERCIAL TERM CONFLICTS ──────────────────────────────────────────
    comm_clauses = normalized_facts.get("commercial_clauses", [])
    if len(comm_clauses) >= 2:
        # Commercial Caps
        comm_by_topic = {}
        for c in comm_clauses:
            c_topic = classify_commercial_topic(c.get("topic", ""), c.get("details", ""))
            if c_topic not in ["OTHER_COMMERCIAL_TERM", "INSURANCE_REQUIREMENT"]:
                c_limit = extract_commercial_limit(c_topic, f"{c.get('topic','')} {c.get('details','')}")
                if c_limit is not None:
                    s_doc = c.get("source_doc", "Doc")
                    comm_by_topic.setdefault(c_topic, []).append((s_doc, c_limit, c))

        for c_topic, c_list in comm_by_topic.items():
            doc_vals = {}
            for item in c_list:
                doc_vals.setdefault(item[0], set()).add(item[1])

            topic_title = c_topic.replace("_", " ").title()

            # Internal inconsistencies
            for s_doc, vals in doc_vals.items():
                if len(vals) > 1:
                    internal_list = [item for item in c_list if item[0] == s_doc]
                    pair = select_opposing_pair(internal_list, lambda x: x[1])
                    if pair:
                        src_a = {"doc": s_doc, "ref": topic_title, "text": pair[0][2].get("details", "")}
                        src_b = {"doc": s_doc, "ref": topic_title, "text": pair[1][2].get("details", "")}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        conflicts.append({
                            "conflict_id": f"CONF-COMM-{conflict_idx}",
                            "conflict_type": "COMMERCIAL_TERM_CONFLICT",
                            "classification": "REVIEW_ITEM",
                            "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                            "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing commercial limitations for {topic_title}.",
                            "source_validity": sv,
                            "topic": f"Internal Discrepancy for {topic_title} in {s_doc}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Discrepancy in commercial thresholds within {s_doc} ({', '.join(str(v) for v in sorted(vals))}).",
                            "recommended_action": "Verify authoritative commercial ceiling with contracting authority."
                        })
                        conflict_idx += 1

            # Cross-document
            single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
            if len(set(single_val_docs.values())) > 1:
                cross_candidates = [item for item in c_list if item[0] in single_val_docs]
                cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                if cross_pair:
                    p_a, p_b = cross_pair
                    src_a = {"doc": p_a[0], "ref": topic_title, "text": p_a[2].get("details", "")}
                    src_b = {"doc": p_b[0], "ref": topic_title, "text": p_b[2].get("details", "")}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                    unique_limits = sorted(set(str(v) for v in single_val_docs.values()))
                    conflicts.append({
                        "conflict_id": f"CONF-COMM-{conflict_idx}",
                        "conflict_type": "COMMERCIAL_TERM_CONFLICT",
                        "classification": classification,
                        "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                        "reason": f"Contradictory commercial limitations specified for {topic_title} across documents.",
                        "source_validity": sv,
                        "topic": f"Differing Commercial Limits ({topic_title})",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Discrepancy in commercial thresholds across procurement documents ({', '.join(unique_limits)}).",
                        "recommended_action": "Verify authoritative commercial ceiling with contracting authority."
                    })
                    conflict_idx += 1

        # Insurance
        ins_terms = [c for c in comm_clauses if "insurance" in f"{c.get('topic','')} {c.get('details','')}".lower() or "liability" in f"{c.get('topic','')} {c.get('details','')}".lower()]
        if len(ins_terms) >= 2:
            ins_by_class = {}
            for it in ins_terms:
                combined_text = f"{it.get('topic','')} {it.get('details','')}"
                ins_class = classify_insurance_class(combined_text)
                amt = extract_monetary_amount(combined_text)
                s_doc = it.get("source_doc", "Doc")
                ins_by_class.setdefault(ins_class, []).append((s_doc, amt, it))

            for ins_class, i_list in ins_by_class.items():
                if len(i_list) >= 2:
                    doc_vals = {}
                    for item in i_list:
                        if item[1] is not None:
                            doc_vals.setdefault(item[0], set()).add(item[1])

                    ins_title = ins_class.replace('_', ' ').title()

                    # Internal inconsistencies
                    for s_doc, vals in doc_vals.items():
                        if len(vals) > 1:
                            internal_list = [item for item in i_list if item[0] == s_doc and item[1] is not None]
                            pair = select_opposing_pair(internal_list, lambda x: x[1])
                            if pair:
                                item_a, item_b = pair
                                src_a = {"doc": s_doc, "ref": item_a[2].get("topic", ""), "text": item_a[2].get("details", "")}
                                src_b = {"doc": s_doc, "ref": item_b[2].get("topic", ""), "text": item_b[2].get("details", "")}
                                sv = validate_conflict_source_validity(src_a, src_b, package_files)
                                amt_strs = [f"${a:,.0f}" for a in sorted(vals)]
                                conflicts.append({
                                    "conflict_id": f"CONF-COMM-{conflict_idx}",
                                    "conflict_type": "COMMERCIAL_TERM_CONFLICT",
                                    "classification": "REVIEW_ITEM",
                                    "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                                    "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing insurance liability monetary thresholds for {ins_title}.",
                                    "source_validity": sv,
                                    "topic": f"Internal Discrepancy for {ins_title} Insurance in {s_doc}",
                                    "source_a": src_a,
                                    "source_b": src_b,
                                    "assessment": f"Discrepancy in required {ins_class.replace('_', ' ').lower()} coverage amounts within {s_doc} ({', '.join(amt_strs)}).",
                                    "recommended_action": "Confirm authoritative insurance coverage limits with contracting authority."
                                })
                                conflict_idx += 1

                    # Cross-document
                    single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
                    if len(set(single_val_docs.values())) > 1:
                        cross_candidates = [item for item in i_list if item[0] in single_val_docs and item[1] is not None]
                        cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                        if cross_pair:
                            item_a, item_b = cross_pair
                            src_a = {"doc": item_a[0], "ref": item_a[2].get("topic", ""), "text": item_a[2].get("details", "")}
                            src_b = {"doc": item_b[0], "ref": item_b[2].get("topic", ""), "text": item_b[2].get("details", "")}
                            sv = validate_conflict_source_validity(src_a, src_b, package_files)
                            classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                            amt_strs = [f"${a:,.0f}" for a in sorted(set(single_val_docs.values()))]
                            conflicts.append({
                                "conflict_id": f"CONF-COMM-{conflict_idx}",
                                "conflict_type": "COMMERCIAL_TERM_CONFLICT",
                                "classification": classification,
                                "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                                "reason": f"Conflicting insurance liability monetary thresholds across documents for {ins_title}.",
                                "source_validity": sv,
                                "topic": f"Conflicting {ins_title} Insurance Limits",
                                "source_a": src_a,
                                "source_b": src_b,
                                "assessment": f"Discrepancy in required {ins_class.replace('_', ' ').lower()} coverage amounts ({', '.join(amt_strs)}).",
                                "recommended_action": "Confirm authoritative insurance coverage limits with contracting authority."
                            })
                            conflict_idx += 1
                    else:
                        amounts = [item[1] for item in i_list if item[1] is not None]
                        prose_pair = select_opposing_pair(i_list, lambda x: x[2].get("details", "").strip())
                        if len(amounts) == 0 and prose_pair:
                            p_a, p_b = prose_pair
                            src_a = {"doc": p_a[0], "ref": p_a[2].get("topic", ""), "text": p_a[2].get("details", "")}
                            src_b = {"doc": p_b[0], "ref": p_b[2].get("topic", ""), "text": p_b[2].get("details", "")}
                            sv = validate_conflict_source_validity(src_a, src_b, package_files)
                            conflicts.append({
                                "conflict_id": f"CONF-COMM-{conflict_idx}",
                                "conflict_type": "COMMERCIAL_TERM_CONFLICT",
                                "classification": "REVIEW_ITEM",
                                "confidence": "LOW",
                                "reason": f"Prose variation in {ins_title} terms across documents without confirmed monetary contradiction.",
                                "source_validity": sv,
                                "topic": f"Term Wording Variation for {ins_title}",
                                "source_a": src_a,
                                "source_b": src_b,
                                "assessment": "Different wording used for insurance terms across documents; coverage amounts unconfirmed.",
                                "recommended_action": "Review insurance wording to confirm underlying coverage limits match."
                            })
                            conflict_idx += 1

    # ── 6. SCOPE CONFLICTS ────────────────────────────────────────────────────
    deliverables = normalized_facts.get("deliverables", [])
    if len(deliverables) >= 2:
        deliv_groups = {}
        for d in deliverables:
            d_ident = _extract_deliverable_identity(d)
            qty_unit = _extract_deliverable_quantity(d)
            if d_ident and qty_unit:
                qty, unit = qty_unit
                s_doc = d.get("source_doc", "Doc")
                d_scope = _extract_deliverable_scope(d, s_doc)
                deliv_groups.setdefault((d_ident, d_scope, unit), []).append((s_doc, qty, d))

        for (d_ident, d_scope, unit), d_list in deliv_groups.items():
            doc_vals = {}
            for item in d_list:
                doc_vals.setdefault(item[0], set()).add(item[1])

            deliv_title = d_ident.replace('DELIVERABLE_', '').replace('_', ' ').title()
            scope_suffix = f" ({d_scope.replace('_', ' ').title()})" if d_scope != "GENERAL_SCOPE" else ""
            topic_label = f"{deliv_title}{scope_suffix}"

            # Internal inconsistencies
            for s_doc, vals in doc_vals.items():
                if len(vals) > 1:
                    internal_list = [item for item in d_list if item[0] == s_doc]
                    pair = select_opposing_pair(internal_list, lambda x: x[1])
                    if pair:
                        c_a, c_b = pair
                        src_a = {"doc": s_doc, "ref": f"Deliverables ({topic_label})", "text": c_a[2].get("description", "")}
                        src_b = {"doc": s_doc, "ref": f"Deliverables ({topic_label})", "text": c_b[2].get("description", "")}
                        sv = validate_conflict_source_validity(src_a, src_b, package_files)
                        conflicts.append({
                            "conflict_id": f"CONF-SCOPE-{conflict_idx}",
                            "conflict_type": "SCOPE_CONFLICT",
                            "classification": "REVIEW_ITEM",
                            "confidence": "HIGH" if sv == "PHYSICAL_BOTH" else "MEDIUM",
                            "reason": f"Internal source inconsistency: same document ({s_doc}) contains differing deliverable quantities for {topic_label}.",
                            "source_validity": sv,
                            "topic": f"Internal Discrepancy for {topic_label} in {s_doc}",
                            "source_a": src_a,
                            "source_b": src_b,
                            "assessment": f"Differing volume counts within {s_doc} ({', '.join(str(q) for q in sorted(vals))} {unit}s).",
                            "recommended_action": "Seek written clarification on authoritative baseline for pricing."
                        })
                        conflict_idx += 1

            # Cross-document
            single_val_docs = {doc: list(vals)[0] for doc, vals in doc_vals.items() if len(vals) == 1}
            if len(set(single_val_docs.values())) > 1:
                cross_candidates = [item for item in d_list if item[0] in single_val_docs]
                cross_pair = select_opposing_pair(cross_candidates, lambda x: x[1], source_fn=lambda x: x[0], require_different_sources=True)
                if cross_pair:
                    c_a, c_b = cross_pair
                    src_a = {"doc": c_a[0], "ref": f"Deliverables ({topic_label})", "text": c_a[2].get("description", "")}
                    src_b = {"doc": c_b[0], "ref": f"Deliverables ({topic_label})", "text": c_b[2].get("description", "")}
                    sv = validate_conflict_source_validity(src_a, src_b, package_files)
                    classification = "TRUE_CONFLICT" if sv == "PHYSICAL_BOTH" and src_a["doc"] != src_b["doc"] else "REVIEW_ITEM"
                    unique_qtys = sorted(set(single_val_docs.values()))
                    conflicts.append({
                        "conflict_id": f"CONF-SCOPE-{conflict_idx}",
                        "conflict_type": "SCOPE_CONFLICT",
                        "classification": classification,
                        "confidence": "HIGH" if classification == "TRUE_CONFLICT" else "MEDIUM",
                        "reason": f"Conflicting deliverable quantities specified for {topic_label} across documents.",
                        "source_validity": sv,
                        "topic": f"Deliverable Volume Discrepancy ({topic_label})",
                        "source_a": src_a,
                        "source_b": src_b,
                        "assessment": f"Differing volume counts across documents ({', '.join(str(q) for q in unique_qtys)} {unit}s).",
                        "recommended_action": "Seek written clarification on authoritative baseline for pricing."
                    })
                    conflict_idx += 1

    # ── 7. FINAL SOURCE VALIDITY ENFORCEMENT ──────────────────────────────────
    for c in conflicts:
        sv = validate_conflict_source_validity(c.get("source_a", {}), c.get("source_b", {}), package_files)
        c["source_validity"] = sv
        doc_a = c.get("source_a", {}).get("doc", "")
        doc_b = c.get("source_b", {}).get("doc", "")
        # A same-file pair must never be classified as TRUE_CONFLICT even if sv == PHYSICAL_BOTH
        if (sv != "PHYSICAL_BOTH" or doc_a == doc_b) and c.get("classification") == "TRUE_CONFLICT":
            c["classification"] = "REVIEW_ITEM"
            c["confidence"] = "MEDIUM"

    return conflicts


# ── 4-STAGE GENUINE PIPELINE IMPLEMENTATION ───────────────────────────────────

def _clean_raw(raw: str) -> str:
    raw = raw.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.IGNORECASE)
    raw = re.sub(r"\s*```$", "", raw)
    return raw.strip()


def _safe_parse_json_with_status(raw: str) -> tuple[dict, str]:
    """
    Parses JSON with explicit status reporting:
    - (dict, "COMPLETE"): fully valid JSON output
    - (dict, "RECOVERED_TRUNCATED"): truncated output repaired via balanced stack recovery
    - ({}, "FAILED"): unparseable
    """
    # 1. Direct json.loads on cleaned text (fast path)
    try:
        cleaned = _clean_raw(raw)
        res = json.loads(cleaned)
        if isinstance(res, dict):
            return res, "COMPLETE"
    except Exception:
        pass

    # 2. Try analyst._parse_json if it returned a complete dict
    try:
        from analyst import _parse_json
        res = _parse_json(raw)
        if isinstance(res, dict) and not res.get("_truncated"):
            return res, "COMPLETE"
    except Exception:
        pass

    # 3. Dedicated stack-based dictionary repair for truncated JSON outputs
    # Strips code fences, scans backwards from cut point, and balances open tokens using LIFO container stack
    t = raw.strip()
    t = re.sub(r"^```[a-z]*\s*\n?", "", t, flags=re.IGNORECASE)
    t = re.sub(r"\n?```\s*$", "", t).strip()
    start_idx = t.find('{')
    if start_idx != -1:
        src = t[start_idx:]
        max_scan = min(len(src), 4000)
        for end_pos in range(len(src) - 1, len(src) - max_scan, -1):
            prefix = src[:end_pos + 1]
            stack = []
            in_str, esc = False, False
            valid = True
            for ch in prefix:
                if esc:
                    esc = False
                    continue
                if ch == "\\" and in_str:
                    esc = True
                    continue
                if ch == '"' and not esc:
                    in_str = not in_str
                    continue
                if in_str:
                    continue
                if ch in ('{', '['):
                    stack.append('}' if ch == '{' else ']')
                elif ch in ('}', ']'):
                    if not stack or stack[-1] != ch:
                        valid = False
                        break
                    stack.pop()
            if not valid or in_str:
                continue
            suffix = "".join(reversed(stack))
            candidate = prefix + suffix
            try:
                res = json.loads(candidate)
                if isinstance(res, dict):
                    return res, "RECOVERED_TRUNCATED"
            except json.JSONDecodeError:
                pass

    return {}, "FAILED"


def _safe_parse_json(raw: str) -> dict:
    """Backward-compatible wrapper returning only parsed dict."""
    return _safe_parse_json_with_status(raw)[0]
# ── STAGE A: DOCUMENT FACT EXTRACTION ─────────────────────────────────────────

# ── STAGE A CHUNKING, AGGREGATION & COVERAGE GUARD ───────────────────────────

_STAGE_A_MAX_CHUNK_CHARS = 12000
_STAGE_A_MARKER_SPLIT_RE = re.compile(r'(\[\[SOURCE:[^\]]+\]\])', re.IGNORECASE)

# General procurement-language signal patterns for coverage auditing
_REQ_SIGNALS = re.compile(
    r'\b(?:shall|must|required|mandatory|condition(?:s)? of participation|supplier will|tenderer shall)\b',
    re.IGNORECASE
)
_DATE_SIGNALS = re.compile(
    r'\b(?:deadline|response deadline|clarification(?:s)?|timescale(?:s)?|closing date|site visit)\b',
    re.IGNORECASE
)
_EVAL_SIGNALS = re.compile(
    r'(?:\b(?:evaluation|award criteria|scoring|weighted|marks|price weighting)\b|%)',
    re.IGNORECASE
)
_SUB_SIGNALS = re.compile(
    r'\b(?:submit|submission|response checklist|annex|attachment|signed|portal)\b',
    re.IGNORECASE
)


def chunk_document_text(doc_text: str, max_chunk_chars: int = _STAGE_A_MAX_CHUNK_CHARS) -> list[str]:
    """
    Deterministically partition doc_text into chunks using source markers.
    If len(doc_text) <= max_chunk_chars, returns [doc_text].
    Splits along source marker boundaries. If an individual block exceeds
    max_chunk_chars, it is split along line boundaries.
    """
    if not doc_text or len(doc_text) <= max_chunk_chars:
        return [doc_text] if doc_text else []

    parts = _STAGE_A_MARKER_SPLIT_RE.split(doc_text)
    blocks = []
    if parts[0].strip():
        blocks.append(parts[0])
    idx = 1
    while idx < len(parts):
        marker = parts[idx]
        content = parts[idx + 1] if idx + 1 < len(parts) else ""
        blocks.append(marker + content)
        idx += 2

    if not blocks:
        return [doc_text]

    chunks = []
    current_chunk = []
    current_len = 0

    for b in blocks:
        b_len = len(b)
        if b_len > max_chunk_chars:
            if current_chunk:
                chunks.append("".join(current_chunk))
                current_chunk = []
                current_len = 0
            # Extract the leading marker of block b if present
            m_match = _STAGE_A_MARKER_SPLIT_RE.search(b)
            leading_marker = m_match.group(0) + "\n" if m_match else ""
            lines = b.splitlines(keepends=True)
            sub = []
            sub_len = 0
            for line in lines:
                if len(line) > max_chunk_chars:
                    if sub:
                        chunks.append("".join(sub))
                        sub = []
                        sub_len = 0
                    p = line
                    prefix = leading_marker if leading_marker and not p.startswith("[[SOURCE:") else ""
                    while len(p) > 0:
                        max_slice = max_chunk_chars - len(prefix)
                        if max_slice <= 0:
                            take = p[:max_chunk_chars]
                            chunks.append(take)
                            p = p[max_chunk_chars:]
                        else:
                            take = p[:max_slice]
                            chunks.append(prefix + take)
                            p = p[max_slice:]
                            prefix = leading_marker
                    continue

                if sub_len + len(line) > max_chunk_chars and sub:
                    chunks.append("".join(sub))
                    sub = [leading_marker, line] if leading_marker and not line.startswith("[[SOURCE:") else [line]
                    sub_len = sum(len(x) for x in sub)
                else:
                    sub.append(line)
                    sub_len += len(line)
            if sub:
                chunks.append("".join(sub))
        else:
            if current_len + b_len > max_chunk_chars and current_chunk:
                chunks.append("".join(current_chunk))
                current_chunk = [b]
                current_len = b_len
            else:
                current_chunk.append(b)
                current_len += b_len

    if current_chunk:
        chunks.append("".join(current_chunk))

    return chunks


def _source_ref_key(ref: dict) -> tuple:
    if not isinstance(ref, dict):
        return ()
    return (
        ref.get("source_doc") or "",
        ref.get("page"),
        ref.get("sheet") or "",
        ref.get("section") or "",
        (ref.get("excerpt") or "").strip()
    )


def aggregate_stage_a_facts(chunk_facts_list: list[dict], filename: str) -> dict:
    """
    Deterministically merge chunk results into a single document-level Stage A result.
    Preserves all valid facts, merges source_refs, conservative metadata merge,
    and deduplicates only materially identical facts.
    """
    merged = {
        "doc_metadata": {},
        "requirements": [],
        "dates": [],
        "evaluation_criteria": [],
        "submission_rules": [],
        "deliverables": [],
        "commercial_clauses": [],
        "contract_risks": []
        ,"typed_observations": []
    }

    # 1. Conservative metadata merge
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        meta = cf.get("doc_metadata") or {}
        if isinstance(meta, dict):
            for k, v in meta.items():
                if v and not merged["doc_metadata"].get(k):
                    merged["doc_metadata"][k] = v

    # 2. Requirements aggregation & deduplication
    req_seen = {}  # map: canon_key -> dict
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for r in cf.get("requirements") or []:
            if not isinstance(r, dict):
                continue
            desc = (r.get("description") or "").strip()
            if not desc:
                continue
            cat = (r.get("category") or "").strip()
            rfso = (r.get("rfso_ref") or "").strip()
            norm_desc = re.sub(r'\W+', '', desc.lower())

            raw_refs = r.get("source_refs") or []
            if not raw_refs:
                raw_refs = [{"source_doc": filename, "page": None, "sheet": None, "section": None, "excerpt": desc[:100]}]

            canon_key = (norm_desc, cat.lower(), rfso.lower())

            explicit = normalize_requirement_type(r.get("requirement_type"))
            candidate = explicit or resolve_requirement_type(r)

            incoming_candidates = set()
            if isinstance(r.get("_semantic_candidates"), (list, set, tuple)):
                for c in r["_semantic_candidates"]:
                    norm_c = normalize_requirement_type(c)
                    if norm_c in SPECIFIC_REQUIREMENT_TYPES:
                        incoming_candidates.add(norm_c)
            if candidate in SPECIFIC_REQUIREMENT_TYPES:
                incoming_candidates.add(candidate)

            if canon_key in req_seen:
                existing = req_seen[canon_key]
                existing_ref_keys = {_source_ref_key(ref) for ref in existing.get("source_refs", [])}
                for ref in raw_refs:
                    rk = _source_ref_key(ref)
                    if rk not in existing_ref_keys:
                        existing.setdefault("source_refs", []).append(ref)
                        existing_ref_keys.add(rk)
                # True N-way order-independent candidate accumulation
                spec_set = set(existing.get("_semantic_candidates") or [])
                spec_set.update(incoming_candidates)
                existing["_semantic_candidates"] = sorted(spec_set)
                existing["requirement_type"] = resolve_candidate_requirement_types(spec_set)
            else:
                r_copy = dict(r)
                r_copy["source_refs"] = list(raw_refs)
                r_copy["_semantic_candidates"] = sorted(incoming_candidates)
                if incoming_candidates:
                    r_copy["requirement_type"] = resolve_candidate_requirement_types(incoming_candidates)
                else:
                    r_copy["requirement_type"] = resolve_requirement_type(r)
                req_seen[canon_key] = r_copy
                merged["requirements"].append(r_copy)

    # 3. Dates aggregation & deduplication
    seen_dates = set()
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for d in cf.get("dates") or []:
            if not isinstance(d, dict):
                continue
            milestone = (d.get("milestone") or "").strip()
            date_val = str(d.get("date") or "").strip()
            canon = (milestone.lower(), date_val)
            if not milestone or canon in seen_dates:
                continue
            seen_dates.add(canon)
            d_copy = dict(d)
            d_copy.setdefault("source_doc", filename)
            merged["dates"].append(d_copy)

    # Typed observations are the authoritative representation for new executive
    # fact families. Preserve each physical occurrence; canonicalization later
    # collapses repeated extraction of the same occurrence deterministically.
    for cf in chunk_facts_list:
        if isinstance(cf, dict):
            merged["typed_observations"].extend(
                copy for copy in (dict(x) for x in (cf.get("typed_observations") or []) if isinstance(x, dict))
            )

    # 4. Evaluation Criteria aggregation & deduplication using evaluation_hierarchy
    from evaluation_hierarchy import deduplicate_evaluation_criteria, normalize_evaluation_criterion
    raw_eval_list = []
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for ec in cf.get("evaluation_criteria") or []:
            if isinstance(ec, dict):
                ec_copy = dict(ec)
                ec_copy.setdefault("source_doc", filename)
                raw_eval_list.append(ec_copy)
    merged["evaluation_criteria"] = deduplicate_evaluation_criteria(raw_eval_list)

    # 5. Submission Rules aggregation & deduplication
    seen_sub = {}
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for sr in cf.get("submission_rules") or []:
            if not isinstance(sr, dict):
                continue
            item = (sr.get("item") or "").strip()
            fmt = str(sr.get("format") or "").strip()
            details = str(sr.get("details") or "").strip()
            if not item:
                continue
            canon = _canonical_submission_item_identity(item) or item.lower()
            
            # Extract and normalize source_refs (sheet-aware, no fabricated fallbacks)
            incoming_srefs = []
            for sref in (sr.get("source_refs") or []):
                if isinstance(sref, dict):
                    sref_copy = dict(sref)
                    sref_copy.setdefault("source_doc", filename)
                    incoming_srefs.append(sref_copy)

            # Observation tracking
            obs_at = sr.get("artifact_type")
            obs_ff = sr.get("file_format")
            obs_sc = sr.get("submission_channel")
            obs_mand = sr.get("mandatory")

            if canon in seen_sub:
                existing_sr = seen_sub[canon]
                # Merge source_refs with sheet-aware identity
                existing_sref_keys = {(s.get("source_doc"), s.get("page"), s.get("sheet"), s.get("section"), s.get("excerpt")) for s in existing_sr.get("source_refs", []) if isinstance(s, dict)}
                for s in incoming_srefs:
                    k = (s.get("source_doc"), s.get("page"), s.get("sheet"), s.get("section"), s.get("excerpt"))
                    if k not in existing_sref_keys:
                        existing_sr.setdefault("source_refs", []).append(s)
                        existing_sref_keys.add(k)
                
                # Track observations & conflicts
                at_obs = existing_sr.setdefault("artifact_type_observations", [])
                if obs_at and obs_at not in at_obs:
                    at_obs.append(obs_at)
                valid_at = [x for x in at_obs if x and x != "Unknown"]
                is_at_conflict = len(set(valid_at)) > 1
                existing_sr["artifact_type_conflict"] = is_at_conflict
                if is_at_conflict:
                    existing_sr["artifact_type"] = "Unknown"
                elif valid_at and (not existing_sr.get("artifact_type") or existing_sr.get("artifact_type") == "Unknown"):
                    existing_sr["artifact_type"] = valid_at[0]

                ff_obs = existing_sr.setdefault("file_format_observations", [])
                if obs_ff and obs_ff not in ff_obs:
                    ff_obs.append(obs_ff)
                valid_ff = [x for x in ff_obs if x and x != "Unspecified"]
                is_ff_conflict = len(set(valid_ff)) > 1
                existing_sr["file_format_conflict"] = is_ff_conflict
                if is_ff_conflict:
                    existing_sr["file_format"] = "Multiple / Mixed"
                elif valid_ff and (not existing_sr.get("file_format") or existing_sr.get("file_format") == "Unspecified"):
                    existing_sr["file_format"] = valid_ff[0]

                sc_obs = existing_sr.setdefault("submission_channel_observations", [])
                if obs_sc and obs_sc not in sc_obs:
                    sc_obs.append(obs_sc)
                valid_sc = [x for x in sc_obs if x and x != "Unspecified"]
                is_sc_conflict = len(set(valid_sc)) > 1
                existing_sr["submission_channel_conflict"] = is_sc_conflict
                if is_sc_conflict:
                    existing_sr["submission_channel"] = "Unspecified"
                elif valid_sc and (not existing_sr.get("submission_channel") or existing_sr.get("submission_channel") == "Unspecified"):
                    existing_sr["submission_channel"] = valid_sc[0]

                mand_obs = existing_sr.setdefault("mandatory_observations", [])
                if obs_mand is not None and obs_mand not in mand_obs:
                    mand_obs.append(obs_mand)
                valid_mand = [x for x in mand_obs if x is not None]
                is_mand_conflict = len(set(valid_mand)) > 1
                existing_sr["mandatory_conflict"] = is_mand_conflict
                if is_mand_conflict:
                    existing_sr["mandatory"] = None
                elif existing_sr.get("mandatory") is None and obs_mand is not None:
                    existing_sr["mandatory"] = obs_mand

                if not existing_sr.get("format") and fmt:
                    existing_sr["format"] = fmt
                if not existing_sr.get("details") and details:
                    existing_sr["details"] = details
            else:
                sr_copy = dict(sr)
                sr_copy.setdefault("source_doc", filename)
                sr_copy["source_refs"] = incoming_srefs
                sr_copy["artifact_type_observations"] = [obs_at] if obs_at else []
                valid_at = [x for x in sr_copy["artifact_type_observations"] if x and x != "Unknown"]
                is_at_conflict = len(set(valid_at)) > 1
                sr_copy["artifact_type_conflict"] = is_at_conflict
                if is_at_conflict:
                    sr_copy["artifact_type"] = "Unknown"

                sr_copy["file_format_observations"] = [obs_ff] if obs_ff else []
                valid_ff = [x for x in sr_copy["file_format_observations"] if x and x != "Unspecified"]
                is_ff_conflict = len(set(valid_ff)) > 1
                sr_copy["file_format_conflict"] = is_ff_conflict
                if is_ff_conflict:
                    sr_copy["file_format"] = "Multiple / Mixed"

                sr_copy["submission_channel_observations"] = [obs_sc] if obs_sc else []
                valid_sc = [x for x in sr_copy["submission_channel_observations"] if x and x != "Unspecified"]
                is_sc_conflict = len(set(valid_sc)) > 1
                sr_copy["submission_channel_conflict"] = is_sc_conflict
                if is_sc_conflict:
                    sr_copy["submission_channel"] = "Unspecified"

                sr_copy["mandatory_observations"] = [obs_mand] if obs_mand is not None else []
                valid_mand = [x for x in sr_copy["mandatory_observations"] if x is not None]
                is_mand_conflict = len(set(valid_mand)) > 1
                sr_copy["mandatory_conflict"] = is_mand_conflict
                if is_mand_conflict:
                    sr_copy["mandatory"] = None

                seen_sub[canon] = sr_copy
                merged["submission_rules"].append(sr_copy)

    # 6. Deliverables aggregation & deduplication
    seen_deliv = set()
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for deliv in cf.get("deliverables") or []:
            if not isinstance(deliv, dict):
                continue
            title = (deliv.get("title") or deliv.get("item") or "").strip()
            desc = str(deliv.get("description") or "").strip()
            canon = (title.lower(), desc.lower())
            if not title or canon in seen_deliv:
                continue
            seen_deliv.add(canon)
            deliv_copy = dict(deliv)
            deliv_copy.setdefault("source_doc", filename)
            merged["deliverables"].append(deliv_copy)

    # 7. Commercial Clauses aggregation & deduplication
    seen_comm = set()
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for cc in cf.get("commercial_clauses") or []:
            if not isinstance(cc, dict):
                continue
            topic = (cc.get("topic") or "").strip()
            details = str(cc.get("details") or "").strip()
            canon = (topic.lower(), details.lower())
            if not topic or canon in seen_comm:
                continue
            seen_comm.add(canon)
            cc_copy = dict(cc)
            cc_copy.setdefault("source_doc", filename)
            merged["commercial_clauses"].append(cc_copy)

    # 8. Contract Risks aggregation & deduplication
    seen_risks = set()
    for cf in chunk_facts_list:
        if not isinstance(cf, dict):
            continue
        for cr in cf.get("contract_risks") or []:
            if not isinstance(cr, dict):
                continue
            risk = (cr.get("risk") or cr.get("title") or "").strip()
            details = str(cr.get("details") or "").strip()
            canon = (risk.lower(), details.lower())
            if not risk or canon in seen_risks:
                continue
            seen_risks.add(canon)
            cr_copy = dict(cr)
            cr_copy.setdefault("source_doc", filename)
            merged["contract_risks"].append(cr_copy)

    return merged


def inspect_stage_a_coverage(doc_text: str, facts: dict, min_signal_count: int = 5) -> dict:
    """
    Inspect whether extracted facts have suspicious under-coverage given the source text signals.
    Does NOT create or modify procurement facts.
    """
    signals = {
        "requirements": len(_REQ_SIGNALS.findall(doc_text)),
        "dates": len(_DATE_SIGNALS.findall(doc_text)),
        "evaluation_criteria": len(_EVAL_SIGNALS.findall(doc_text)),
        "submission_rules": len(_SUB_SIGNALS.findall(doc_text)),
    }

    suspicious = []
    if signals["requirements"] >= min_signal_count and not facts.get("requirements"):
        suspicious.append("requirements")
    if signals["dates"] >= min_signal_count and not facts.get("dates"):
        suspicious.append("dates")
    if signals["evaluation_criteria"] >= min_signal_count and not facts.get("evaluation_criteria"):
        suspicious.append("evaluation_criteria")
    if signals["submission_rules"] >= min_signal_count and not facts.get("submission_rules"):
        suspicious.append("submission_rules")

    return {
        "is_suspicious": len(suspicious) > 0,
        "suspicious_families": suspicious,
        "signal_counts": signals,
    }


def _extract_chunk_facts(chunk_text: str, filename: str, api_key: str, client=None) -> dict:
    """Run Stage A fact extraction on a single chunk with temperature=0."""
    if client is None:
        client = get_anthropic_client(api_key=api_key)
    model = "claude-haiku-4-5-20251001"

    content = [
        {
            "type": "text",
            "text": STAGE_A_FACT_EXTRACTION_PROMPT + f"\n\nDOCUMENT TO PROCESS ({filename}):\n" + chunk_text
        }
    ]

    response = client.messages.create(
        model=model,
        max_tokens=8000,
        temperature=0.0,
        messages=[{"role": "user", "content": content}],
    )

    data, parse_status = _safe_parse_json_with_status(response.content[0].text)
    if not isinstance(data, dict):
        data = {}
    data["_parse_status"] = parse_status

    for r in data.get("requirements", []):
        r.setdefault("source_refs", [{"source_doc": filename, "page": None, "sheet": None, "section": None, "excerpt": r.get("description", "")[:100]}])
    for d in data.get("dates", []):
        d.setdefault("source_doc", filename)
    for e in data.get("evaluation_criteria", []):
        e.setdefault("source_doc", filename)
    for s in data.get("submission_rules", []):
        s.setdefault("source_doc", filename)
    for d in data.get("deliverables", []):
        d.setdefault("source_doc", filename)
    for c in data.get("commercial_clauses", []):
        c.setdefault("source_doc", filename)
    for cr in data.get("contract_risks", []):
        cr.setdefault("source_doc", filename)

    return data


# ── STAGE A: DOCUMENT FACT EXTRACTION ─────────────────────────────────────────

def extract_document_facts(doc_text: str, filename: str, api_key: str) -> dict:
    """
    STAGE A: Process document to extract factual procurement data ONLY.
    Uses deterministic marker-aware chunking, conservative aggregation,
    and a coverage guard with bounded recovery.
    """
    client = get_anthropic_client(api_key=api_key)
    chunks = chunk_document_text(doc_text, max_chunk_chars=_STAGE_A_MAX_CHUNK_CHARS)

    chunk_results = []
    has_recovered_truncation = False
    has_parse_failure = False

    for chunk in chunks:
        res = _extract_chunk_facts(chunk, filename, api_key, client=client)
        pstatus = res.get("_parse_status")
        if pstatus == "RECOVERED_TRUNCATED":
            target_size = max(500, len(chunk) // 2) if len(chunk) > 1000 else len(chunk)
            retry_subchunks = chunk_document_text(chunk, max_chunk_chars=target_size)
            if len(retry_subchunks) <= 1:
                half_pt = len(chunk) // 2
                split_candidates = [chunk[:half_pt], chunk[half_pt:]] if len(chunk) > 1000 else [chunk]
                retry_subchunks = [c for c in split_candidates if c.strip()]

            sub_results = []
            sub_all_complete = True
            for sc in retry_subchunks:
                sres = _extract_chunk_facts(sc, filename, api_key, client=client)
                if sres.get("_parse_status") != "COMPLETE":
                    sub_all_complete = False
                sub_results.append(sres)
            if sub_all_complete:
                res = aggregate_stage_a_facts(sub_results, filename)
                res["_parse_status"] = "COMPLETE"
            else:
                res = aggregate_stage_a_facts([res] + sub_results, filename)
                res["_parse_status"] = "RECOVERED_TRUNCATED"
                has_recovered_truncation = True
        elif pstatus == "FAILED":
            # Directive 4: Bounded retry for FAILED chunks using smaller subchunks
            target_size = max(500, len(chunk) // 2) if len(chunk) > 1000 else len(chunk)
            retry_subchunks = chunk_document_text(chunk, max_chunk_chars=target_size)
            if len(retry_subchunks) <= 1:
                half_pt = len(chunk) // 2
                split_candidates = [chunk[:half_pt], chunk[half_pt:]] if len(chunk) > 1000 else [chunk]
                retry_subchunks = [c for c in split_candidates if c.strip()]

            sub_results = []
            sub_all_complete = True
            for sc in retry_subchunks:
                sres = _extract_chunk_facts(sc, filename, api_key, client=client)
                if sres.get("_parse_status") != "COMPLETE":
                    sub_all_complete = False
                sub_results.append(sres)
            if sub_all_complete and sub_results:
                res = aggregate_stage_a_facts(sub_results, filename)
                res["_parse_status"] = "COMPLETE"
            else:
                # Retain failure and preserve any recovered partial items
                res = aggregate_stage_a_facts([res] + sub_results, filename)
                res["_parse_status"] = "FAILED"
                has_parse_failure = True
        chunk_results.append(res)

    aggregated = aggregate_stage_a_facts(chunk_results, filename)

    # Coverage Guard inspection
    coverage = inspect_stage_a_coverage(doc_text, aggregated, min_signal_count=5)
    recovery_attempts = 0
    max_recoveries = 2

    # If suspicious and only 1 chunk was extracted, subdivide with smaller chunks for recovery
    while coverage["is_suspicious"] and recovery_attempts < max_recoveries:
        recovery_attempts += 1
        smaller_chunks = chunk_document_text(doc_text, max_chunk_chars=8000)
        if len(smaller_chunks) > len(chunks):
            retry_results = []
            for schunk in smaller_chunks:
                rres = _extract_chunk_facts(schunk, filename, api_key, client=client)
                if rres.get("_parse_status") == "RECOVERED_TRUNCATED":
                    has_recovered_truncation = True
                elif rres.get("_parse_status") == "FAILED":
                    has_parse_failure = True
                retry_results.append(rres)
            retry_aggregated = aggregate_stage_a_facts(retry_results, filename)
            merged_recovery = aggregate_stage_a_facts([aggregated, retry_aggregated], filename)
            merged_cov = inspect_stage_a_coverage(doc_text, merged_recovery, min_signal_count=5)
            aggregated = merged_recovery
            coverage = merged_cov
            if not coverage["is_suspicious"]:
                break
        else:
            break

    # Directive 4: Document diagnostic handling
    if has_parse_failure:
        aggregated["_extraction_diagnostic"] = {
            "status": "PARSE_FAILURE",
            "recovery_attempts": recovery_attempts,
            "has_parse_failure": True,
            "has_recovered_truncation": has_recovered_truncation,
            "note": "Document facts extraction encountered unrecovered chunk parse failure",
        }
    elif coverage["is_suspicious"]:
        aggregated["_extraction_diagnostic"] = {
            "status": "SUSPICIOUS_UNDER_COVERAGE",
            "suspicious_families": coverage["suspicious_families"],
            "signal_counts": coverage["signal_counts"],
            "recovery_attempts": recovery_attempts,
            "has_recovered_truncation": has_recovered_truncation,
        }
    elif has_recovered_truncation:
        aggregated["_extraction_diagnostic"] = {
            "status": "RECOVERED_TRUNCATED",
            "recovery_attempts": recovery_attempts,
            "has_recovered_truncation": True,
            "note": "Document facts recovered from truncated JSON parser; bounded retries completed without complete parse",
        }
    else:
        aggregated["_extraction_diagnostic"] = {
            "status": "VERIFIED_ADEQUATE",
            "recovery_attempts": recovery_attempts,
        }

    return aggregated


# ── STAGE B: PACKAGE NORMALIZATION ───────────────────────────────────────────

def normalize_package_facts(doc_facts_list: list[dict], package_metadata: dict) -> dict:
    """
    STAGE B: Combine document-level facts, deduplicate identical items,
    merge source references, and validate source provenance.
    """
    normalized = {
        "doc_metadata": {},
        "requirements": [],
        "dates": [],
        "evaluation_criteria": [],
        "submission_rules": [],
        "deliverables": [],
        "commercial_clauses": [],
        "contract_risks": []
        ,"_canonical_opportunity": None
    }

    req_seen = {}
    for df in doc_facts_list:
        # Merge metadata (first non-empty)
        meta = df.get("doc_metadata", {})
        for k, v in meta.items():
            if v and not normalized["doc_metadata"].get(k):
                normalized["doc_metadata"][k] = v

        # Normalize and deduplicate requirements
        for r in df.get("requirements", []):
            if not isinstance(r, dict):
                continue
            raw_desc = r.get("description")
            if not isinstance(raw_desc, str):
                continue
            desc = raw_desc.strip()
            if not desc:
                continue
            desc_key = re.sub(r'\W+', '', desc.lower())
            if not desc_key:
                continue
            # Validate source references
            validated_refs = validate_source_refs(r.get("source_refs", []), package_metadata)
            
            explicit = normalize_requirement_type(r.get("requirement_type"))
            candidate = explicit or resolve_requirement_type(r)

            incoming_candidates = set()
            if isinstance(r.get("_semantic_candidates"), (list, set, tuple)):
                for c in r["_semantic_candidates"]:
                    norm_c = normalize_requirement_type(c)
                    if norm_c in SPECIFIC_REQUIREMENT_TYPES:
                        incoming_candidates.add(norm_c)
            if candidate in SPECIFIC_REQUIREMENT_TYPES:
                incoming_candidates.add(candidate)

            if desc_key in req_seen:
                # Merge source references into existing requirement
                existing_r = req_seen[desc_key]
                existing_refs = existing_r.get("source_refs", [])
                for vref in validated_refs:
                    if not any(e.get("source_doc") == vref.get("source_doc") and e.get("page") == vref.get("page") for e in existing_refs):
                        existing_refs.append(vref)
                existing_r["source_refs"] = existing_refs
                # True N-way order-independent candidate accumulation across all documents
                spec_set = set(existing_r.get("_semantic_candidates") or [])
                spec_set.update(incoming_candidates)
                existing_r["_semantic_candidates"] = sorted(spec_set)
                existing_r["requirement_type"] = resolve_candidate_requirement_types(spec_set)
            else:
                r_copy = dict(r)
                r_copy["source_refs"] = validated_refs
                r_copy.setdefault("qual_status", "UNKNOWN")
                r_copy.setdefault("evidence_status", "MISSING")
                r_copy["_semantic_candidates"] = sorted(incoming_candidates)
                if incoming_candidates:
                    r_copy["requirement_type"] = resolve_candidate_requirement_types(incoming_candidates)
                else:
                    r_copy["requirement_type"] = resolve_requirement_type(r)
                req_seen[desc_key] = r_copy
                normalized["requirements"].append(r_copy)

        # Merge other facts
        normalized["dates"].extend(df.get("dates", []))
        raw_pkg_eval = normalized.setdefault("_raw_evaluation_criteria", [])
        raw_pkg_eval.extend(df.get("evaluation_criteria", []))
        normalized["deliverables"].extend(df.get("deliverables", []))
        normalized["commercial_clauses"].extend(df.get("commercial_clauses", []))
        normalized["contract_risks"].extend(df.get("contract_risks", []))

        # Deduplicate and merge submission_rules across package documents
        sub_seen = {}  # canon -> rule dict in normalized["submission_rules"]
        for ex in normalized["submission_rules"]:
            c = _canonical_submission_item_identity(ex.get("item", "")) or ex.get("item", "").lower()
            sub_seen[c] = ex

        for sr in df.get("submission_rules", []):
            if not isinstance(sr, dict):
                continue
            item = (sr.get("item") or "").strip()
            fmt = str(sr.get("format") or "").strip()
            details = str(sr.get("details") or "").strip()
            if not item:
                continue
            canon = _canonical_submission_item_identity(item) or item.lower()
            
            raw_srefs = sr.get("source_refs") or []
            validated_srefs = validate_source_refs(raw_srefs, package_metadata) if raw_srefs else []

            incoming_at_obs = sr.get("artifact_type_observations") or ([sr.get("artifact_type")] if sr.get("artifact_type") else [])
            incoming_ff_obs = sr.get("file_format_observations") or ([sr.get("file_format")] if sr.get("file_format") else [])
            incoming_sc_obs = sr.get("submission_channel_observations") or ([sr.get("submission_channel")] if sr.get("submission_channel") else [])
            incoming_mand_obs = sr.get("mandatory_observations") or ([sr.get("mandatory")] if sr.get("mandatory") is not None else [])

            if canon in sub_seen:
                existing = sub_seen[canon]
                existing_sref_keys = {(s.get("source_doc"), s.get("page"), s.get("sheet"), s.get("section"), s.get("excerpt")) for s in existing.get("source_refs", []) if isinstance(s, dict)}
                for s in validated_srefs:
                    k = (s.get("source_doc"), s.get("page"), s.get("sheet"), s.get("section"), s.get("excerpt"))
                    if k not in existing_sref_keys:
                        existing.setdefault("source_refs", []).append(s)
                        existing_sref_keys.add(k)
                
                # Merge observations and update conflicts
                at_obs = existing.setdefault("artifact_type_observations", [])
                for obs in incoming_at_obs:
                    if obs and obs not in at_obs:
                        at_obs.append(obs)
                valid_at = [x for x in at_obs if x and x != "Unknown"]
                is_at_conflict = len(set(valid_at)) > 1
                existing["artifact_type_conflict"] = is_at_conflict
                if is_at_conflict:
                    existing["artifact_type"] = "Unknown"
                elif valid_at and (not existing.get("artifact_type") or existing.get("artifact_type") == "Unknown"):
                    existing["artifact_type"] = valid_at[0]

                ff_obs = existing.setdefault("file_format_observations", [])
                for obs in incoming_ff_obs:
                    if obs and obs not in ff_obs:
                        ff_obs.append(obs)
                valid_ff = [x for x in ff_obs if x and x != "Unspecified"]
                is_ff_conflict = len(set(valid_ff)) > 1
                existing["file_format_conflict"] = is_ff_conflict
                if is_ff_conflict:
                    existing["file_format"] = "Multiple / Mixed"
                elif valid_ff and (not existing.get("file_format") or existing.get("file_format") == "Unspecified"):
                    existing["file_format"] = valid_ff[0]

                sc_obs = existing.setdefault("submission_channel_observations", [])
                for obs in incoming_sc_obs:
                    if obs and obs not in sc_obs:
                        sc_obs.append(obs)
                valid_sc = [x for x in sc_obs if x and x != "Unspecified"]
                is_sc_conflict = len(set(valid_sc)) > 1
                existing["submission_channel_conflict"] = is_sc_conflict
                if is_sc_conflict:
                    existing["submission_channel"] = "Unspecified"
                elif valid_sc and (not existing.get("submission_channel") or existing.get("submission_channel") == "Unspecified"):
                    existing["submission_channel"] = valid_sc[0]

                mand_obs = existing.setdefault("mandatory_observations", [])
                for obs in incoming_mand_obs:
                    if obs is not None and obs not in mand_obs:
                        mand_obs.append(obs)
                valid_mand = [x for x in mand_obs if x is not None]
                is_mand_conflict = len(set(valid_mand)) > 1
                existing["mandatory_conflict"] = is_mand_conflict
                if is_mand_conflict:
                    existing["mandatory"] = None
                elif valid_mand and existing.get("mandatory") is None:
                    existing["mandatory"] = valid_mand[0]

                if not existing.get("format") and fmt:
                    existing["format"] = fmt
                if not existing.get("details") and details:
                    existing["details"] = details
            else:
                sr_copy = dict(sr)
                sr_copy["source_refs"] = validated_srefs
                sr_copy["artifact_type_observations"] = list(incoming_at_obs)
                valid_at = [x for x in incoming_at_obs if x and x != "Unknown"]
                is_at_conflict = len(set(valid_at)) > 1
                sr_copy["artifact_type_conflict"] = is_at_conflict
                if is_at_conflict:
                    sr_copy["artifact_type"] = "Unknown"

                sr_copy["file_format_observations"] = list(incoming_ff_obs)
                valid_ff = [x for x in incoming_ff_obs if x and x != "Unspecified"]
                is_ff_conflict = len(set(valid_ff)) > 1
                sr_copy["file_format_conflict"] = is_ff_conflict
                if is_ff_conflict:
                    sr_copy["file_format"] = "Multiple / Mixed"

                sr_copy["submission_channel_observations"] = list(incoming_sc_obs)
                valid_sc = [x for x in incoming_sc_obs if x and x != "Unspecified"]
                is_sc_conflict = len(set(valid_sc)) > 1
                sr_copy["submission_channel_conflict"] = is_sc_conflict
                if is_sc_conflict:
                    sr_copy["submission_channel"] = "Unspecified"

                sr_copy["mandatory_observations"] = list(incoming_mand_obs)
                valid_mand = [x for x in incoming_mand_obs if x is not None]
                is_mand_conflict = len(set(valid_mand)) > 1
                sr_copy["mandatory_conflict"] = is_mand_conflict
                if is_mand_conflict:
                    sr_copy["mandatory"] = None

                sub_seen[canon] = sr_copy
                normalized["submission_rules"].append(sr_copy)

    # Clean up internal candidate accumulation tracking from final normalized requirements
    for r in normalized["requirements"]:
        r.pop("_semantic_candidates", None)
        r.pop("_specific_types", None)

    # Deterministically deduplicate package-level evaluation criteria
    from evaluation_hierarchy import deduplicate_evaluation_criteria
    raw_pkg_eval = normalized.pop("_raw_evaluation_criteria", [])
    normalized["evaluation_criteria"] = deduplicate_evaluation_criteria(raw_pkg_eval)

    from canonical_opportunity import build_canonical_opportunity
    normalized["_canonical_opportunity"] = build_canonical_opportunity(doc_facts_list, package_metadata)

    return normalized


# ── STAGE C: RECONCILIATION & CONFLICT ANALYSIS ───────────────────────────────

def reconcile_package_facts(normalized_facts: dict, package_files: list[str]) -> list[dict]:
    """
    STAGE C: Compare normalized facts to detect cross-document conflicts & addenda overrides.
    """
    conflicts = detect_document_conflicts(normalized_facts, package_files)
    canonical = normalized_facts.get("_canonical_opportunity")
    if canonical:
        from canonical_opportunity import resolve_canonical_opportunity
        canonical = resolve_canonical_opportunity(canonical)
        normalized_facts["_canonical_opportunity"] = canonical
        conflicts.extend(canonical.get("conflicts", []))
    return conflicts


# -- SUBMISSION DOCUMENT PROJECTION -----------------------------------------
#
# PURPOSE
# -------
# The documents table drives a file / submission-package checklist that can
# become a Stage-Submit readiness blocker. Only project a document record when
# normalized evidence establishes a DISCRETE, INDEPENDENTLY TRACKED submission
# artefact or package item -- something a bid manager must locate, upload, or
# sign off as a separate deliverable.
#
# DO NOT project:
#   - fields embedded inside another form/response
#   - workbook tabs that are part of a parent workbook
#   - Yes/No response columns
#   - page / format / font constraints
#   - portal / process / delivery-mode instructions
#   - reference-only material
#   - external-link restrictions
#   - pricing rules (not independent files)
#
# PREFER OMISSION OVER INVENTION.
# A rule that is not projected remains visible in Stage D / UNDERSTAND
# via the submission_rules list.

import re as _re

# ---------------------------------------------------------------------------
# Format-field signals (examined before indicator-word scanning)
# ---------------------------------------------------------------------------

# Format phrases indicating the item is NOT an independent file.
# Covers embedded content, form entries, checklist columns, portal containers,
# and mixed/ambiguous formats (e.g. "Embedded or Separate File").
_FORMAT_NEGATIVE_PHRASES = frozenset([
    'form entry',
    'integrated',
    'embedded',
    'proponent response column',
    'yes/no confirmation',
    'yes / no confirmation',
    'yes/no format',
    'not permitted',
    'reference only',
    'for reference',
    'available for reference',
])

# Signals indicating physical file format
_FILE_FORMAT_SIGNALS = {
    'excel spreadsheet': 'Spreadsheet',
    'excel workbook': 'Spreadsheet',
    'spreadsheet': 'Spreadsheet',
    'online form': 'Online Form',
    'hard copy': 'Hard Copy',
    'excel': 'Spreadsheet',
    'docx': 'DOCX',
    'xlsx': 'XLSX',
    'pdf': 'PDF',
    'xls': 'XLS',
}

# Signals indicating packaging or standalone independence (NOT physical format)
_ARTIFACT_INDEPENDENCE_SIGNALS = frozenset([
    'separate file',
    'separate files',
    'separate submission',
    'separate document',
    'single document',
    'standalone document',
    'attachment',
    'attachments',
    'attached file',
    'attached files',
    'attached samples',
])

# Controlled channel mapping
_CONTROLLED_CHANNEL_MAP = {
    'electronic bid submission': 'E-procurement',
    'electronic submission': 'E-procurement',
    'portal': 'Portal',
    'merx': 'Portal',
    'buyandsell': 'Portal',
    'email': 'Email',
    'e-procurement': 'E-procurement',
    'eprocurement': 'E-procurement',
    'courier': 'Courier',
    'physical delivery': 'Physical',
    'physical': 'Physical',
    'upload': 'Upload',
    'upload file': 'Upload',
    'upload proposal': 'Upload',
}

# Generic document formats that qualify as independent when paired with
# a strong standalone submission artifact noun in the item title.
_GENERIC_DOCUMENT_FORMAT_PHRASES = frozenset([
    'document',
    'documents',
])

# ---------------------------------------------------------------------------
# Item / details exclusion signals
# ---------------------------------------------------------------------------

# Phrases in item text that indicate a process or format instruction.
_PROCESS_ITEM_PHRASES = frozenset([
    'submit through', 'submit via', 'submitted through', 'submitted via',
    'submit using', 'submitted using',
    'electronic submission only', 'portal only', 'online submission',
    'registration required', 'register on',
    'no external links', 'external links',
    'page limit', 'maximum pages', 'not exceed', 'page count',
    'font size', 'margin',
    'maximum response length', 'maximum length',
    'response length', 'response limit',
    'portal submission', 'online portal',
    'submission portal',
    'currency and tax',
    'assumptions and restrictions',
    'response format',
])

# Phrases in details text that reveal reference-only / process intent.
_PROCESS_DETAILS_PHRASES = frozenset([
    'not mandatory but available for reference',
    'available for reference',
    'for reference only',
    'reference only',
    'will not be evaluated',
    'external to the form will not',
    'links to websites',
    'links to external',
])

# Regex for quantified page/word/item constraints in item text.
_QUANTITY_CONSTRAINT_RE = _re.compile(
    r'(?:maximum|max|no more than|not exceed|limit of?)\s+\d+\s*'
    r'(?:pages?|words?|lines?|items?)',
    _re.IGNORECASE,
)

# Regex for workbook tabs or worksheets that are part of another submitted workbook.
_WORKBOOK_TAB_RE = _re.compile(
    r"(?i)\b(?:tabs?|worksheets?)\b"
)

# Strong standalone artifact nouns with morphology (singular / plural alternatives).
# These represent discrete submission packages/forms even without explicit file format,
# provided format is generic document, empty, or standalone (and not negative).
_STRONG_STANDALONE_ARTIFACT_RE = _re.compile(
    r"(?i)\b(?:"
    r"forms?|proposals?|questionnaires?|templates?|annex(?:es)?|schedules?|"
    r"pricing\s+forms?|rate\s+cards?|work\s*plans?|accounts?|checklists?"
    r")\b"
)

def _canonical_submission_item_identity(item: str) -> str:
    """
    Extract a stable canonical identifier for a submission item.
    Normalizes official identifiers (Annex, Form, Schedule, Envelope) and punctuation
    so variations in format/details text do not split the same logical artifact.
    """
    if not item:
        return ""
    text = item.strip().lower()
    m = _re.search(r'\b(annex|appendix|schedule|form|envelope)\s+([a-z0-9]+)\b', text)
    if m:
        return f"{m.group(1)} {m.group(2)}"
    return _re.sub(r'[\W_]+', ' ', text).strip()


def classify_submission_rule(
    item: str,
    fmt: str | None = None,
    details: str | None = None,
    artifact_type: str | None = None,
    file_format: str | None = None,
    submission_channel: str | None = None,
    **kwargs
) -> dict:
    '''
    Classify a submission rule across the three orthogonal dimensions:
      1. Artifact Identity (What is to be submitted)
      2. File Format (Physical representation: PDF, DOCX, XLSX, etc.)
      3. Submission Channel (Delivery mechanism: Portal, Email, Courier, etc.)

    Returns a dict containing:
      - is_concrete_document: bool
      - classification: 'ARTIFACT' | 'PROCESS_ONLY' | 'EMBEDDED_RESPONSE' | 'UNKNOWN'
      - artifact_type: str
      - file_format: str
      - submission_channel: str
    '''
    item_clean = (item or '').strip()
    fmt_raw = (fmt or kwargs.get('format') or '').strip()
    details_clean = (details or '').strip()

    if not item_clean:
        return {
            'is_concrete_document': False,
            'classification': 'UNKNOWN',
            'artifact_type': artifact_type or 'Unknown',
            'file_format': file_format or 'Unspecified',
            'submission_channel': submission_channel or 'Unspecified',
        }

    item_lower = item_clean.lower()
    fmt_lower = fmt_raw.lower()
    details_lower = details_clean.lower()

    # Determine submission channel from controlled mapping or text scan
    resolved_channel = 'Unspecified'
    if submission_channel and submission_channel.strip() and submission_channel.strip() != 'Unspecified':
        raw_sc = submission_channel.strip().lower()
        resolved_channel = _CONTROLLED_CHANNEL_MAP.get(raw_sc, submission_channel.strip().title())
    else:
        # Check fmt and details
        for raw_k, norm_v in sorted(_CONTROLLED_CHANNEL_MAP.items(), key=lambda x: len(x[0]), reverse=True):
            if raw_k in fmt_lower or raw_k in details_lower:
                resolved_channel = norm_v
                break

    # Determine file format (strictly physical format, never 'Separate Document' / 'Separate File')
    resolved_format = 'Unspecified'
    has_independence_signal = False
    if file_format and file_format.strip() and file_format.strip() != 'Unspecified':
        ff_str = file_format.strip()
        ff_lower = ff_str.lower()
        if ff_lower in _ARTIFACT_INDEPENDENCE_SIGNALS:
            has_independence_signal = True
            resolved_format = 'Unspecified'
        elif ff_lower in _FILE_FORMAT_SIGNALS:
            resolved_format = _FILE_FORMAT_SIGNALS[ff_lower]
        else:
            resolved_format = ff_str
    
    if resolved_format == 'Unspecified' and fmt_raw:
        # Scan for physical formats first (descending length so 'xlsx' matches before 'xls')
        for sig, norm_val in sorted(_FILE_FORMAT_SIGNALS.items(), key=lambda x: len(x[0]), reverse=True):
            if sig in fmt_lower:
                resolved_format = norm_val
                break
        for ind in _ARTIFACT_INDEPENDENCE_SIGNALS:
            if ind in fmt_lower:
                has_independence_signal = True
                break

    # Determine artifact type
    resolved_artifact_type = artifact_type or 'Unknown'
    if resolved_artifact_type == 'Unknown':
        if any(kw in item_lower for kw in ('pricing', 'financial', 'rate card', 'cost')):
            resolved_artifact_type = 'Pricing / Financial'
        elif any(kw in item_lower for kw in ('proposal', 'response')):
            resolved_artifact_type = 'Proposal / Response'
        elif any(kw in item_lower for kw in ('questionnaire',)):
            resolved_artifact_type = 'Questionnaire / Workbook'
        elif any(kw in item_lower for kw in ('annex', 'form', 'template', 'schedule', 'checklist')):
            resolved_artifact_type = 'Form / Annex'
        elif any(kw in item_lower for kw in ('declaration', 'certif')):
            resolved_artifact_type = 'Declaration / Certification'
        elif any(kw in item_lower for kw in ('sample', 'evidence', 'accounts', 'attachment')):
            resolved_artifact_type = 'Evidence / Attachment'

    # Hard negative format signals (embedded content, form entry, etc.)
    if any(p in fmt_lower for p in _FORMAT_NEGATIVE_PHRASES):
        return {
            'is_concrete_document': False,
            'classification': 'EMBEDDED_RESPONSE',
            'artifact_type': resolved_artifact_type,
            'file_format': resolved_format if resolved_format != 'Unspecified' else 'Online Form',
            'submission_channel': resolved_channel,
        }

    # Workbook tab inside another workbook
    if _WORKBOOK_TAB_RE.search(item_clean):
        return {
            'is_concrete_document': False,
            'classification': 'EMBEDDED_RESPONSE',
            'artifact_type': resolved_artifact_type,
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Process / portal / formatting instruction in item text or envelope container
    is_envelope_container = bool(_re.search(r'(?i)^\s*envelope\s+\d+\b', item_clean))
    is_pure_bid_submission_header = (item_lower == 'electronic bid submission')

    if is_envelope_container or is_pure_bid_submission_header or any(p in item_lower for p in _PROCESS_ITEM_PHRASES):
        return {
            'is_concrete_document': False,
            'classification': 'PROCESS_ONLY',
            'artifact_type': 'Process Instruction',
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Quantified constraint in item text
    if _QUANTITY_CONSTRAINT_RE.search(item_clean):
        return {
            'is_concrete_document': False,
            'classification': 'PROCESS_ONLY',
            'artifact_type': 'Process Instruction',
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Details reveal reference-only or non-evaluated intent
    if any(p in details_lower for p in _PROCESS_DETAILS_PHRASES):
        return {
            'is_concrete_document': False,
            'classification': 'PROCESS_ONLY',
            'artifact_type': 'Process Instruction',
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Explicit normalized orthogonal inputs
    if artifact_type == 'Process Instruction':
        return {
            'is_concrete_document': False,
            'classification': 'PROCESS_ONLY',
            'artifact_type': 'Process Instruction',
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    if resolved_format == 'Online Form' and resolved_artifact_type in ('Questionnaire / Workbook', 'Form / Annex') and not has_independence_signal:
        return {
            'is_concrete_document': False,
            'classification': 'EMBEDDED_RESPONSE',
            'artifact_type': resolved_artifact_type,
            'file_format': 'Online Form',
            'submission_channel': resolved_channel,
        }

    has_explicit_artifact_type = bool(artifact_type and artifact_type.strip() and artifact_type.strip() != 'Unknown')
    if has_explicit_artifact_type and resolved_artifact_type in ('Proposal / Response', 'Declaration / Certification', 'Evidence / Attachment', 'Pricing / Financial', 'Form / Annex') and resolved_format != 'Online Form':
        return {
            'is_concrete_document': True,
            'classification': 'ARTIFACT',
            'artifact_type': resolved_artifact_type,
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Explicit independent file/package format signal
    if has_independence_signal or resolved_format != 'Unspecified':
        return {
            'is_concrete_document': True,
            'classification': 'ARTIFACT',
            'artifact_type': resolved_artifact_type,
            'file_format': resolved_format,
            'submission_channel': resolved_channel,
        }

    # Normalize residual format after stripping channel phrases
    residual_fmt = fmt_lower
    for ch_phrase in ('electronic portal or email', 'portal/email', 'electronic portal', 'online portal', 'electronic submission', 'electronic bid submission', 'portal', 'email', 'merx', 'buyandsell', 'upload'):
        residual_fmt = residual_fmt.replace(ch_phrase, '')
    residual_fmt = residual_fmt.strip(' /,-')

    # 7. Residual format is empty, generic document, or unspecified AND item title has strong standalone artifact noun
    if (not residual_fmt or residual_fmt in _GENERIC_DOCUMENT_FORMAT_PHRASES):
        if _STRONG_STANDALONE_ARTIFACT_RE.search(item_clean):
            return {
                'is_concrete_document': True,
                'classification': 'ARTIFACT',
                'artifact_type': resolved_artifact_type,
                'file_format': resolved_format,
                'submission_channel': resolved_channel,
            }

    # 8. Fallback
    return {
        'is_concrete_document': False,
        'classification': 'PROCESS_ONLY' if (resolved_channel != 'Unspecified' or not fmt_raw) else 'UNKNOWN',
        'artifact_type': resolved_artifact_type,
        'file_format': resolved_format,
        'submission_channel': resolved_channel,
    }


def _is_concrete_submission_document(
    item: str,
    fmt: str | None = None,
    details: str | None = None,
) -> bool:
    '''
    Backward-compatible predicate: Return True only when normalized evidence establishes
    the rule as a DISCRETE, INDEPENDENTLY TRACKED submission artefact.
    '''
    return classify_submission_rule(item, fmt, details)['is_concrete_document']


def build_submission_documents(
    submission_rules: list[dict],
    submission_deadline: str | None = None,
) -> list[dict]:
    '''
    SUBMISSION DOCUMENT PROJECTION -- deterministic, provenance-preserving.

    Projects normalized submission rules into document records for the bid
    registry. Only rules representing a DISCRETE, INDEPENDENTLY TRACKED
    submission artefact or package item are included.

    Deduplicates identical artifacts while merging source_refs.

    mandatory semantics:
      1   -> REQUIRED
      0   -> OPTIONAL
      absent/None -> UNKNOWN (key omitted from record)
    '''
    documents = []
    seen_docs = {}  # name.lower() -> index in documents

    for sr in (submission_rules or []):
        if not isinstance(sr, dict):
            continue
        item = (sr.get('item') or '').strip()
        fmt = sr.get('format') or ''
        details = sr.get('details') or ''
        
        classification = classify_submission_rule(
            item=item,
            fmt=fmt,
            details=details,
            artifact_type=sr.get('artifact_type'),
            file_format=sr.get('file_format'),
            submission_channel=sr.get('submission_channel'),
        )
        if not classification['is_concrete_document']:
            continue

        # Determine orthogonal conflict states and observation lists
        at_obs = list(sr.get('artifact_type_observations') or ([sr.get('artifact_type')] if sr.get('artifact_type') else []))
        ff_obs = list(sr.get('file_format_observations') or ([sr.get('file_format')] if sr.get('file_format') else []))
        sc_obs = list(sr.get('submission_channel_observations') or ([sr.get('submission_channel')] if sr.get('submission_channel') else []))
        mand_obs = list(sr.get('mandatory_observations') or ([sr.get('mandatory')] if sr.get('mandatory') is not None else []))

        valid_at_obs = [x for x in at_obs if x and x != 'Unknown']
        at_conflict = bool(sr.get('artifact_type_conflict') or len(set(valid_at_obs)) > 1)

        valid_ff_obs = [x for x in ff_obs if x and x != 'Unspecified']
        ff_conflict = bool(sr.get('file_format_conflict') or len(set(valid_ff_obs)) > 1)

        valid_sc_obs = [x for x in sc_obs if x and x != 'Unspecified']
        sc_conflict = bool(sr.get('submission_channel_conflict') or len(set(valid_sc_obs)) > 1)

        valid_mand_obs = [x for x in mand_obs if x is not None]
        mand_conflict = bool(sr.get('mandatory_conflict') or len(set(valid_mand_obs)) > 1)

        # Conservative conflict resolution for projected metadata
        if at_conflict:
            proj_artifact_type = 'Unknown'
            # doc_type must not be determined from arbitrary conflicting artifact_type
            if any(kw in item.lower() for kw in ('pricing', 'financial', 'rate card', 'cost')):
                doc_type = 'Financial'
            else:
                doc_type = 'Submission'
        elif classification.get('artifact_type') == 'Pricing / Financial':
            proj_artifact_type = 'Pricing / Financial'
            doc_type = 'Financial'
        elif any(kw in item.lower() for kw in ('pricing', 'financial', 'rate card', 'cost')):
            proj_artifact_type = classification['artifact_type']
            doc_type = 'Financial'
        else:
            proj_artifact_type = classification['artifact_type']
            doc_type = 'Submission'

        proj_file_format = 'Multiple / Mixed' if ff_conflict else classification['file_format']
        proj_submission_channel = 'Unspecified' if sc_conflict else classification['submission_channel']

        # mandatory: None if conflict or absent
        if mand_conflict:
            proj_mandatory = None
        else:
            proj_mandatory = sr.get('mandatory')

        name_key = _canonical_submission_item_identity(item) or item.lower()

        # Extract source_refs from submission rule (no manufactured fallbacks)
        srefs = []
        for sref in (sr.get('source_refs') or []):
            if isinstance(sref, dict):
                srefs.append(dict(sref))

        if name_key in seen_docs:
            existing_doc = documents[seen_docs[name_key]]
            existing_srefs = existing_doc.setdefault('source_refs', [])
            existing_keys = {(s.get('source_doc'), s.get('page'), s.get('sheet'), s.get('section'), s.get('excerpt')) for s in existing_srefs if isinstance(s, dict)}
            for s in srefs:
                k = (s.get('source_doc'), s.get('page'), s.get('sheet'), s.get('section'), s.get('excerpt'))
                if k not in existing_keys:
                    existing_srefs.append(s)
                    existing_keys.add(k)
            if existing_srefs:
                existing_doc['provenance_state'] = 'VERIFIED'

            if not existing_doc.get('notes') and details:
                existing_doc['notes'] = details

            # Merge observation history & conflict flags
            ex_at_obs = existing_doc.setdefault('artifact_type_observations', [])
            for obs in at_obs:
                if obs and obs not in ex_at_obs:
                    ex_at_obs.append(obs)
            v_at = [x for x in ex_at_obs if x and x != 'Unknown']
            is_at_c = at_conflict or existing_doc.get('artifact_type_conflict', False) or (len(set(v_at)) > 1)
            existing_doc['artifact_type_conflict'] = is_at_c
            if is_at_c:
                existing_doc['artifact_type'] = 'Unknown'
                if any(kw in existing_doc['name'].lower() for kw in ('pricing', 'financial', 'rate card', 'cost')):
                    existing_doc['doc_type'] = 'Financial'
                else:
                    existing_doc['doc_type'] = 'Submission'
            elif v_at and (not existing_doc.get('artifact_type') or existing_doc.get('artifact_type') == 'Unknown'):
                existing_doc['artifact_type'] = v_at[0]
                if existing_doc['artifact_type'] == 'Pricing / Financial':
                    existing_doc['doc_type'] = 'Financial'

            ex_ff_obs = existing_doc.setdefault('file_format_observations', [])
            for obs in ff_obs:
                if obs and obs not in ex_ff_obs:
                    ex_ff_obs.append(obs)
            v_ff = [x for x in ex_ff_obs if x and x != 'Unspecified']
            is_ff_c = ff_conflict or existing_doc.get('file_format_conflict', False) or (len(set(v_ff)) > 1)
            existing_doc['file_format_conflict'] = is_ff_c
            if is_ff_c:
                existing_doc['file_format'] = 'Multiple / Mixed'
            elif v_ff and (not existing_doc.get('file_format') or existing_doc.get('file_format') == 'Unspecified'):
                existing_doc['file_format'] = v_ff[0]

            ex_sc_obs = existing_doc.setdefault('submission_channel_observations', [])
            for obs in sc_obs:
                if obs and obs not in ex_sc_obs:
                    ex_sc_obs.append(obs)
            v_sc = [x for x in ex_sc_obs if x and x != 'Unspecified']
            is_sc_c = sc_conflict or existing_doc.get('submission_channel_conflict', False) or (len(set(v_sc)) > 1)
            existing_doc['submission_channel_conflict'] = is_sc_c
            if is_sc_c:
                existing_doc['submission_channel'] = 'Unspecified'
            elif v_sc and (not existing_doc.get('submission_channel') or existing_doc.get('submission_channel') == 'Unspecified'):
                existing_doc['submission_channel'] = v_sc[0]

            ex_mand_obs = existing_doc.setdefault('mandatory_observations', [])
            for obs in mand_obs:
                if obs is not None and obs not in ex_mand_obs:
                    ex_mand_obs.append(obs)
            v_mand = [x for x in ex_mand_obs if x is not None]
            is_mand_c = mand_conflict or existing_doc.get('mandatory_conflict', False) or (len(set(v_mand)) > 1)
            existing_doc['mandatory_conflict'] = is_mand_c
            if is_mand_c:
                existing_doc.pop('mandatory', None)
            elif v_mand and existing_doc.get('mandatory') is None:
                existing_doc['mandatory'] = v_mand[0]

            continue

        prov_state = 'VERIFIED' if srefs else 'UNVERIFIED'
        doc: dict = {
            'name':                            item,
            'doc_type':                        doc_type,
            'owner':                           None,
            'due_date':                        submission_deadline,
            'status':                          'Expected',
            'notes':                           details,
            'file_format':                     proj_file_format,
            'submission_channel':              proj_submission_channel,
            'artifact_type':                   proj_artifact_type,
            'source_refs':                     srefs,
            'provenance_state':                prov_state,
            'artifact_type_observations':      at_obs,
            'artifact_type_conflict':          at_conflict,
            'file_format_observations':        ff_obs,
            'file_format_conflict':            ff_conflict,
            'submission_channel_observations': sc_obs,
            'submission_channel_conflict':     sc_conflict,
            'mandatory_observations':          mand_obs,
            'mandatory_conflict':              mand_conflict,
        }
        if proj_mandatory is not None:
            doc['mandatory'] = proj_mandatory
        
        seen_docs[name_key] = len(documents)
        documents.append(doc)

    return documents


# ── STAGE D: BID BRIEF SYNTHESIS ─────────────────────────────────────────────


_REQUIREMENT_FIELDS = (
    "req_id", "category", "requirement_type", "description", "rfso_ref",
    "weight", "evidence", "source_refs", "qual_status", "evidence_status"
)
_SOURCE_REF_FIELDS = ("source_doc", "page", "sheet", "section", "excerpt")
_EXCERPT_CAP = 500  # characters per source_ref excerpt — size control without dropping requirements

# Conservative preflight limit: serialized JSON characters of the full Stage D context.
# Claude Haiku has a 200k-token context window; at ~3 chars/token this gives ~600k chars
# of usable input after subtracting the fixed prompt (~5k chars) and output reservation (8k tokens).
# We use 580_000 chars as a safe conservative limit.
_STAGE_D_CONTEXT_CHAR_LIMIT = 580_000


class StageDContextTooLargeError(RuntimeError):
    """
    Raised when the serialized Stage D context exceeds the safe model input budget.

    This error guarantees:
    - The context is COMPLETE — no requirements were silently omitted.
    - The failure is explicit — no partial synthesis attempt is made.
    """
    pass


def _compact_requirement(r: dict) -> dict:
    """
    Return a compact but complete requirement record preserving all decision-critical fields.
    Excerpts are capped at _EXCERPT_CAP characters; no other truncation is applied.
    Caller must ensure r is a dict (raises if not).
    """
    out = {}
    for f in _REQUIREMENT_FIELDS:
        if f == "source_refs":
            srefs = []
            for sref in (r.get("source_refs") or []):
                if isinstance(sref, dict):
                    compact_sref = {k: sref.get(k) for k in _SOURCE_REF_FIELDS}
                    if compact_sref.get("excerpt") and len(compact_sref["excerpt"]) > _EXCERPT_CAP:
                        compact_sref["excerpt"] = compact_sref["excerpt"][:_EXCERPT_CAP]
                    srefs.append(compact_sref)
            out["source_refs"] = srefs
        else:
            if f in r:
                out[f] = r[f]
    return out


def build_stage_d_context(normalized_facts: dict, conflicts: list[dict]) -> dict:
    """
    STAGE D CONTEXT BUILDER — deterministic, lossless, strictly integrity-checked.

    Produces the complete evidence model that Stage D (AI synthesis) receives.
    Never slices or truncates the requirement list.  All Mandatory, Financial,
    Rated and Supporting requirements are included in full.

    Integrity rules:
    - source_requirement_count is derived from the ORIGINAL list length
      (including any non-dict entries).
    - Non-dict entries in the requirements list raise RuntimeError immediately
      (silent exclusion followed by reporting 0 omissions is forbidden).
    - all_mandatory_included and all_financial_included are computed from
      actual source vs included category counts, not hard-coded True.

    Returns a dict containing context_integrity with completeness audit data.
    """
    raw_reqs = normalized_facts.get("requirements") or []
    source_count = len(raw_reqs)

    # Strict integrity: non-dict entries are an extraction defect — fail explicitly
    for idx, r in enumerate(raw_reqs):
        if not isinstance(r, dict):
            raise RuntimeError(
                f"Stage D context integrity failure: requirements[{idx}] is "
                f"{type(r).__name__!r}, not a dict. "
                "Non-dict requirement entries must not be silently excluded. "
                "Fix the upstream normalization stage that produced this entry."
            )

    mandatory  = [_compact_requirement(r) for r in raw_reqs
                  if r.get("category", "").strip().lower() == "mandatory"]
    financial  = [_compact_requirement(r) for r in raw_reqs
                  if r.get("category", "").strip().lower() == "financial"]
    rated      = [_compact_requirement(r) for r in raw_reqs
                  if r.get("category", "").strip().lower() == "rated"]
    supporting = [_compact_requirement(r) for r in raw_reqs
                  if r.get("category", "").strip().lower() == "supporting"]
    other      = [_compact_requirement(r) for r in raw_reqs
                  if r.get("category", "").strip().lower()
                  not in ("mandatory", "financial", "rated", "supporting")]

    included_count = len(mandatory) + len(financial) + len(rated) + len(supporting) + len(other)
    omitted_count  = source_count - included_count  # must always be 0

    if omitted_count != 0:
        raise RuntimeError(
            f"Stage D context builder omitted {omitted_count} requirements — "
            "silent omission is forbidden. Check category filter logic for bugs."
        )

    from collections import Counter
    cat_counts = Counter(r.get("category", "").strip() for r in raw_reqs)

    # Compute inclusion flags from actual counts (not hard-coded)
    src_mandatory = cat_counts.get("Mandatory", 0) + cat_counts.get("mandatory", 0)
    src_financial = cat_counts.get("Financial", 0) + cat_counts.get("financial", 0)
    all_mandatory_included = (len(mandatory) == src_mandatory)
    all_financial_included = (len(financial) == src_financial)

    context = {
        "metadata": normalized_facts.get("doc_metadata", {}),
        "requirements": {
            "mandatory":  mandatory,
            "financial":  financial,
            "rated":      rated,
            "supporting": supporting,
            **({"other": other} if other else {}),
        },
        "dates":               normalized_facts.get("dates", []),
        "evaluation_criteria": normalized_facts.get("evaluation_criteria", []),
        "submission_rules":    normalized_facts.get("submission_rules", []),
        "deliverables":        normalized_facts.get("deliverables", []),
        "commercial_clauses":  normalized_facts.get("commercial_clauses", []),
        "contract_risks":      normalized_facts.get("contract_risks", []),
        "detected_conflicts":  list(conflicts),
        "context_integrity": {
            "source_requirement_count":    source_count,
            "included_requirement_count":  included_count,
            "counts_by_category":          dict(cat_counts),
            "omitted_requirement_count":   omitted_count,
            "all_mandatory_included":      all_mandatory_included,
            "all_financial_included":      all_financial_included,
        },
    }
    return context


def apply_stage_d_authoritative_sections(synth_data: dict, normalized_facts: dict) -> dict:
    """
    POST-SYNTHESIS AUTHORITATIVE SECTION APPLICATOR.

    After Stage D AI returns its synthesis, this function deterministically
    rebuilds evidence-backed structured sections from the normalized facts so
    that AI omission cannot destroy procurement-critical data.

    Deduplication uses material canonical tuple keys so that two entries with
    the same title/stage/milestone but different data values are BOTH preserved.
    Only exact duplicates (identical on all canonical fields) are collapsed.

    Sections replaced deterministically:
      qualification_gates      <- Mandatory Supplier Qualification requirements only
      evaluation_breakdown     <- normalized evaluation_criteria
      submission_requirements  <- normalized submission_rules
      key_dates                <- normalized dates
      commercial_structure     <- normalized commercial_clauses
      contract_risks           <- normalized contract_risks
      deliverables_summary     <- normalized deliverables

    No default values are invented. Fields absent in normalized data are set to None.

    AI-synthesized interpretation fields are preserved unchanged:
      executive_summary, opportunity_type, contract_term, procurement_model,
      scope_categories, source_citations, bid fields, outline.
    """
    # Robust brief coercion: handle None, [], "", non-dict
    if not isinstance(synth_data, dict):
        synth_data = {}
    raw_brief = synth_data.get("brief")
    brief = dict(raw_brief) if isinstance(raw_brief, dict) else {}

    all_reqs = [r for r in (normalized_facts.get("requirements") or []) if isinstance(r, dict)]

    # ── qualification_gates — ONLY Mandatory Supplier Qualification requirements ───
    # Canonical key: (description, rfso_ref, req_id) — preserves same description
    # from different RFSO refs or different doc versions.
    qual_reqs = [r for r in all_reqs if is_supplier_qualification(r)]
    seen_gates: set[tuple] = set()
    gates = []
    for r in qual_reqs:
        desc    = (r.get("description") or "").strip()
        rfso    = r.get("rfso_ref")
        req_id  = r.get("req_id")
        canon   = (desc, rfso, req_id)
        if not desc or canon in seen_gates:
            continue
        seen_gates.add(canon)
        ref_parts = [
            sref["source_doc"]
            for sref in (r.get("source_refs") or [])
            if isinstance(sref, dict) and sref.get("source_doc")
        ]
        rfp_ref = rfso or (", ".join(ref_parts) if ref_parts else None)
        gates.append({
            "requirement":          desc,
            "type":                 "Supplier Qualification",
            "rfp_ref":              rfp_ref,
            "req_id":               req_id,
            "disqualification_risk":"High",
        })
    brief["qualification_gates"] = gates

    # ── evaluation_breakdown — normalized evaluation_criteria ─────────────────
    # Uses evaluation_hierarchy to preserve full hierarchy while maintaining backward-compatible keys
    from evaluation_hierarchy import deduplicate_evaluation_criteria
    eval_criteria = normalized_facts.get("evaluation_criteria", [])
    brief["evaluation_breakdown"] = deduplicate_evaluation_criteria(eval_criteria)

    # ── submission_requirements — normalized submission_rules ─────────────────
    # Canonical key: (item, format, details, mandatory, first source_doc)
    sub_rules = normalized_facts.get("submission_rules", [])
    seen_sub: set[tuple] = set()
    sub_reqs = []
    for sr in sub_rules:
        if not isinstance(sr, dict):
            continue
        item    = (sr.get("item") or "").strip()
        fmt     = sr.get("format")
        details = sr.get("details")
        mand    = sr.get("mandatory")
        src_doc = None
        for sref in (sr.get("source_refs") or []):
            if isinstance(sref, dict) and sref.get("source_doc"):
                src_doc = sref["source_doc"]
                break
        canon = (item, str(fmt), str(details), str(mand), src_doc)
        if not item or canon in seen_sub:
            continue
        seen_sub.add(canon)
        entry = {
            "item":    item,
            "format":  fmt,
            "details": details,
        }
        if mand is not None:
            entry["mandatory"] = mand
        if sr.get("artifact_type") is not None:
            entry["artifact_type"] = sr["artifact_type"]
        if sr.get("file_format") is not None:
            entry["file_format"] = sr["file_format"]
        if sr.get("submission_channel") is not None:
            entry["submission_channel"] = sr["submission_channel"]
        if sr.get("source_refs") is not None:
            entry["source_refs"] = sr["source_refs"]
        if sr.get("artifact_type_observations") is not None:
            entry["artifact_type_observations"] = sr["artifact_type_observations"]
        if sr.get("artifact_type_conflict") is not None:
            entry["artifact_type_conflict"] = sr["artifact_type_conflict"]
        if sr.get("file_format_observations") is not None:
            entry["file_format_observations"] = sr["file_format_observations"]
        if sr.get("file_format_conflict") is not None:
            entry["file_format_conflict"] = sr["file_format_conflict"]
        if sr.get("submission_channel_observations") is not None:
            entry["submission_channel_observations"] = sr["submission_channel_observations"]
        if sr.get("submission_channel_conflict") is not None:
            entry["submission_channel_conflict"] = sr["submission_channel_conflict"]
        if sr.get("mandatory_observations") is not None:
            entry["mandatory_observations"] = sr["mandatory_observations"]
        if sr.get("mandatory_conflict") is not None:
            entry["mandatory_conflict"] = sr["mandatory_conflict"]
        sub_reqs.append(entry)
    brief["submission_requirements"] = sub_reqs

    # ── key_dates — normalized dates ──────────────────────────────────────────
    # Canonical key: (milestone, date, first source_doc)
    # Two dates with same milestone label but different values are BOTH preserved.
    dates = normalized_facts.get("dates", [])
    seen_dates: set[tuple] = set()
    key_dates = []
    for d in dates:
        if not isinstance(d, dict):
            continue
        milestone = (d.get("milestone") or "").strip()
        date_val  = d.get("date", "")
        src_doc   = None
        for sref in (d.get("source_refs") or []):
            if isinstance(sref, dict) and sref.get("source_doc"):
                src_doc = sref["source_doc"]
                break
        # If no source_refs, use source_doc field directly
        if src_doc is None:
            src_doc = d.get("source_doc")
        canon = (milestone, str(date_val), src_doc)
        if not milestone or canon in seen_dates:
            continue
        seen_dates.add(canon)
        key_dates.append({
            "milestone":  milestone,
            "date":       date_val,
            "source_doc": src_doc,
        })
    brief["key_dates"] = key_dates

    # ── commercial_structure — normalized commercial_clauses ──────────────────
    # Canonical key: (topic, details, first source_doc)
    comm_clauses = normalized_facts.get("commercial_clauses", [])
    seen_comm: set[tuple] = set()
    comm_struct = []
    for cc in comm_clauses:
        if not isinstance(cc, dict):
            continue
        topic   = (cc.get("topic") or "").strip()
        details = cc.get("details")
        src_doc = None
        for sref in (cc.get("source_refs") or []):
            if isinstance(sref, dict) and sref.get("source_doc"):
                src_doc = sref["source_doc"]
                break
        if src_doc is None:
            src_doc = cc.get("source_doc")
        canon = (topic, str(details), src_doc)
        if not topic or canon in seen_comm:
            continue
        seen_comm.add(canon)
        comm_struct.append({
            "topic":   topic,
            "details": details,
        })
    brief["commercial_structure"] = comm_struct

    # ── contract_risks — normalized contract_risks ────────────────────────────
    # Canonical key: (risk_title, severity, details) — no invented default severity
    risks = normalized_facts.get("contract_risks", [])
    seen_risks: set[tuple] = set()
    risk_list = []
    for risk in risks:
        if not isinstance(risk, dict):
            continue
        risk_title = (risk.get("risk") or risk.get("title") or "").strip()
        severity   = risk.get("severity")          # None if not present — never invent
        details    = risk.get("details") or risk.get("description")
        canon      = (risk_title, str(severity), str(details))
        if not risk_title or canon in seen_risks:
            continue
        seen_risks.add(canon)
        entry = {
            "risk":    risk_title,
            "details": details,
        }
        if severity is not None:
            entry["severity"] = severity
        risk_list.append(entry)
    brief["contract_risks"] = risk_list

    # ── deliverables_summary — normalized deliverables ────────────────────────
    # Canonical key: (title, description, category, first source_doc)
    deliverables = normalized_facts.get("deliverables", [])
    seen_deliv: set[tuple] = set()
    deliv_summary = []
    for d in deliverables:
        if not isinstance(d, dict):
            continue
        title   = (d.get("title") or d.get("item") or "").strip()
        desc    = d.get("description") or d.get("details")
        cat     = d.get("category")    # None if not present — never invent
        src_doc = None
        for sref in (d.get("source_refs") or []):
            if isinstance(sref, dict) and sref.get("source_doc"):
                src_doc = sref["source_doc"]
                break
        if src_doc is None:
            src_doc = d.get("source_doc")
        canon = (title, str(desc), str(cat), src_doc)
        if not title or canon in seen_deliv:
            continue
        seen_deliv.add(canon)
        entry = {
            "title":       title,
            "description": desc,
        }
        if cat is not None:
            entry["category"] = cat
        deliv_summary.append(entry)
    brief["deliverables_summary"] = deliv_summary

    result = dict(synth_data)
    result["brief"] = brief
    return result


def synthesize_bid_brief(normalized_facts: dict, conflicts: list[dict], api_key: str) -> dict:
    """Stage D only: complete projection, bounded retry, strict authoritative output."""
    from stage_d_checkpoints import checkpoint_run
    with checkpoint_run() as checkpoint:
        return _synthesize_projected_bid_brief(normalized_facts, conflicts, api_key, checkpoint)


def _synthesize_projected_bid_brief(normalized_facts, conflicts, api_key, checkpoint):
    import copy
    from stage_d_projection import (
        build_stage_d_synthesis_projection, finalize_stage_d_request,
        validate_stage_d_response, ProjectionValidationError, AUTHORITATIVE_SECTIONS,
    )
    try:
        projection = build_stage_d_synthesis_projection(normalized_facts, conflicts)
    except ProjectionValidationError as exc:
        checkpoint.write("stage-d/validation.json", {"status": "FAILED", "code": exc.code})
        raise
    checkpoint.write("stage-d/projection.json", projection["prompt_context"])
    checkpoint.write("stage-d/sidecar.json", projection["sidecar"])
    correction = ""
    client = None
    for attempt in (1, 2):
        finalized = finalize_stage_d_request(projection, STAGE_D_SYNTHESIS_PROMPT + correction,
                                            guard_chars=_STAGE_D_CONTEXT_CHAR_LIMIT)
        checkpoint.write("stage-d/size-diagnostics.json", finalized["diagnostics"])
        checkpoint.write("stage-d/request.txt", finalized["request_text"], text=True)
        checkpoint.write("stage-d/output-config.json", finalized["output_config"])
        prefix = f"stage-d/attempt-{attempt:02d}"
        checkpoint.write(prefix + "/size-diagnostics.json", finalized["diagnostics"])
        checkpoint.write(prefix + "/request.txt", finalized["request_text"], text=True)
        checkpoint.write(prefix + "/output-config.json", finalized["output_config"])
        if not finalized["diagnostics"]["dispatch_allowed"]:
            checkpoint.write(prefix + "/validation.json", {"status": "BLOCKED", "code": "CONTEXT_TOO_LARGE"})
            raise StageDContextTooLargeError(
                "Stage D context is complete but too large for safe synthesis. "
                "No requirements were silently omitted. "
                f"Source requirement count: {projection['diagnostics']['requirement_count']}. "
                f"Serialized context size: {finalized['diagnostics']['total_prompt_chars']:,} characters "
                f"(limit: {_STAGE_D_CONTEXT_CHAR_LIMIT:,} characters)."
            )
        try:
            if client is None:
                # Disable SDK retries only for D; a single loop owns the budget.
                client = get_anthropic_client(api_key=api_key).with_options(max_retries=0)
            response = client.messages.create(
                model="claude-haiku-4-5-20251001", max_tokens=8000,
                output_config=finalized["output_config"],
                messages=[{"role": "user", "content": [{"type": "text", "text": finalized["request_text"]}]}],
            )
            text = "".join(block.text for block in response.content if block.type == "text")
            checkpoint.write(prefix + "/response.json", {"text": text, "stop_reason": response.stop_reason})
            if response.stop_reason != "end_turn":
                raise ProjectionValidationError("INCOMPLETE_MODEL_RESPONSE")
            canonical = normalized_facts.get("_canonical_opportunity")
            if canonical and canonical.get("observations"):
                from canonical_opportunity import remove_authoritative_values_for_validation
                from stage_d_projection import strict_json, canonical_json
                text = canonical_json(remove_authoritative_values_for_validation(strict_json(text), canonical))
            validated = validate_stage_d_response(text, projection)
        except (ProjectionValidationError, anthropic.APIError) as exc:
            code = exc.code if isinstance(exc, ProjectionValidationError) else type(exc).__name__
            if isinstance(exc, anthropic.APIStatusError):
                # Private checkpoint only: preserve provider contract diagnostics, never headers.
                checkpoint.write(prefix + "/provider-error.json", {"status_code": exc.status_code, "body": exc.body})
            checkpoint.write(prefix + "/validation.json", {"status": "FAILED", "code": code})
            if attempt == 2:
                raise
            correction = "\nPrevious attempt rejected: " + code + ". Return complete valid JSON with valid field citations."
            continue
        checkpoint.write(prefix + "/validation.json", validated)
        # Always use a copy of the original authoritative inputs, never the prompt projection.
        result = apply_stage_d_authoritative_sections(validated["synthesis"], copy.deepcopy(normalized_facts))
        canonical = normalized_facts.get("_canonical_opportunity")
        if canonical and canonical.get("observations"):
            from canonical_opportunity import apply_authoritative_values
            result = apply_authoritative_values(result, canonical)
        expected = apply_stage_d_authoritative_sections({}, copy.deepcopy(normalized_facts))
        if any(result["brief"].get(k) != expected["brief"][k] for k in AUTHORITATIVE_SECTIONS):
            checkpoint.write(prefix + "/assembly-validation.json", {"status": "FAILED", "code": "AUTHORITATIVE_INVARIANT"})
            raise ProjectionValidationError("AUTHORITATIVE_INVARIANT")
        if normalized_facts != projection["sidecar"]["authoritative_inputs"]["normalized_facts"] or conflicts != projection["sidecar"]["authoritative_inputs"]["conflicts"]:
            raise ProjectionValidationError("AUTHORITATIVE_INPUT_MUTATION")
        checkpoint.write("stage-d/synthesis-result.json", result)
        return result


# ── MAIN ORCHESTRATOR ─────────────────────────────────────────────────────────

def extract_procurement_package(package_files: list[tuple[str, bytes]], api_key: str) -> tuple[dict, str]:
    from stage_d_checkpoints import checkpoint_run
    with checkpoint_run(package_files) as checkpoint:
        return _extract_procurement_package(package_files, api_key, checkpoint)


def _extract_procurement_package(package_files, api_key, checkpoint):
    """
    Genuine 4-Stage Ingestion & Extraction Pipeline:
    1. STAGE A: Document Fact Extraction with Deterministic Markers.
    2. STAGE B: Package Normalization & Strict Provenance Validation.
    3. STAGE C: Cross-Document Reconciliation (6 Conflict Classes).
    4. STAGE D: Executive Bid Brief Synthesis from Normalized Facts Model.
    """
    # Step 0: Preprocess and extract text & metadata for each document
    package_metadata = {
        "files": [f[0] for f in package_files],
        "doc_metadata": {},
        "doc_texts": {}
    }
    
    extracted_docs = []
    for fname, fbytes in package_files:
        doc_text, dmeta = extract_document_with_metadata(fbytes, fname)
        package_metadata["doc_metadata"][fname] = dmeta
        package_metadata["doc_texts"][fname] = doc_text
        extracted_docs.append((fname, doc_text))

    if checkpoint.mode != "off":
        from stage_d_checkpoints import pack_preprocessed
        checkpoint.write("preprocessed-inputs.json", pack_preprocessed(package_metadata))

    # STAGE A: Extract facts per document
    doc_facts_list = []
    for index, (fname, doc_text) in enumerate(extracted_docs, 1):
        facts = extract_document_facts(doc_text, fname, api_key)
        doc_facts_list.append(facts)
        checkpoint.write(f"stage-a/document-{index:04d}.json", facts)
    checkpoint.write("stage-a/document-facts.json", doc_facts_list)

    # STAGE B: Package Normalization & Provenance Validation
    normalized_facts = normalize_package_facts(doc_facts_list, package_metadata)
    checkpoint.write("stage-b/normalized-facts.json", normalized_facts)
    checkpoint.write("stage-b/canonical-observation-ledger.json", normalized_facts.get("_canonical_opportunity"))

    # STAGE C: Reconciliation & Conflict Analysis
    conflicts = reconcile_package_facts(normalized_facts, package_metadata["files"])
    checkpoint.write("stage-c/conflicts.json", conflicts)
    checkpoint.write("stage-c/canonical-opportunity.json", normalized_facts.get("_canonical_opportunity"))

    # STAGE D: Executive Bid Brief Synthesis
    synth_output = synthesize_bid_brief(normalized_facts, conflicts, api_key)

    result = _assemble_procurement_result(synth_output, normalized_facts, conflicts)
    checkpoint.write("stage-d/final-result.json", result)
    return result, "claude-haiku-4-5-20251001 (Staged Pipeline A->B->C->D)"


def _assemble_procurement_result(synth_output, normalized_facts, conflicts):
    # Existing final assembly; normalized records and deterministic helpers own data.
    import copy
    from stage_d_projection import ProjectionValidationError
    original_facts, original_conflicts = copy.deepcopy(normalized_facts), copy.deepcopy(conflicts)
    bid = synth_output.get("bid", {})
    brief = dict(synth_output.get("brief", {}))
    outline = synth_output.get("outline", [])

    # Ensure document_conflicts is attached to brief
    brief["document_conflicts"] = conflicts

    # Build concrete submission documents from normalized submission rules.
    # Only rules representing explicit submission components are included.
    # An empty result is a valid state — it means the procurement package
    # did not establish concrete submission documents in the normalized facts.
    # The old default of Technical Proposal.pdf / Financial Envelope.pdf is
    # removed: those names were invented, not sourced from the procurement package.
    documents = build_submission_documents(
        normalized_facts.get("submission_rules", []),
        submission_deadline=bid.get("submission_deadline"),
    )

    result = {
        "bid": bid,
        "brief": brief,
        "requirements": normalized_facts.get("requirements", []),
        "documents": documents,
        "outline": outline
    }

    expected_documents = build_submission_documents(
        original_facts.get("submission_rules", []),
        submission_deadline=bid.get("submission_deadline"),
    )
    if (result["requirements"] != original_facts.get("requirements", [])
            or result["documents"] != expected_documents
            or result["brief"]["document_conflicts"] != original_conflicts
            or normalized_facts != original_facts or conflicts != original_conflicts):
        raise ProjectionValidationError("FINAL_ASSEMBLY_INVARIANT")
    return result


# Legacy Single-File Compatibility Alias
def extract_rfp(file_bytes: bytes, filename: str, api_key: str) -> tuple:
    """Backward-compatible single file extraction invoking staged package ingestion."""
    return extract_procurement_package([(filename, file_bytes)], api_key)
